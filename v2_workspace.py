"""V2 evaluation workspace. UI actions reuse an immutable calculation run."""
import datetime as dt
import html
import io
import json
import zipfile
import hashlib
import pandas as pd
import streamlit as st
import workspace_app as W
from valuation import legacy as L
from valuation.case import Case, inspect_case, compare_cases
from valuation.service import calculate, refresh_run, calculation_key, export_bundle, AMOUNT_LABELS
from valuation.presentation import label, issue_rows
from valuation.explain import VERSION, event_rows, export_blockers, safe_filename
from ui_format import dataframe

STAGES=['입력·시장자료','평가·분석','계산내역','조서 출력']
STAGE_NAMES={'입력·시장자료':'01  입력','평가·분석':'02  평가결과','계산내역':'03  계산내역','조서 출력':'04  조서 출력'}
CSS='''<style>
.stApp{background:#F4F6FA;color:#182238}
[data-testid="stHeader"]{background:transparent;pointer-events:none}
[data-testid="stHeader"] button{pointer-events:auto}
[data-testid="stAppDeployButton"]{display:none}
[data-testid="stMainBlockContainer"]{max-width:1440px;padding:4.2rem 2.2rem 4rem}
[data-testid="stSidebar"]{background:#f9faff;border-right:1px solid #e0e4ee;min-width:232px;max-width:270px}
[data-testid="stSidebar"] p,[data-testid="stSidebar"] label,[data-testid="stSidebar"] h3{color:#202939}
[data-testid="stSidebar"] [data-testid="stCaptionContainer"] p{color:#58657b}
[data-testid="stSidebar"] [role="radiogroup"] label{padding:8px 10px;border-radius:8px}
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked){background:#ebe7ff;box-shadow:inset 3px 0 #6651d6}
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) p{color:#382777;font-weight:650}
h1{font-size:1.85rem!important;letter-spacing:-.04em}h2{font-size:1.4rem!important}h3{font-size:1.12rem!important}
[data-testid="stMetric"]{background:white;padding:20px;border:1px solid #e0e5ef;border-radius:12px}
[data-testid="stMetricValue"]{font-variant-numeric:tabular-nums;font-size:1.8rem}
[data-testid="stExpander"]{background:white;border:1px solid #dce2ee;border-radius:10px}
[data-testid="stDataFrame"]{background:white;border-radius:8px}
.v2-hero{background:white;border:1px solid #dce2ee;border-radius:14px;padding:24px;margin:8px 0 14px}
.v2-hero .eyebrow{color:#526176;font-size:14px}.v2-hero .amount{font-size:clamp(25px,3.5vw,36px);font-weight:750;color:#272152;font-variant-numeric:tabular-nums;overflow-wrap:anywhere}
.v2-hero .unit{font-size:15px;color:#526176;font-weight:400}.v2-status{font-size:13px;color:#526176;margin-bottom:10px}
.v2-kicker{font-size:12px;letter-spacing:1.6px;color:#5145cd;font-weight:700}.v2-title{font-size:19px;font-weight:700;margin:0}
.v2-formula{background:#f0eefb;padding:18px;border-radius:8px;line-height:1.9;overflow-wrap:anywhere;font-variant-numeric:tabular-nums}
button{min-height:42px}button:focus-visible,input:focus-visible{outline:3px solid #5145cd!important;outline-offset:2px}
@media(max-width:768px){[data-testid="stMainBlockContainer"]{padding:4rem .9rem 3rem}.v2-hero{padding:18px}[data-testid="stMetricValue"]{font-size:1.3rem}}
</style>'''


def goto(stage, topic=None):
    st.session_state['_workflow_stage']=stage
    if topic:st.session_state['_calc_topic']=topic


def sample_case():
    from dataclasses import asdict
    from valuation.case import import_legacy
    t=L.Terms(inst='RCPS',d_issue='2026-01-01',d_base='2026-06-30',d_mat='2031-01-01',
        S0=48000.,K0=60000.,issue_px=60000.,face_total=6_000_000_000.,sig=.35,
        par=500.,cpn=.02,div_basis=0,div_mode=0,ipay=12.,mat_mode=1,ytm=.06,ytm_cmp=1,
        p_mode='accrue',p_yield=.06,p_cmp=1,p_s=24.,p_e=59.,p_f=12.,rfx_mode=0,
        issuer_call=0,k_w=0.,cv_s=1.,cv_e=60.,view='holder',gap_m=1.,
        rf_curve=[[.25,.025],[1,.027],[3,.03],[10,.032]],cr_curve=[[.25,.07],[1,.075],[3,.08],[10,.085]])
    c=import_legacy(asdict(t),'샘플기업 제1회 RCPS · 합성사례')
    c.exercise_styles={'cv':'any','p_f':'periodic'}
    c.sources={k:'사용성 시험을 위한 합성 가정 · 실제 시장자료 아님' for k in ('S0','sig','rf_curve','cr_curve')}
    return c


def sidebar():
    with st.sidebar:
        st.markdown('### VALUATION\nRCPS WORKSPACE')
        st.caption('평가 · 계산근거 · 조서')
        case=st.session_state.get('case')
        if case:
            st.markdown('**현재 평가**')
            st.write(case.name)
            st.radio('평가 진행',STAGES,format_func=STAGE_NAMES.get,key='_workflow_stage')
        with st.expander('평가파일 열기·새로 만들기',expanded=case is None):
            upload=st.file_uploader('평가파일 불러오기',type='json')
            if upload:
                digest=hashlib.sha256(upload.getvalue()).hexdigest()
                if st.session_state.get('upload_digest')!=digest:
                    try:
                        W.install_case(W.read_case(upload.getvalue(),upload.name.removesuffix('.json')))
                        st.session_state.upload_digest=digest
                        st.rerun()
                    except (ValueError,TypeError) as exc:st.error(str(exc))
            with st.form('new_case'):
                name=st.text_input('평가 건명')
                instrument=st.selectbox('평가 상품',['RCPS','CB','BW','SHA'])
                if st.form_submit_button('빈 입력안 만들기'):
                    if not name.strip():st.error('평가 건명을 입력하십시오.')
                    else:
                        W.install_case(Case(name=name.strip(),contract=dict(inst=instrument,rfx_mode=0,cpn=0.,cv_s=99.,cv_e=0.,p_s=99.,p_e=0.,k_w=0.,sha_put_s=99.,sha_put_e=0.,sha_call_s=99.,sha_call_e=0.),method=dict(model='TF',view='holder',gap_m=1.)))
                        st.rerun()
            if st.button('합성사례로 시작하기',use_container_width=True):
                W.install_case(sample_case());st.rerun()
        previous=None
        with st.expander('전기 평가와 비교'):
            upload=st.file_uploader('전기 평가파일(선택)',type='json',key='previous_upload')
            if upload:
                try:previous=W.read_case(upload.getvalue(),upload.name)
                except (ValueError,TypeError) as exc:st.error(str(exc))
            if previous and st.button('전기 입력을 새 평가로 복사'):
                previous.name+=' — 갱신';W.install_case(previous);st.rerun()
        case=st.session_state.get('case')
        if case:
            st.download_button('평가파일 저장',json.dumps(case.to_dict(),ensure_ascii=False,indent=2),
                '평가입력_v2.json','application/json',use_container_width=True)
            st.caption('현재 세션에 반영됨 · 서버에 영구 보관되지 않습니다. 종료 전 파일로 저장해 주세요.')
            with st.expander('기존 버전과 비교용 파일'):
                data=case.to_dict();data.pop('memo_context',None)
                st.download_button('기존 앱 호환 입력 저장',json.dumps(data,ensure_ascii=False,indent=2),'평가입력_비교용.json','application/json')
                st.caption('계산 입력은 같고 새 메모의 조건 식별정보만 제외합니다. 원본 파일은 별도로 보관하세요.')
    return previous


def context(case,run,current):
    inst=case.contract.get('inst','')
    st.markdown('<div class="v2-kicker">VALUATION / '+html.escape(inst)+'</div>',unsafe_allow_html=True)
    st.markdown('<div class="v2-title">'+html.escape(case.name)+'</div>',unsafe_allow_html=True)
    v=case.effective()
    st.caption(f"{v.get('d_base','기준일 미입력')} 기준 · {'투자자' if v.get('view')=='holder' else '발행자'} 관점 · {v.get('model','TF')} 모형")
    if run and not current:
        st.warning('입력이 변경되었습니다. 아래는 이전 결과입니다. 현재 입력으로 다시 평가해야 조서를 만들 수 있습니다.')
    elif run:
        st.markdown(f'<div class="v2-status">현재 입력으로 계산됨 · 평가 {run.summary["run_id"]} · {run.summary["calculation_seconds"]:.2f}초</div>',unsafe_allow_html=True)


def formula_box(text):
    st.markdown('<div class="v2-formula">'+html.escape(text).replace('\n','<br>')+'</div>',unsafe_allow_html=True)


def component_evidence(run):
    t,s=run.terms,run.summary
    if L.is_sha(t):
        W.sha_result_panel(run);return
    a=s['amounts_100'];factor=t.face_total/100
    st.subheader('총액과 구성요소가 연결되는 방법')
    st.caption('동일 실행의 원금 100 기준 값에서 환산합니다. 각 권리를 순차적으로 추가한 차액이며 회계상 인식액과 구별합니다.')
    formula_box(f"주계약 {a['host_reference']:,.2f}\n+ 상환권 증분 {a['put_increment']:,.2f}\n+ 전환권 증분 {a['conversion_increment']:,.2f}\n− 콜 영향 {a['call_deduction']:,.2f}\n= 순포지션 {a['net']:,.2f}")
    formula_box(f"총액 = {a['net']:,.2f} × {t.face_total:,.2f} ÷ 100\n= {s['amounts_total']['net']:,.2f} 원")
    if s['amounts_per_share']:
        st.write(f"주당금액 = 원금 100 기준 × 1주당 발행가 {t.issue_px:,.2f} ÷ 100 = {s['amounts_per_share']['net']:,.2f}원")
    st.caption('표시값은 반올림됩니다. 합계는 반올림 전 원값으로 계산합니다. 음의 증분도 임의로 0으로 바꾸지 않습니다.')
    st.button('전체 노드 계산표 보기',on_click=goto,args=('계산내역','노드 계산표'),key='root_jump')
    st.caption('조서: 기본 값은 평가요약, 상세 수식은 결과 및 계산대사. 생성한 파일의 실제 연결 위치는 조서 출력 후 표시됩니다.')


def calculation_panel(run):
    topics=['노드 계산표','노드 스케줄','부트스트래핑','구성요소','행사·지급일정']
    old=st.session_state.get('_calc_topic')
    if old not in topics:st.session_state['_calc_topic']='노드 스케줄' if old=='금리·할인계수' else topics[0]
    topic=st.radio('계산내역 선택',topics,horizontal=True,key='_calc_topic')
    if topic=='구성요소':component_evidence(run)
    elif topic in topics[:3]:
        from valuation.calculation_view import node_values, schedule_values, bootstrap_values, worksheet_html
        import streamlit.components.v1 as components
        if L.is_sha(run.terms) and topic!='부트스트래핑':
            W.sha_result_panel(run);return
        if topic=='노드 계산표':
            st.caption('전체 격자를 표로 확인하고 셀을 눌러 수식·참조값을 여세요. 본체의 가치이며 별도 콜 차감액은 구성요소에서 확인합니다.')
            # Preserve older sessions safely when a new valuation has fewer steps.
            for key in ('trace_i','trace_j'):
                if key in st.session_state:st.session_state[key]=min(st.session_state[key],run.terms.n)
            start,end=0,run.terms.n
            if run.terms.n>250:
                spans=list(range(0,run.terms.n+1,32))
                start=st.selectbox('계산표 시점 구간',spans,format_func=lambda i:f'{i}–{min(i+31,run.terms.n)} 시점',key='tree_span_'+run.summary['run_id'])
                end=min(start+31,run.terms.n)
                st.caption('큰 격자는 32시점씩 표시합니다. 구간을 바꾸면 만기까지 모든 노드를 확인할 수 있습니다.')
            payload=node_values(run,start,end)
        elif topic=='노드 스케줄':
            payload=schedule_values(run)
        else:
            payload=bootstrap_values(run)
            for k in ('rf_curve','cr_curve'):st.caption(f'{label(k)} 출처: {run.case.sources.get(k,"미기록")}')
        if hasattr(st,'iframe'):
            st.iframe(worksheet_html(payload),height=730)
        else:
            components.html(worksheet_html(payload),height=730,scrolling=True)
    else:
        st.subheader('행사금액과 지급가치')
        dataframe(pd.DataFrame(event_rows(run)),hide_index=True,use_container_width=True)
        W.dp_panel(run)
        st.caption('격자일과 계약상 날짜는 다를 수 있습니다. 상세 조서의 계약일 목록·행사일 대조표에서 배정 규칙을 함께 확인하십시오.')


def result_panel(run,case,current,pending):
    s,t=run.summary,run.terms
    units=['총액 · 원']+(['주당 · 원'] if s['amounts_per_share'] else [])+['원금 100']
    unit=st.radio('표시 단위',units,horizontal=True,key='result_unit')
    vals=s['amounts_total'] if unit==units[0] else s['amounts_per_share'] if unit=='주당 · 원' else s['amounts_100']
    key='put' if L.is_sha(t) else 'net'
    amount=f'{vals[key]:,.2f}'
    st.markdown(f'<div class="v2-hero"><div class="eyebrow">{AMOUNT_LABELS[key]}'+(' · 콜 차감 후' if key=='net' else '')+f'</div><div class="amount">{amount} <span class="unit">{unit}</span></div></div>',unsafe_allow_html=True)
    st.button('계산표·수식 보기 →',on_click=goto,args=('계산내역','노드 계산표'),type='primary')
    c1,c2,c3=st.columns(3)
    precision=2
    c1.metric('본체 · 콜 차감 전' if key=='net' else '콜 가치',f"{vals.get('whole_before_call',vals.get('call',0)):,.{precision}f}")
    c2.metric('콜 차감' if key=='net' else '주당 기준가격(원)',f"{vals.get('call_deduction',t.K0):,.{precision}f}")
    c3.metric('계산 구간',f'{t.n:,}');c3.caption(f"평균 {s['grid']['average_days']:.2f}일")
    if L.is_sha(t):W.sha_result_panel(run)
    table=[{'구성요소':AMOUNT_LABELS[k],'금액':vals[k]} for k in ('host_reference','put_increment','conversion_increment','call_deduction','net') if k in vals]
    if table:
        st.subheader('평가금액 구성')
        dataframe(pd.DataFrame(table),hide_index=True,use_container_width=True,column_config={'금액':st.column_config.NumberColumn(format='%,.2f')})
        st.caption('권리를 순차적으로 추가한 가치 차이입니다. 회계상 인식액은 회계 참고에서 확인합니다.')
    notices=[i for i in run.issues if i.code not in {'source','engine_defaults','legacy_defaults','judgement_scope','market_date','coverage_review'}]
    blocked=export_blockers(run)
    for message in blocked:st.warning(message)
    with st.expander(f'확인할 사항 {len(notices)}건 · 계산 점검',expanded=bool(blocked)):
        if notices:dataframe(pd.DataFrame(issue_rows(notices)),hide_index=True,use_container_width=True)
        else:st.write('이 실행에서 표시할 수치 경고가 없습니다.')
        dataframe(pd.DataFrame(s['checks']),hide_index=True,use_container_width=True)
    with st.expander('행사·지급일정'):
        dataframe(pd.DataFrame(event_rows(run)),hide_index=True,use_container_width=True)
        W.dp_panel(run)
    with st.expander('민감도 · 기준 평가 유지'):
        variable=st.selectbox('민감도 변수',['S0','sig','rf_curve','cr_curve'],format_func=label)
        magnitude=st.number_input('변화폭(주당가치는 %, 변동성·금리는 %p)',min_value=.01,max_value=50. if variable=='S0' else 10.,value=10. if variable=='S0' else 1.)
        if st.button('민감도 계산',disabled=not current or pending):
            from valuation.analysis import sensitivity
            try:
                with st.spinner('기준 입력을 보존하고 감소·증가 조건을 계산합니다.'):
                    st.session_state.analysis=sensitivity(run,variable,magnitude)
            except (ValueError,ArithmeticError) as exc:st.error(str(exc))
        a=st.session_state.get('analysis')
        if a and a['calculation_key']==s['calculation_key'] and current:
            dataframe(pd.DataFrame(a['rows']).rename(columns=AMOUNT_LABELS),hide_index=True,use_container_width=True)
            st.caption(f"계산된 변수 {label(a['variable'])} · 변화폭 ±{a['change']:.2f}. 기준 평가와 입력은 유지됩니다.")
    with st.expander('회계 참고 · 선택한 분류 가정'):
        st.caption('평가금액의 분해와 회계상 인식은 구별합니다. K-IFRS 참고표이며 일반기업회계기준·자기신용 OCI 등 미지원 처리는 별도 검토합니다.')
        W.day1_panel(run,case)
        W.split_panel(run)
    if current:
        with st.expander('기존 상세 분석 도구'):
            mode=st.selectbox('분석 도구',['결과 요약','상세 계산·회계 참고표'])
            if mode=='상세 계산·회계 참고표':
                from application import detailed
                detailed(run)
    with st.expander('평가 정보'):
        st.write({'실행번호':s['run_id'],'문서 개정':s['document_id'],'버전':VERSION,'계산시각':s['calculated_at'],'입력 식별값':s['case_sha256']})


def export_panel(run,case,current,pending,previous):
    st.header('조서 출력')
    st.caption('검토 목적에 맞게 선택하세요. 결과와 입력은 같은 평가 실행으로 묶습니다.')
    option=st.radio('조서 구성',['기본 값 조서','상세 계산 값 조서','상세 계산 수식 조서'],horizontal=True)
    explanations={'기본 값 조서':'권장 · 요약과 주요 근거를 읽고 대사·보관할 때 사용합니다. 입력을 바꿔도 재계산되지 않습니다.',
        '상세 계산 값 조서':'전체 노드와 계약 일정을 확인합니다. 생성 당시의 값으로 고정됩니다.',
        '상세 계산 수식 조서':'수식을 추적하고 직접 입력한 주가·변동성을 바꿔 재계산합니다. 날짜·금리·권리조건·간격·모형 등 다른 입력은 앱에서 변경 후 다시 생성해야 합니다.'}
    st.info(explanations[option])
    accounting=st.checkbox('회계처리·분개·상각표 포함 (초안)',value=False)
    judgment=st.checkbox('평가자 메모 포함',value=True)
    st.caption('분리 비교 등 수치 계산근거는 메모 선택과 관계없이 상세 조서에 포함됩니다. 원문 저작물은 포함하지 않습니다.')
    blocked=export_blockers(run) if run else []
    for message in blocked:st.warning(message)
    if not current or pending:st.warning('현재 입력으로 평가한 결과가 필요합니다. 입력 오류를 해결하고 다시 평가해 주세요.')
    if run:st.caption(f"평가 {run.summary['run_id']} · {run.terms.n:,}구간 · {run.terms.d_base} 기준")
    if st.button('조서 생성',type='primary',disabled=not current or pending or bool(blocked)):
        st.session_state.pop('bundle',None)
        try:
            with st.status('조서 생성 중',expanded=True) as status:
                st.write('입력과 결과의 일치·출력 조건을 확인합니다.')
                data=export_bundle(run,formula=option=='상세 계산 수식 조서',detail=option!='기본 값 조서',previous=previous,accounting=accounting,judgment=judgment)
                st.session_state.bundle=data
                st.session_state.bundle_key=(case.fingerprint(),option,accounting,judgment,previous.fingerprint() if previous else None)
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    from openpyxl import load_workbook
                    name=next(n for n in z.namelist() if n.endswith('.xlsx'))
                    wb=load_workbook(io.BytesIO(z.read(name)),read_only=True)
                    st.session_state.workbook_locations={name:True for name in wb.sheetnames};wb.close()
                    st.session_state.locations_run=run.summary['run_id']
                status.update(label='조서 생성 완료',state='complete',expanded=False)
        except (ValueError,ArithmeticError,OSError,RuntimeError,MemoryError) as exc:
            st.session_state.pop('bundle',None)
            st.error('조서를 생성하지 못했습니다. 입력과 기존 결과는 유지됩니다. '+str(exc))
    key=(case.fingerprint(),option,accounting,judgment,previous.fingerprint() if previous else None)
    if current and not pending and not blocked and st.session_state.get('bundle_key')==key and st.session_state.get('bundle'):
        data=st.session_state.bundle
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            name=next(n for n in z.namelist() if n.endswith('.xlsx'))
            st.download_button('Excel 조서만 저장',z.read(name),safe_filename(run,option,'xlsx'),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
        st.download_button('평가 조서 묶음 저장',data,safe_filename(run,option,'zip'),'application/zip')
        st.success('조서 생성 완료 · 같은 실행의 Excel·입력·적용값·시장자료·결과 기록을 생성했습니다.')
        st.caption('다운로드 파일은 기기에 저장됩니다. 이 앱은 서버 영구 보관을 제공하지 않습니다.')


def main():
    if st.session_state.get('_workflow_stage') not in STAGES:st.session_state['_workflow_stage']=STAGES[0]
    previous=sidebar()
    case=st.session_state.get('case')
    if case is None:
        st.title('평가에서 근거 확인까지, 한 작업 공간에서')
        st.write('RCPS · CB · BW · 주주간계약의 입력, 평가, 계산 추적과 Excel 조서를 연결합니다.')
        st.info('왼쪽에서 기존 평가파일을 열거나 합성사례로 새 화면을 체험하세요.')
        st.caption('동료 사용 시험에는 합성·익명 사례를 사용하세요. 평가파일을 저장하면 다음에 이어서 작업할 수 있습니다.')
        return
    issues=inspect_case(case);errors=[x for x in issues if x.severity=='error']
    run=st.session_state.get('run');pending=st.session_state.get('_input_pending',False)
    current=run is not None and not errors and run.summary['calculation_key']==calculation_key(case)
    context(case,run,current)
    stage=st.session_state['_workflow_stage']
    if stage=='입력·시장자료':
        st.header('입력')
        area=st.radio('입력 항목',['계약·평가 입력','주가·변동성·금리 자료','출처·평가가정'],horizontal=True,key='_input_area')
        if area=='계약·평가 입력':W.input_editor(case,autosave=True)
        elif area=='주가·변동성·금리 자료':
            from market_tools_ui import main as market
            market(case)
        else:
            W.evidence_editor(case)
            if L.dp_active(L.Terms(**case.effective())):
                with st.form('dp_missing_policy'):
                    reason=st.text_area('미입력 발생연도의 상환재원 가정·근거',case.sources.get('dp_missing_assumption',''),help='비워 두면 해당 연도가 실제 지급에 사용될 때 조서가 차단됩니다. 미입력 연도를 제한 없이 지급 가능한 것으로 가정할 근거가 있을 때만 기록하세요.')
                    if st.form_submit_button('재원 가정 저장'):
                        candidate=Case.from_dict(case.to_dict());candidate.sources['dp_missing_assumption']=reason.strip();W.save_case(candidate)
        st.button('평가결과로 이동 →',on_click=goto,args=('평가·분석',),type='primary')
        return
    if stage=='평가·분석':
        st.header('평가결과')
        if errors:dataframe(pd.DataFrame(issue_rows(errors)),hide_index=True,use_container_width=True)
        if pending:st.warning('입력화면에 반영하지 못한 값이 있습니다. 해당 항목을 확인해 주세요.')
        if st.button('현재 입력으로 평가',type='primary',disabled=bool(errors) or pending):
            try:
                with st.status('현재 입력을 평가합니다.',expanded=True) as status:
                    st.write('계약조건·금리곡선을 확인하고 가치와 구성요소를 계산합니다.')
                    new=refresh_run(run,case) if current else calculate(case)
                    st.session_state.run=new
                    st.session_state.pop('bundle',None);st.session_state.pop('analysis',None)
                    status.update(label='평가 완료',state='complete',expanded=False)
                st.rerun()
            except (ValueError,ArithmeticError) as exc:
                st.error('평가를 완료하지 못했습니다. 입력과 이전 결과는 보존됩니다. '+str(exc))
        if run:result_panel(run,case,current,pending)
        else:st.info('현재 입력으로 평가하면 결과와 계산근거가 표시됩니다.')
        if previous:
            with st.expander('전기 대비 입력 변경'):
                dataframe(pd.DataFrame(compare_cases(previous,case)),hide_index=True,use_container_width=True)
    elif stage=='계산내역':
        st.header('계산내역')
        if run:calculation_panel(run)
        else:st.info('평가를 먼저 실행해 주세요. 계산하지 않은 금액은 표시하지 않습니다.')
        st.button('← 평가결과로',on_click=goto,args=('평가·분석',))
    else:export_panel(run,case,current,pending,previous)
