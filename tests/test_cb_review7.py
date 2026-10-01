"""전환사채 평가조서 검토(예약 7) 반영 — 가상 수치.

· 조기상환권을 주계약에 두면 상각표가 첫 조기상환 가능일에 끝난다. 그날 이자는 «행사일 이자 별도지급» 을 따른다
  (행사금액에 포함이면 마지막 회차 이자 0) — 격자의 조기상환 갈래와 같다. 값 조서·수식 조서가 같은 값이다.
"""
import io
import os
import shutil
import subprocess
import tempfile
import zipfile
import pytest
from openpyxl import load_workbook
from valuation import legacy
from valuation.case import import_legacy
from valuation.service import calculate, export_bundle

SOFFICE = shutil.which("libreoffice") or shutil.which("soffice")
BASE = dict(inst="CB", conv_class="liability", emb_approach=2, p_sep=0, k_sep=0, model="TF",
            d_issue="2025-07-03", d_base="2025-07-03", d_mat="2029-07-03", S0=1300., K0=1500., floor=1050.,
            cpn=.02, ipay=3., ytm=.07, ytm_cmp=4, cv_s=12., cv_e=47., rfx_mode=0,
            p_s=12., p_e=27., p_f=3., p_mode="accrue", p_yield=.07, p_cmp=4, k_w=0.,
            sig=.5, rf_curve=[[1, .023], [3, .025]], cr_curve=[[1, .06], [3, .088]], face_total=2e9, gap_m=1.)


@pytest.mark.parametrize("add", [1, 0])
def test_expected_maturity_last_coupon(add):
    run = calculate(import_legacy(dict(BASE, p_cpn_add=add), "가상 전환사채"))
    t = run.terms
    assert legacy.eir_expect(t) is not None                     # 첫 조기상환 가능일이 기대만기
    eir = legacy.eir_or_none(t, *[run.raw[k] for k in ("full", "b0", "b1", "b2", "ca")])
    rows = eir[1]
    c = 100*legacy.eff_cpn(t)*t.ipay/12
    assert rows[-1][4] == pytest.approx(c if add else 0.0)
    assert rows[-1][5] == pytest.approx(eir[2], abs=1e-6)       # 기말 = 그날 행사금액
    if not SOFFICE: return
    z = zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=True, detail=True, accounting=True)))
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "f.xlsx"); open(src, "wb").write(z.read("formula_review.xlsx"))
        out = os.path.join(d, "o"); os.mkdir(out)
        subprocess.run([SOFFICE, "-env:UserInstallation=file://" + os.path.join(d, "p"), "--headless",
                        "--convert-to", "xlsx", "--outdir", out, src], capture_output=True, timeout=900)
        M = load_workbook(os.path.join(out, "f.xlsx"), data_only=True)["상각표"]
    last = 15 + len(rows) - 1
    assert M.cell(last, 7).value == pytest.approx(rows[-1][4], abs=1e-9)
    assert M["C10"].value == pytest.approx(eir[0], abs=1e-9)
    assert M.cell(last, 8).value == pytest.approx(rows[-1][5], abs=1e-6)
