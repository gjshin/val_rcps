"""Exercise semantics, one-case UI, privacy and contract-review regressions."""
import copy
import datetime as dt
import hashlib
import io
import json
import math
import zipfile
from dataclasses import asdict
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from test_v2_workflow import synthetic, ROOT
from test_practical_controls import prices
from valuation import legacy, controls
from valuation.case import Case, inspect_case
from valuation.service import calculate, calculation_key, refresh_run
from valuation.report import basic_workbook
from valuation.bridge import apply_volatility
from valuation.market_data import query_spec, make_pack
from valuation.evidence import evidence_cards, judgment_record


def draft_review():
    return dict(documents=[dict(name='Synthetic contract', sha256=hashlib.sha256(b'synthetic').hexdigest())],
                findings=[dict(document='Synthetic contract', clause='Article 1 / page 1', quote='Principal is 100.',
                               fields=['face_total'], topics=['classification'], interpretation='Synthetic review only', action='Confirm signature')],
                open_items=['Confirm signed original'])


@pytest.mark.parametrize('gap', [.25, 1., 3., 12.])
@pytest.mark.parametrize('inst', ['RCPS', 'CB', 'BW', 'SHA'])
def test_anytime_put_call_remain_open_at_every_step(gap, inst):
    case = synthetic()
    case.contract.update(inst=inst, p_s=0., p_e=12., k_s=0., k_e=12.,
                         sha_put_s=0., sha_put_e=12., sha_call_s=0., sha_call_e=12.)
    case.method['gap_m'] = gap
    case.exercise_styles = {k:'any' for k in ['p_f','k_f','sha_put_f','sha_call_f']}
    terms = legacy.derive(legacy.Terms(**case.effective()))
    # Independent count: one opportunity per node, even for the minimum 4-node grid.
    assert terms.p_f == pytest.approx(terms.T * 12 / terms.n)
    assert terms.k_f == terms.p_f
    if inst == 'SHA':
        raw = legacy.sha_engine(terms)
        assert all(raw['p_on'](i) and raw['c_on'](i) for i in range(terms.n+1))
    else:
        raw = legacy.engine(terms)
        assert all(raw['kstrike'](i) is not None for i in range(terms.n+1))
        tree = legacy.bdt_parts(terms)
        assert all(tree['in_put'](i) for i in range(terms.n+1))


def test_anytime_immediate_redemption_has_hand_computed_value():
    case = synthetic()
    case.contract.update(mat_mode=1, cv_s=99., cv_e=0., p_s=0., p_e=12., p_mode='fixed', p_rate=120., cpn=0.)
    case.exercise_styles = {'p_f':'any'}
    # With positive discounting and every date paying 120, immediate exercise wins.
    assert calculate(case).summary['amounts_100']['net'] == pytest.approx(120.)
    case.contract.update(issuer_call=1, k_w=1., k_s=0., k_e=12., k_prem=0., p_s=99., p_e=0., cv_s=0., cv_e=12.)
    case.market['S0'] = 100.
    case.exercise_styles = {'k_f':'any'}
    assert calculate(case).summary['amounts_100']['net'] == pytest.approx(100.)


def test_styles_roundtrip_boundaries_conflict_and_key():
    case = synthetic(); baseline = calculation_key(case)
    case.exercise_styles = {'k_f':'any', 'cv':'single'}
    case.contract.update(k_s=3., k_e=9., cv_s=6., cv_e=12.)
    assert Case.from_dict(case.to_dict()).to_dict() == case.to_dict()
    assert case.effective()['cv_e'] == 6.
    assert calculation_key(case) != baseline
    terms = legacy.derive(legacy.Terms(**case.effective()))
    raw = legacy.engine(terms)
    assert raw['kstrike'](0) is None and raw['kstrike'](terms.n) is None
    case.contract['k_sched'] = '6 100'
    assert any(i.code == 'exercise_conflict' for i in inspect_case(case))
    case.exercise_styles = {'unknown':'any'}
    with pytest.raises(ValueError): Case.from_dict(case.to_dict())


def test_metadata_does_not_reprice_but_changes_review_key():
    case = synthetic(); run = calculate(case)
    case.judgments['classification'] = judgment_record(case, '검토 완료','Conclusion','Basis','Reviewer','Clause')
    case.contract_review = draft_review()
    assert calculation_key(case) == run.summary['calculation_key']
    updated = refresh_run(run, case)
    assert updated.raw is run.raw
    assert next(c for c in evidence_cards(case) if c['id']=='classification')['review_stale']
    assert any(x['code']=='contract_open_items' for x in controls.blockers(updated))
    wb = load_workbook(io.BytesIO(basic_workbook(updated)))
    assert wb['계약검토안'].cell(2,3).value == 'Principal is 100.'
    assert wb['추가확인자료'].cell(2,1).value == 'Confirm signed original'


@pytest.mark.parametrize('mutation', [
    lambda r:r['findings'][0].update(fields=['S0; run command']),
    lambda r:r['findings'][0].update(topics=['unknown']),
    lambda r:r['findings'][0].update(quote=''),
    lambda r:r['documents'][0].update(sha256='bad'),
    lambda r:r.update(open_items=[42]),
])
def test_malformed_contract_review_rejected(mutation):
    case = synthetic(); case.contract_review = draft_review(); mutation(case.contract_review)
    with pytest.raises(ValueError): Case.from_dict(case.to_dict())


def test_volatility_can_apply_before_other_inputs_and_base_date():
    query = query_spec(False, ['TEST,Example'], 30, '2025-01-01')
    pack = make_pack([['Example', prices()]], query, tdays=250, drop=False, pick='median', source='Synthetic', retrieved_at='2025-01-01T00:00:00+00:00')
    case = Case('Incomplete', contract={'inst':'RCPS'})
    updated = apply_volatility(case, pack)
    assert updated.market['sig'] == pack['sigma']
    assert 'S0' not in updated.market and 'd_base' not in updated.method
    updated.method['d_base'] = '2025-01-02'
    with pytest.raises(ValueError, match='기준일'): apply_volatility(updated, pack)


def app_with_case(case=None):
    app = AppTest.from_file(str(ROOT/'app.py'), default_timeout=30)
    app.session_state.case = case or synthetic()
    return app.run()


def test_all_steps_and_market_tools_open_without_pricing_or_network(monkeypatch):
    def forbidden(*a, **kw): raise AssertionError('Read-only step invoked calculation/network')
    for name in ['engine','decompose','sha_engine','fetch_prices','fetch_close','backsolve']:
        monkeypatch.setattr(legacy, name, forbidden)
    app = app_with_case()
    original = copy.deepcopy(app.session_state.case.to_dict())
    for step in ['조서 출력','평가·분석','입력·시장자료']:
        app.radio(key='_workflow_stage').set_value(step).run()
        assert not app.exception
    app.radio(key='_input_area').set_value('주가·변동성·금리 자료').run()
    for tool in ['기준일 주가 조회','주가 역산','금리곡선 불러오기','BDT 금리 변동성','변동성 산출']:
        next(w for w in app.selectbox if w.label=='시장자료 도구').set_value(tool).run()
        assert not app.exception
    assert app.session_state.case.to_dict() == original
    assert 'run' not in app.session_state


def test_volatility_apply_updates_same_case_and_visible_field(monkeypatch):
    monkeypatch.setattr(legacy,'fetch_prices',lambda *a: (prices(), 'Synthetic adjusted close'))
    app = app_with_case()
    app.session_state.case.market.pop('S0')
    app.session_state.case.market.pop('rf_curve')
    app.radio(key='_input_area').set_value('주가·변동성·금리 자료').run()
    next(w for w in app.text_area if w.label.startswith('피어 목록')).set_value('TEST,Example')
    next(w for w in app.number_input if w.label=='조회 일수 (거래일)').set_value(30)
    app.run()
    next(b for b in app.button if b.label=='주가 받기').click().run()
    next(b for b in app.button if b.label=='이 변동성을 현재 평가에 적용').click().run()
    assert not app.exception
    sigma = app.session_state.case.market['sig']
    assert sigma == app.session_state.case.market_evidence['sig']['sigma']
    assert 'S0' not in app.session_state.case.market
    app.radio(key='_input_area').set_value('계약·평가 입력').run()
    assert next(w for w in app.number_input if w.label=='주가 변동성(연, %)').value == pytest.approx(sigma*100)


def test_anytime_ui_persists_after_grid_change():
    case = synthetic()
    case.contract.update(p_s=0.,p_e=12.,issuer_call=1,k_w=1.,k_s=0.,k_e=12.)
    app = app_with_case(case)
    for title in ['상환청구 방식','콜 행사 방식']:
        next(w for w in app.selectbox if w.label==title).set_value('any').run()
    next(w for w in app.number_input if w.label=='계산 간격(개월)').set_value(1.).run()
    assert app.session_state.case.exercise_styles == {'p_f':'any','k_f':'any'}
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    assert not app.exception
    terms = app.session_state.run.terms
    assert terms.p_f == pytest.approx(12*terms.T/terms.n)
    assert terms.k_f == terms.p_f


def test_new_rcps_rights_stay_checked_until_both_date_ranges_are_entered():
    case = Case(name='New RCPS', contract=dict(inst='RCPS', d_issue='2025-03-07',
        d_mat='2035-03-06', rfx_mode=0, cpn=0., cv_s=99., cv_e=0.,
        p_s=99., p_e=0., k_w=0.),
        method=dict(d_base='2026-06-30', model='TF', view='holder', gap_m=1.))
    app = app_with_case(case)
    for title, detail in [('전환·신주인수권 행사 가능', '전환·신주인수권 행사 시작일'),
                          ('투자자 상환청구권 있음', '투자자 상환청구 시작일')]:
        next(w for w in app.checkbox if w.label == title).check().run()
        assert not app.exception
        assert next(w for w in app.checkbox if w.label == title).value
        assert any(w.label == detail for w in app.get('date_input'))
    for title, day in [('전환·신주인수권 행사 시작일', dt.date(2025, 3, 8)),
                       ('전환·신주인수권 행사 종료일', dt.date(2035, 3, 6)),
                       ('투자자 상환청구 시작일', dt.date(2028, 3, 8)),
                       ('투자자 상환청구 종료일', dt.date(2035, 3, 5))]:
        next(w for w in app.get('date_input') if w.label == title).set_value(day).run()
    assert not app.exception
    assert all(next(w for w in app.checkbox if w.label == title).value for title in
               ('전환·신주인수권 행사 가능', '투자자 상환청구권 있음'))
    assert app.session_state.case.contract['cv_s'] < app.session_state.case.contract['cv_e']
    assert app.session_state.case.contract['p_s'] < app.session_state.case.contract['p_e']


def test_shared_detail_basic_views_reuse_run_without_transfer(monkeypatch):
    app = app_with_case()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    case = copy.deepcopy(app.session_state.case.to_dict())
    def forbidden(*a, **kw): raise AssertionError('Existing result was repriced')
    for name in ['engine','decompose','sha_engine','validate']:
        monkeypatch.setattr(legacy, name, forbidden)
    next(w for w in app.selectbox if w.label=='분석 도구').set_value('상세 계산·회계 참고표').run()
    for section in ['이자율곡선','주가·변동성','의사결정']:
        next(w for w in app.selectbox if w.label=='상세 분석 항목').set_value(section).run()
        assert not app.exception
    assert app.session_state.case.to_dict() == case
    assert not any('가져오기' in b.label or '상세 입력으로 계산' == b.label for b in app.button)


def test_bdt_detail_does_not_decide_model_from_exercise_ratio_alone():
    case = synthetic()
    case.contract.update(p_s=0., p_e=12.)
    app = app_with_case(case)
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    next(w for w in app.selectbox if w.label=='분석 도구').set_value('상세 계산·회계 참고표').run()
    next(w for w in app.selectbox if w.label=='상세 분석 항목').set_value('판단·근거').run()
    assert not app.exception
    messages = [str(w.value) for kind in ('info','warning','success') for w in app.get(kind)]
    # 행사 진단은 앱 판정(초안)으로만 표시하고 BDT 를 켜라고 지시하지 않는다.
    assert any('행사금액 ÷ 계속보유가치' in m and '앱 판정(초안)' in m for m in messages)
    assert 'BDT 를 켜십시오' not in ' '.join(messages)


def test_detailed_exports_are_available_without_changing_grid():
    app = app_with_case()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    original_key = app.session_state.run.summary['calculation_key']
    app.radio(key='_workflow_stage').set_value('조서 출력').run()
    for option in ['상세 계산 수식 조서', '상세 계산 값 조서']:
        next(w for w in app.radio if w.label=='조서 구성').set_value(option).run()
        assert not next(b for b in app.button if b.label=='조서 생성').disabled
        next(b for b in app.button if b.label=='조서 생성').click().run()
        assert not app.exception
        assert app.session_state.run.summary['calculation_key'] == original_key
    next(w for w in app.radio if w.label=='조서 구성').set_value('기본 값 조서').run()
    next(b for b in app.button if b.label=='조서 생성').click().run()
    assert not app.exception
    assert any('조서 생성 완료' in w.value for w in app.success)
    assert app.session_state.run.summary['calculation_key'] == original_key
    with zipfile.ZipFile(io.BytesIO(app.session_state.bundle)) as z:
        assert 'value_review.xlsx' in z.namelist()


def test_contract_review_screen_has_no_document_uploader_or_ai_endpoint():
    app = app_with_case()
    assert app.radio(key='_workflow_stage').options == ['입력·시장자료','평가·분석','조서 출력']
    assert not app.exception
    labels = [u.label for u in app.get('file_uploader')]
    assert labels == ['평가파일 불러오기', '전기 평가파일(선택)']
    assert not any(x in (ROOT/'contract_ui.py').read_text() for x in ['api.openai.com', 'anthropic.com'])


def test_new_third_party_call_uses_agreed_initial_method():
    app = app_with_case()
    next(w for w in app.selectbox if w.label=='RCPS 콜 권리').set_value(2).run()
    assert not app.exception
    assert next(w for w in app.selectbox if w.label=='콜 평가방법').value == 2
    assert next(w for w in app.selectbox if w.label=='콜 행사가액 분해방법').value == 1
    next(w for w in app.number_input if w.label=='콜 대상 비율(%)').set_value(30.).run()
    assert not app.exception
    assert app.session_state.case.method['k_method'] == 2
    assert app.session_state.case.method['k_split'] == 1


def test_anytime_refixing_uses_every_step_and_keeps_label():
    case = synthetic(); case.method['gap_m'] = .7; case.exercise_styles = {'rfx_cyc':'any'}
    terms = legacy.derive(legacy.Terms(**case.effective()))
    assert round(terms.rfx_cyc*terms.n/(terms.T*12)) == 1
    assert legacy.rfx_any(terms)


@pytest.mark.parametrize('formula', [False,True])
def test_detailed_workpapers_omit_contract_review_keep_exercise_mode(formula):
    import zipfile
    from valuation.service import export_bundle
    case = synthetic(); case.exercise_styles = {'p_f':'any'}; case.contract_review = draft_review()
    data = export_bundle(calculate(case), detail=True, formula=formula)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        name = 'formula_review.xlsx' if formula else 'value_review.xlsx'
        wb = load_workbook(io.BytesIO(z.read(name)))
        assert '계약검토안' not in wb
        assert json.loads(z.read('case.json'))['contract_review'] == case.contract_review
        info = [[c.value for c in row] for row in wb['조서 정보']]
        at = next(i for i, r in enumerate(info) if r[0] == '행사방식')
        assert info[at+2][1] == '기간 중 언제든지'
