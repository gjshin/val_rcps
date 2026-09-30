"""엑셀에서 입력을 바꾼 뒤 다시 계산하면 앱과 같은 값이 나오는가.

조서대조.py 는 «앱에서 입력을 바꿔 조서를 새로 만든 뒤» 견준다. 이 시험은 반대 방향이다 —
조서를 한 번 만든 뒤 가정 시트의 노란 칸(주가 · 변동성 · 표면이자율 · 보장수익률 등)을
**엑셀 안에서** 바꾸고 다시 계산한 값을, 같은 입력으로 앱이 새로 계산한 값과 견준다.
수식 조서의 칸 하나라도 값으로 굳어 있으면(입력을 따라오지 않으면) 여기서 걸린다.

견주는 칸은 재계산 검사 도구와 같은 목록이다 (legacy.formula_key_cells — 결과 시트의
전체·주계약·부채요소·조기상환청구권·매도청구권·전환권대가와 상각표 유효이자율).

재계산은 기본으로 formulas(파이썬 계산기)로 한다. ``--libreoffice`` 를 주면 LibreOffice(Calc)로
다시 계산한다 — 둘 다 Microsoft Excel 이 아니므로 실제 Excel 에서의 확인은 따로 해야 한다.

    python3 tests/엑셀입력변경.py                   # 전부 (formulas)
    python3 tests/엑셀입력변경.py 부채               # 이름에 «부채» 가 든 경우만
    python3 tests/엑셀입력변경.py --libreoffice      # LibreOffice 로 재계산
"""
import copy
import importlib.util
import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_spec = importlib.util.spec_from_file_location("josu", os.path.join(ROOT, "tests", "조서대조.py"))
J = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(J)

# (이름, 계약 설정, 엑셀에서 바꿀 칸 {가정 시트 B열 라벨: (Terms 항목, 새 값)})
CASES = [
    ("CB · 주가 +10% · 변동성 +5%p", {},
     {"평가기준일 주가": ("S0", 853*1.1), "변동성 σ": ("sig", 0.413 + 0.05)}),
    # 전환권이 부채면 주계약이 «100 − 파생상품» 의 잔여라 주가가 바뀌면 상각표 유효이자율도 바뀐다
    ("CB 전환권 부채 · 주가 −15%", {"conv_class": "liability"},
     {"평가기준일 주가": ("S0", 853*0.85)}),
    ("CB 표면 2% → 3% · 만기보장 5% → 6%", {"cpn": 0.02, "ytm": 0.05},
     {"표면이자율": ("cpn", 0.03), "만기보장수익률": ("ytm", 0.06)}),
    # 조기상환금액을 보장수익률로 붙이는 방식(accrue)이어야 그 칸이 계산에 쓰인다
    ("CB 조기상환 보장 4% → 7%", {"p_mode": "accrue", "p_yield": 0.04},
     {"조기상환 보장수익률": ("p_yield", 0.07)}),
    ("RCPS · 주가 +20% · 우선배당 1% → 2%", {"inst": "RCPS", "cpn": 0.01},
     {"평가기준일 주가": ("S0", 853*1.2), "우선배당률 (계약)": ("cpn", 0.02)}),
    ("BW 현금납입 · 변동성 +10%p", {"inst": "BW", "bw_pay": 1},
     {"변동성 σ": ("sig", 0.513)}),
    ("GS · 주가 −10%", {"model": "GS"},
     {"평가기준일 주가": ("S0", 853*0.9)}),
]


def terms(G, over):
    t = G["Terms"](); t.rf_curve = [(1, .0226), (3, .0240), (5, .0252)]
    t.cr_curve = [(1, .1409), (3, .1740), (5, .1905)]
    t.gap_m = 6.0
    for k, v in over.items(): setattr(t, k, v)
    G["derive"](t)
    return t


def app_values(G, t):
    """앱이 그 입력으로 새로 계산한 값 — 수식 조서가 따라와야 할 답."""
    full, b0, b1, b2, ca, conv = G["decompose"](t)
    eir = G["eir_or_none"](t, full, b0, b1, b2, ca)
    raw = dict(full=full, b0=b0, b1=b1, b2=b2, ca=ca, conv=conv)
    return G["formula_key_cells"](t, raw, eir), (full, b0, b1, b2, ca, conv, eir)


def solve_libreoffice(path, sheets_cells):
    """LibreOffice 로 다시 계산해 {시트: {칸: 값}} — 값이 없는 수식 칸은 불러올 때 계산된다."""
    import shutil, subprocess
    import openpyxl
    from pathlib import Path
    exe = os.environ.get("VALUATION_SOFFICE") or shutil.which("libreoffice") or shutil.which("soffice")
    if not exe:
        raise SystemExit("LibreOffice 가 없습니다 — --libreoffice 없이 돌리십시오.")
    base = Path(path).parent; out = base/"lo"; out.mkdir(exist_ok=True)
    r = subprocess.run([exe, "-env:UserInstallation=" + (base/"profile").as_uri(), "--headless",
                        "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                       capture_output=True, text=True, timeout=1800)
    tgt = out/Path(path).name
    if r.returncode or not tgt.exists():
        raise SystemExit("LibreOffice 재계산 실패: " + (r.stderr or r.stdout)[-800:])
    wb = openpyxl.load_workbook(tgt, data_only=True)
    got = {}
    for sh, cells in sheets_cells.items():
        got[sh] = {}
        for c in cells:
            v = wb[sh][c].value if sh in wb.sheetnames else None
            if isinstance(v, (int, float)) and not isinstance(v, bool): got[sh][c] = float(v)
    return got


def main():
    import openpyxl
    G = J.load_app()
    use_lo = "--libreoffice" in sys.argv
    args = [a for a in sys.argv[1:] if a != "--libreoffice"]
    only = args[0] if args else ""
    bad = 0
    for lbl, over, change in CASES:
        if only and only not in lbl: continue
        t0 = terms(G, over)
        _, (full, b0, b1, b2, ca, conv, eir) = app_values(G, t0)
        x = G["build_xlsx_formula"](t0, full, b0, b1, b2, ca, conv, eir)
        # 같은 입력으로 앱이 새로 계산한 답
        t1 = copy.deepcopy(t0)
        for k, (attr, v) in change.items(): setattr(t1, attr, v)
        G["derive"](t1)
        want, _ = app_values(G, t1)
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "wb.xlsx")
            open(path, "wb").write(x)
            wb = openpyxl.load_workbook(path)
            A = wb["가정"]
            rows = {c.value: c.row for (c,) in A.iter_rows(min_col=2, max_col=2)
                    if isinstance(c.value, str)}
            miss = [k for k in change if k not in rows]
            if miss:
                print(f"\n[{lbl}]  ★ 가정 시트에 없는 칸: {miss}"); bad += 1; continue
            for k, (_, v) in change.items():
                A.cell(row=rows[k], column=3).value = v
            if use_lo:
                # LibreOffice 는 시트 이름·COMBIN 을 그대로 계산한다 — 바꾸지 않고 넘긴다
                wb.save(path)
                mp = {s_: s_ for s_ in wb.sheetnames}
                need = {}
                for _, sh, cell, _ in want: need.setdefault(sh, []).append(cell)
                got = solve_libreoffice(path, need)
            else:
                # 조서대조와 같은 방식으로 시트 이름을 짧게 바꾸고 COMBIN 을 값으로 푼다
                mp = {nm: f"S{i:02d}" for i, nm in enumerate(wb.sheetnames)}
                for ws in wb.worksheets:
                    for row in ws.iter_rows():
                        for c in row:
                            if isinstance(c.value, str) and c.value.startswith("="):
                                f = c.value
                                for o, nn in sorted(mp.items(), key=lambda z: -len(z[0])):
                                    f = f.replace(f"'{o}'!", f"{nn}!").replace(f"{o}!", f"{nn}!")
                                c.value = J._combin(f)
                for o, nn in mp.items(): wb[o].title = nn
                wb.save(path)
                got = J.solve(path, tuple(sorted({mp[s_] for _, s_, _, _ in want if s_ in mp})))
        _report(lbl, change, want, got, mp)
        bad += _count_bad(want, got, mp)
    print("\n" + ("모든 항목 일치 — 엑셀에서 바꾼 입력이 결과까지 따라온다" if bad == 0
                  else f"★ {bad}건 불일치")
          + (" (LibreOffice 재계산)" if use_lo else " (formulas 재계산)"))
    return 0 if bad == 0 else 1


def _ok(h, w):
    return h is not None and abs(h - w) <= 1e-7*max(1.0, abs(w))


def _count_bad(want, got, mp):
    return sum(not _ok(got.get(mp.get(sh, ""), {}).get(cell), w) for _, sh, cell, w in want)


def _report(lbl, change, want, got, mp):
    print(f"\n[{lbl}]  엑셀에서 바꾼 칸: " + ", ".join(f"{k} → {v:g}" for k, (_, v) in change.items()))
    for name, sheet, cell, w in want:
        h = got.get(mp.get(sheet, ""), {}).get(cell)
        fmt = (lambda v: f"{v*100:.6f}%") if "유효이자율" in name else (lambda v: f"{v:.6f}")
        print("   %-24s 엑셀 %14s · 앱 %14s  %s"
              % (name, fmt(h) if h is not None else "없음", fmt(w), "" if _ok(h, w) else "★"))


if __name__ == "__main__":
    sys.exit(main())
