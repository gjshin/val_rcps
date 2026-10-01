"""콜 권리자 — 전환사채·신주인수권부사채 콜의 첫 선택 (가상 수치).

«발행회사 본인만 / 지정 가능 / 사전 특정» 하나를 고르면 저장 칸(제3자 지정 · 콜 유형)과 기본 평가방법·회계처리가
정해지고, 회계처리 문장·조서 가정 시트·분리 판단이 그 선택을 따른다.
"""
import io
import zipfile
import pytest
from openpyxl import load_workbook
from valuation import legacy
from valuation.case import import_legacy
from valuation.service import calculate, export_bundle, CaseError

BASE = dict(inst="CB", conv_class="equity", model="TF", d_issue="2026-01-15", d_base="2026-01-15", d_mat="2029-01-15",
            S0=10000., K0=10000., floor=7000., cpn=0.01, ytm=.03, ytm_cmp=4, ipay=3., cv_s=12., cv_e=35.,
            rfx_mode=1, rfx_cyc=3., p_s=18., p_e=33., p_f=3., p_mode="accrue", p_yield=.03, p_cmp=4,
            k_s=12., k_e=18., k_f=3., k_prem=.03, k_cmp=4, k_w=.35, sig=.5,
            rf_curve=[[1, .025], [3, .027]], cr_curve=[[1, .08], [3, .095]], face_total=2_000_000_000)


def run_for(h, **over):
    return calculate(import_legacy(dict(BASE, **legacy.call_holder_fields(h, "TF"), **over), f"콜 권리자 {h}"))


@pytest.mark.parametrize("h", [0, 1, 2])
def test_round_trip(h):
    f = legacy.call_holder_fields(h, "TF")
    assert legacy.call_holder(f) == h
    assert f["k_sep"] == (0 if h == 0 else 1)
    assert f["k_method"] == (0 if h == 0 else 2)
    assert legacy.call_holder_fields(h, "GS")["k_method"] == 0      # GS 는 옵션차익법을 쓰지 않는다


def test_prespecified_without_third_blocked():
    with pytest.raises(CaseError):
        calculate(import_legacy(dict(BASE, k_third=0, k_kind=1, k_sep=1), "모순"))


@pytest.mark.parametrize("h,word", [(0, "발행회사만 행사하는 매도청구권은 거래상대방이 그대로인 내재파생상품"),
                                    (1, "발행회사가 제3자를 지정할 수 있어"),
                                    (2, "발행 시 정해진 제3자의 매도청구권")])
def test_accounting_sentence_and_assumption(h, word):
    run = run_for(h)
    for formula in (False, True):
        z = zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=formula, detail=True, accounting=True)))
        wb = load_workbook(io.BytesIO(z.read("formula_review.xlsx" if formula else "value_review.xlsx")))
        txt = str(wb["회계처리"]["B3"].value)
        assert word in txt and ("제3자에게 이전될 수 있어" not in txt)
        if formula:
            A = wb["가정"]
            assert any(A.cell(r, 3).value == legacy.CALL_HOLDERS[h] for r in range(1, A.max_row + 1))


def test_issuer_only_split_judgement():
    run = run_for(0)
    t = run.terms
    sp = legacy.split_test(t, *[run.raw[k] for k in ("full", "b0", "b1", "b2", "ca")], [])
    assert sp["call"]["결론"] != "별도의 금융상품"          # 발행회사만 → 내재파생 판정
    assert sp["call"]["설정일치"]                          # 기본 회계처리(내재파생 포함)와 맞는다


def test_no_call_sentence_empty():
    t = legacy.Terms(k_w=0.0)
    assert legacy.call_alloc_note(t) == ""


def test_screen_call_holder_first_choice():
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    root = Path(__file__).resolve().parents[1]
    case = import_legacy(dict(BASE, **legacy.call_holder_fields(1, "TF")), "콜 권리자 화면")
    app = AppTest.from_file(str(root / "app.py"), default_timeout=60)
    app.session_state["case"] = case
    app.run()
    assert not app.exception
    sel = next(w for w in app.selectbox if w.label == "콜 권리자")
    assert sel.value == 1
    sel.set_value(0).run()
    assert not app.exception
    # 발행회사 본인만 → 기본 평가방법 유무가치비교법 · 회계처리 복합내재파생에 포함
    meth = next(w for w in app.selectbox if w.label.startswith("콜") and "평가방법" in w.label)
    assert meth.value == 0
    assert not any(w.label.startswith("제3자 콜 유형") for w in app.selectbox)   # 상세 조건에서 빠졌다
