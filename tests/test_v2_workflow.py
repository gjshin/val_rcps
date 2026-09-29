"""Contract invariants, rejection paths and shared UI/CLI evidence tests."""
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile
from dataclasses import asdict

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from valuation import legacy
from valuation.case import Case, import_legacy, inspect_case, compare_cases
from valuation.service import CaseError, calculate, export_bundle, input_rows


def synthetic():
    terms = legacy.Terms(inst="RCPS", d_issue="2025-01-01", d_base="2025-01-01",
        d_mat="2026-01-01", gap_m=3, issue_px=100., S0=40., K0=100., sig=.25,
        rfx_mode=0, cpn=0., div_mode=1, issuer_call=0, p_s=99., p_e=0.,
        cv_s=0., cv_e=12., k_w=0., rf_curve=[[1., .02], [5., .02]],
        cr_curve=[[1., .1], [5., .1]])
    return import_legacy(asdict(terms), "Synthetic — no client data")


def test_auto_conversion_without_cash_rights_equals_common_share():
    run = calculate(synthetic())
    assert run.summary["amounts_100"]["net"] == pytest.approx(40., abs=1e-10)
    assert any(i.code == "auto_conversion_split" for i in run.issues)
    assert run.summary["status"] == "review_only"


def test_contract_facts_are_unchanged_by_assumptions_and_running():
    case = synthetic()
    case.assumptions = [dict(field="p_s", value=108., rationale="Illustrative forecast constraint")]
    original = copy.deepcopy(case.to_dict())
    run = calculate(case)
    assert case.to_dict() == original
    assert case.contract["p_s"] == 99.
    assert run.terms.p_s == 108.
    case.contract["p_s"] = 200.
    assert run.case.contract["p_s"] == 99.
    row = next(r for r in input_rows(run) if r["항목"] == "p_s")
    assert (row["원본 입력"], row["적용값"]) == (99., 108.)


@pytest.mark.parametrize("group,key,value,code", [
    ("market", "sig", float("nan"), "type"),
    ("market", "S0", -1., "positive"),
    ("market", "rf_curve", [[1., .02], [1., .03]], "curve"),
    ("market", "cr_curve", [[1., float("inf")], [5., .1]], "curve"),
    ("method", "gap_m", 12/52, "legacy_week_grid"),
    ("method", "d_base", "2027-01-01", "date_order"),
    ("method", "d_base", "2024-01-01", "date_order"),
    ("method", "model", "unknown", "enum"),
    ("method", "carry", 42, "enum"),
    ("contract", "k_w", 2., "fraction"),
    ("contract", "mat_mode", "0", "type"),
    ("contract", "S0", 20., "wrong_section"),
])
def test_invalid_inputs_block_calculation(group, key, value, code):
    case = synthetic()
    getattr(case, group)[key] = value
    with pytest.raises(CaseError) as exc:
        calculate(case)
    assert any(i.code == code for i in exc.value.issues)


def test_missing_required_market_input_never_falls_back_to_example():
    case = synthetic()
    del case.market["S0"]
    with pytest.raises(CaseError):
        calculate(case)


def test_unknown_legacy_input_is_not_silently_dropped():
    with pytest.raises(ValueError, match="알 수 없는"):
        import_legacy({"S_zero": 100})


def test_import_records_every_default_and_roundtrips():
    case = import_legacy({"inst": "RCPS"})
    assert "S0" in case.imported_defaults
    assert Case.from_dict(case.to_dict()).to_dict() == case.to_dict()


@pytest.mark.parametrize("treatment", ["exact", "included", "modeled"])
def test_unsupported_right_cannot_claim_exact_modeling(treatment):
    case = synthetic()
    case.additional_rights = [dict(kind="distributable_profit", clause="Synthetic clause",
        treatment=treatment, rationale="test", assumption_fields=[])]
    with pytest.raises(CaseError) as exc:
        calculate(case)
    assert any(i.code == "unsupported_exact" for i in exc.value.issues)


def test_scenario_requires_linked_assumption_and_retains_review_issue():
    case = synthetic()
    case.additional_rights = [dict(kind="distributable_profit", clause="Synthetic clause",
        treatment="scenario", rationale="Forecast approximation", assumption_fields=["p_s"])]
    with pytest.raises(CaseError):
        calculate(case)
    case.assumptions = [dict(field="p_s", value=108., rationale="Forecast approximation")]
    assert any(i.code == "additional_right" for i in calculate(case).issues)


def test_assumption_requires_unique_key_and_rationale():
    case = synthetic()
    case.assumptions = [dict(field="p_s", value=108., rationale="")]*2
    assert {"missing_rationale", "duplicate_assumption"} <= {i.code for i in inspect_case(case)}


def test_prior_comparison_separates_contract_market_and_assumption_changes():
    old, new = synthetic(), synthetic()
    new.contract["K0"] = 99.
    new.market["S0"] = 42.
    new.assumptions = [dict(field="p_s", value=108., rationale="test")]
    assert {"contract", "market", "assumptions"} == {r["section"] for r in compare_cases(old, new)}


@pytest.mark.parametrize("formula", [False, True])
def test_export_keeps_input_evidence_and_manifest(formula):
    case = synthetic()
    case.sources["S0"] = '=HYPERLINK("https://invalid.example","untrusted")'
    run = calculate(case)
    data = export_bundle(run, formula=formula, previous=synthetic())
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        manifest = json.loads(z.read("manifest.json"))
        assert all(hashlib.sha256(z.read(k)).hexdigest() == v for k, v in manifest.items())
        assert json.loads(z.read("case.json")) == case.to_dict()
        assert json.loads(z.read("result.json"))["case_sha256"] == case.fingerprint()
        assert "changes.json" in z.namelist()
        name = "formula_review.xlsx" if formula else "value_review.xlsx"
        wb = load_workbook(io.BytesIO(z.read(name)), data_only=False)
        assert wb.sheetnames[:3] == (["V2_검토기록", "V2_계약과가정", "V2_추가권리"] if formula else ["평가요약", "계약조건", "시장자료"])
        sheet = wb['V2_계약과가정'] if formula else wb['시장자료']
        cells = [cell for row in sheet for cell in row if cell.value == case.sources["S0"]]
        assert cells and all(cell.data_type == "s" for cell in cells)


def test_formula_export_never_silently_switches_refixing_method():
    case = synthetic()
    case.contract.update(rfx_mode=1, floor=20., rfx_cyc=3.)
    case.method["carry"] = 0
    run = calculate(case)
    with pytest.raises(ValueError, match="동일한 수식"):
        export_bundle(run, formula=True)


def test_cli_and_service_use_identical_calculation_and_evidence(tmp_path):
    case = synthetic()
    source, output = tmp_path / "case.json", tmp_path / "result.zip"
    source.write_text(json.dumps(case.to_dict()), encoding="utf-8")
    process = subprocess.run([sys.executable, str(ROOT / "tools/run_case.py"), str(source), "--out", str(output)],
                             cwd=tmp_path, capture_output=True, text=True)
    assert process.returncode == 0, process.stderr
    with zipfile.ZipFile(output) as z:
        result = json.loads(z.read("result.json"))
    direct = calculate(case).summary
    assert result["amounts_100"] == direct["amounts_100"]
    assert result["case_sha256"] == direct["case_sha256"]
    assert result["code_sha256"] == direct["code_sha256"]


def test_engine_can_import_without_streamlit():
    script = "import sys; from valuation import legacy; assert 'streamlit' not in sys.modules"
    proc = subprocess.run([sys.executable, "-c", script], cwd=ROOT, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr


@pytest.mark.parametrize("instrument", ["CB", "BW", "SHA"])
def test_other_instruments_keep_their_existing_path_and_export(instrument):
    case = synthetic()
    case.contract["inst"] = instrument
    run = calculate(case)
    if instrument == "SHA":
        reference = legacy.sha_engine(copy.deepcopy(run.terms))
        assert run.summary["amounts_100"]["put"] == reference["put"]
        assert run.summary["amounts_100"]["call"] == reference["call"]
    else:
        _, _, _, whole, call, _ = legacy.decompose(copy.deepcopy(run.terms))
        assert run.summary["amounts_100"]["net"] == whole-call
    with zipfile.ZipFile(io.BytesIO(export_bundle(run))) as z:
        assert z.testzip() is None


@pytest.mark.parametrize("filename", ["app.py", "vol_app.py"])
def test_existing_app_entrypoints_still_start(filename):
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(str(ROOT / filename), default_timeout=30).run()
    assert not app.exception


def test_workspace_loads_and_runs_shared_service():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(str(ROOT / "workspace_app.py"), default_timeout=30)
    app.session_state["case"] = synthetic()
    app.run()
    assert not app.exception
    app.radio(key="_workflow_stage").set_value("평가·분석").run()
    button = next(b for b in app.button if b.label == "현재 입력으로 평가")
    button.click().run()
    assert not app.exception
    assert app.session_state["run"].summary["amounts_100"]["net"] == pytest.approx(40.)
    # A saved edit retains the previous snapshot, clearly marked stale, and blocks export.
    app.radio(key="_workflow_stage").set_value("입력·시장자료").run()
    field = next(w for w in app.number_input if w.label == "현재 전환가액(원)")
    field.set_value(120.).run()
    app.radio(key="_workflow_stage").set_value("평가·분석").run()
    assert not app.exception
    assert app.session_state['run'].case.contract['K0'] == 100.
    assert any('변경 전 입력' in w.value for w in app.warning)
    app.radio(key='_workflow_stage').set_value('조서 출력').run()
    assert next(b for b in app.button if b.label == '조서 생성').disabled


def test_new_case_keeps_market_inputs_blank():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(str(ROOT / "workspace_app.py"), default_timeout=30).run()
    next(w for w in app.text_input if w.label == "평가 건명").set_value("Blank synthetic")
    next(b for b in app.button if b.label == "빈 입력안 만들기").click().run()
    assert not app.exception
    assert "S0" not in app.session_state["case"].market
    assert "rf_curve" not in app.session_state["case"].market
    app.radio(key="_workflow_stage").set_value("평가·분석").run()
    assert next(b for b in app.button if b.label == "현재 입력으로 평가").disabled
