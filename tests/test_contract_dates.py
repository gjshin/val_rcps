"""00 계약일 목록 — 조정일·지급 회수·행사월이 숫자가 아니라 계약일을 노드에 배정한 수식인가 (가상 수치).

00 격자 공통 6행(조정일)·9행(지급 회수)·20·27행(행사월)에 앱이 배정한 결과를 숫자로 넣으면 배정이 맞는지
엑셀에서 따라갈 수 없다. 이제 그 칸은 «00 계약일 목록» 을 찾는 수식이고, LibreOffice 로 다시 계산한 값이
엔진의 refix_steps · pay_steps · exercise_amounts 와 같아야 한다.
"""
import os
import shutil
import subprocess
import tempfile
import pytest
from openpyxl import load_workbook
from valuation import legacy

SOFFICE = shutil.which("libreoffice") or shutil.which("soffice")
CASES = [
    ("발행일 평가 · 분기 이자 · 하향 조정", dict(cpn=.02, ipay=3., rfx_mode=1, rfx_cyc=3.)),
    ("중간 평가 7개월 · 반기 이자 · 상향 조정 5개월", dict(d_base="2025-12-20", cpn=.03, ipay=6., rfx_mode=2, rfx_cyc=5.,
                                                 p_f=6., k_f=2.)),
    ("주 격자 · 최초 조정 12개월 · 이후 7개월", dict(grid_days=7., rfx_first=12., rfx_cyc=7., cpn=.01, ipay=3.)),
]


def build(over):
    t = legacy.Terms(rf_curve=[(1, .0226), (3, .0240), (5, .0252)], cr_curve=[(1, .1409), (3, .1740), (5, .1905)],
                     carry=1, gap_m=1.0, **over)
    legacy.derive(t)
    full, b0, b1, b2, ca, conv = legacy.decompose(t)
    wb = legacy.build_xlsx_formula(t, full, b0, b1, b2, ca, conv, legacy.eir_table(t, b0), as_workbook=True)
    return t, wb


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice 없음")
@pytest.mark.parametrize("name,over", CASES, ids=[c[0] for c in CASES])
def test_assignment_rows_are_formulas_and_match_engine(name, over):
    t, wb = build(over)
    n, dt_ = t.n, t.T/t.n
    C = wb["00 격자 공통"]
    assert "00 계약일 목록" in wb.sheetnames
    for i in range(1, n+1):
        for r in (6, 9):
            v = C.cell(r, 3+i).value
            assert isinstance(v, str) and v.startswith("="), (r, i, v)     # 숫자로 박지 않는다
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "f.xlsx"); wb.save(src); out = os.path.join(d, "o"); os.mkdir(out)
        subprocess.run([SOFFICE, "-env:UserInstallation=file://" + os.path.join(d, "p"), "--headless",
                        "--convert-to", "xlsx", "--outdir", out, src], capture_output=True, timeout=900)
        V = load_workbook(os.path.join(out, "f.xlsx"), data_only=True)["00 격자 공통"]
    rfx = set(legacy.refix_steps(t, n, dt_))
    pay = legacy.pay_steps(t, n, dt_)
    EA = legacy.exercise_amounts(t, n, dt_)
    cpn = 100*legacy.eff_cpn(t)*t.ipay/12
    assert {i for i in range(n+1) if V.cell(6, 3+i).value == 1} == rfx
    for i in range(n+1):
        assert abs((V.cell(9, 3+i).value or 0) - cpn*pay.get(i, 0)) < 1e-9, i
    got_p = {i: V.cell(20, 3+i).value for i in range(n+1) if isinstance(V.cell(20, 3+i).value, (int, float))}
    got_k = {i: V.cell(27, 3+i).value for i in range(n+1) if isinstance(V.cell(27, 3+i).value, (int, float))}
    assert set(got_p) == set(EA["p_dates"]) and all(abs(got_p[i] - EA["p_dates"][i]) < 1e-6 for i in got_p)
    assert set(got_k) == set(EA["k_dates"]) and all(abs(got_k[i] - EA["k_dates"][i]) < 1e-6 for i in got_k)
    for i in got_p:
        assert abs(V.cell(26, 3+i).value - (EA["put"](i) + (cpn*pay.get(i, 0) if (i < n and t.p_cpn_add) else 0))) < 1e-6
