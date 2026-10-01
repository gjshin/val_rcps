"""신주인수권부사채 조서 검토 반영 — 가상 수치로 만든 조서의 표시·잠금·판정 비교.

· 금리격자를 쓰면 금리곡선과 금리변동성은 노란 칸(엑셀 입력)이 아니다 — 기준금리 a 가 고정값이다.
· 입력 칸의 노란색은 한 가지다.
· 분리 판단에 «수치 판정 · 이용자 설정 → 일치 / 검토 필요» 줄이 있고, 수식 조서는 판정 수식에 이어진다.
· 상품 칸은 BW, 회계처리 줄은 «신주인수권대가», 판단 문장의 «전환사채의 전환과 같은» 은 바뀌지 않는다.
· 최초 인식 차이 세 가지 구분 표에 이 평가가 표시된다.
"""
import io
import zipfile
import pytest
from openpyxl import load_workbook
from valuation import legacy
from valuation.case import import_legacy
from valuation.service import calculate, export_bundle


def bw_case(**over):
    obj = dict(inst="BW", bw_pay=1, bw_detach=0, conv_class="equity", model="TF",
               d_issue="2026-01-29", d_base="2026-01-29", d_mat="2031-01-29",
               S0=5000., K0=5200., floor=3640., cpn=0.0, ytm=.03, ytm_cmp=4, ipay=3.,
               cv_s=12., cv_e=59., rfx_mode=2, rfx_cyc=3.,
               p_s=24., p_e=57., p_f=3., p_mode="accrue", p_yield=.03, p_cmp=4,
               k_s=12., k_e=24., k_f=3., k_prem=.03, k_cmp=4, k_w=.30,
               sig=.45, put_bdt=1, bdt_sig=.20,
               rf_curve=[[1, .025], [3, .027], [5, .029]],
               cr_curve=[[1, .08], [3, .095], [5, .105]],
               face_total=10_000_000_000)
    obj.update(over)
    return calculate(import_legacy(obj, "가상 신주인수권부사채"))


def book(run, formula):
    z = zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=formula, detail=True, accounting=True)))
    return load_workbook(io.BytesIO(z.read("formula_review.xlsx" if formula else "value_review.xlsx")))


def find(ws, label, col=2):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise AssertionError(f"{ws.title} 에 «{label}» 줄이 없다")


@pytest.fixture(scope="module")
def run():
    return bw_case()


@pytest.fixture(scope="module")
def fwb(run):
    return book(run, True)


@pytest.fixture(scope="module")
def vwb(run):
    return book(run, False)


def _yellow(ws):
    return {c.fill.fgColor.rgb[-6:] for row in ws.iter_rows() for c in row
            if c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb
            and c.fill.fgColor.rgb[-6:-4] in ("FD", "FF") and c.fill.fgColor.rgb[-2:] < "E8"
            and c.fill.fgColor.rgb[-4:-2] >= "E0"}


def test_bdt_locks_curve_and_sigma(fwb):
    A = fwb["가정"]
    r = find(A, "BDT 변동성 σ (앱에서만 변경)")
    assert A.cell(r, 3).fill.fill_type in (None, "none")
    assert not isinstance(A.cell(r, 3).value, str)          # 산출내역 시트에 잇지 않고 값으로
    I = fwb["IR 입력곡선"]
    assert all(I.cell(8, c).fill.fill_type in (None, "none") for c in (2, 3, 5, 6))
    assert "금리곡선 수익률, BDT 금리변동성 σ" in fwb["해설"]["C6"].value
    assert "금리곡선 수익률 등" not in fwb["해설"]["C5"].value


def test_curve_stays_yellow_without_bdt():
    wb = book(bw_case(put_bdt=0), True)
    assert wb["IR 입력곡선"].cell(8, 3).fill.fgColor.rgb[-6:] == legacy.INPUT_FILL


def test_one_input_colour(fwb):
    seen = set()
    for ws in fwb.worksheets:
        seen |= _yellow(ws)
    assert seen <= {legacy.INPUT_FILL}, seen


def test_labels(fwb):
    A = fwb["가정"]
    assert A.cell(find(A, "상품"), 4).value == "BW"
    E = fwb["회계처리"]
    labels = [str(E.cell(r, 2).value) for r in range(1, E.max_row + 1)]
    assert any("신주인수권대가" in x for x in labels)
    assert not any('"전환권대가' in x for x in labels)
    c2 = fwb["판단·근거"]["C2"].value
    assert "전환사채의 전환과 같은 계산 구조" in c2 and "행사과" not in c2


@pytest.mark.parametrize("formula", [False, True])
def test_split_compare_row(run, formula, fwb, vwb):
    J = (fwb if formula else vwb)["분리 판단"]
    r = find(J, "판정과 설정 비교")
    v = J.cell(r, 3).value
    if formula:
        assert v.startswith("=") and "C" in v                 # 판정 수식에 이어진다
        assert J.cell(find(J, "결론"), 3).value.startswith("=C")
    else:
        assert "분리하지 않을 여지" in v and "검토 필요" in v


def test_split_compare_agree():
    # 주계약에 포함을 고르면 «여지» 와 일치한다
    r = bw_case(p_sep=0); t = r.terms
    sp = legacy.split_test(t, *[r.raw[k] for k in ("full", "b0", "b1", "b2", "ca")], [])
    assert legacy.put_in_host(t)
    assert legacy.split_compare(t, "put", sp["put"]).endswith("→ 일치")


def test_day1_three_cases(fwb):
    E = fwb["회계처리"]
    r = find(E, "최초 인식 차이의 세 가지 구분")
    rows = [(E.cell(r + i, 2).value, E.cell(r + i, 4).value) for i in (1, 2, 3)]
    assert [x[1] for x in rows] == ["◀ 이 평가", None, None]
    cases = legacy.issuer_day1_cases(dict(hybrid=True, pl=False))
    assert [c[2] for c in cases] == ["", "", "◀ 이 평가"]


def test_reference_values_marked(fwb):
    R = fwb["결과"]
    txt = " ".join(str(R.cell(r, 2).value) for r in range(1, R.max_row + 1))
    assert "앱에서 생성 당시 계산한 참고값" in txt


def test_cb_relabel_unchanged():
    # 전환사채 조서에는 신주인수권부사채 예외가 끼어들지 않는다
    assert legacy.relabel_text("전환사채와 같은", []) == "전환사채와 같은"
    assert legacy.relabel_text("전환사채와 같은", legacy._BW_WORDS) == "전환사채와 같은"
    assert legacy.relabel_text("전환사채의 전환권", legacy._BW_WORDS) == "신주인수권부사채의 신주인수권"


def test_split_compare_call_room_is_review():
    # 매도청구권은 어느 설정이든 분리해 처리하므로 «분리하지 않을 여지» 면 검토 필요다
    t = bw_case().terms
    out = legacy.split_compare(t, "call", {"결론": "분리하지 않을 여지", "설정일치": True})
    assert "검토 필요" in out and "지원하지 않음" in out


def test_formula_compare_follows_workbook_setting(fwb):
    # 이용자 설정 글자도 가정 시트의 분류·매도청구권 처리 칸을 보고 정한다 (값으로 굳히지 않는다)
    J = fwb["분리 판단"]
    v = J.cell(find(J, "판정과 설정 비교"), 3).value
    assert "신주인수권과 묶어 분리" in v and "매도청구권과 묶어 분리" in v and "분리 — 파생상품부채" in v
