"""2026.10.07 수정 — 조서 입력 기록 자릿수 · 메모 조건 · 미입력 연도 안내 · 수식 조서 열린 칸 · 표시 자릿수 · 실행번호."""
import io
import re
import zipfile

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from test_commercial_v2 import case
from test_v2_workflow import ROOT
from valuation import legacy as L
from valuation.case import Case
from valuation.explain import KINDS, export_blockers, memo_key, memo_status, _memo_key_v21
from valuation.presentation import display_value, exact_number
from valuation.service import calculate, export_bundle


def _book(run, formula=False):
    z = zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=formula, detail=formula)))
    return load_workbook(io.BytesIO(z.read('formula_review.xlsx' if formula else 'value_review.xlsx')))


def test_workpaper_input_record_is_not_rounded():
    # 2.125% 를 2.12% 로, 32.587% 를 32.59% 로 적으면 조서가 실제 입력과 다른 숫자를 남긴다.
    assert display_value('cpn', .02125) == '2.125%'
    assert display_value('sig', .32587) == '32.587%'
    assert exact_number(6e9) == '6,000,000,000' and exact_number(12.016427) == '12.016427'
    wb = _book(calculate(case(cpn=.02125, sig=.32587)))
    found = {r[0]: r[1:3] for ws in wb for r in ws.iter_rows(values_only=True) if r and r[0] in
             ('표면이자·우선배당률(연, %)', '주가 변동성(연, %)')}
    assert found['표면이자·우선배당률(연, %)'] == ('2.125%', '2.125%')
    assert found['주가 변동성(연, %)'] == ('32.587%', '32.587%')


def test_memos_stay_current_when_only_market_data_change():
    c = case()
    for t in ('refixing', 'call_method', 'split_put', 'priority'):
        c.memos[t] = {'decision': '앱 판정에 동의', 'reason': '근거'}; c.memo_context[t] = memo_key(c, t)
    c.market['S0'] = 21000.; c.method['d_base'] = '2026-03-31'; c.market['sig'] = .4
    assert all(memo_status(c, t) == '현재 조건의 기록' for t in c.memos)
    c.contract['rfx_mode'] = 1                                   # 계약 조항이 바뀌면 다시 확인
    assert memo_status(c, 'refixing').startswith('이전 조건')


def test_memo_saved_by_previous_version():
    # 이전 판이 모든 입력으로 만든 식별값(검토 주제)은 지금 보는 항목을 다 덮으므로 그대로 인정한다.
    # 이전 판이 몇 항목만 본 주제(콜 방법·우선순위 등)는 바뀌었는지 알 수 없어 다시 확인하게 한다.
    c = case()
    import hashlib, json
    from dataclasses import asdict
    from valuation.explain import _V21_GROUPS
    v = {**asdict(L.Terms()), **c.effective()}
    c.memos['refixing'] = c.memos['call_method'] = {'decision': '앱 판정에 동의', 'reason': '근거'}
    c.memo_context['refixing'] = hashlib.sha256(json.dumps({k: v[k] for k in sorted(v)}, sort_keys=True,
                                                           ensure_ascii=False).encode()).hexdigest()
    c.memo_context['call_method'] = hashlib.sha256(json.dumps({k: v[k] for k in _V21_GROUPS['call_method']},
                                                              sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    assert memo_status(c, 'refixing') == '현재 조건의 기록'
    assert memo_status(c, 'call_method').startswith('이전 조건')


def test_memo_signatures_cover_the_inputs_each_judgment_uses():
    c = case(issuer_call=2, k_w=.3, k_s=3., k_e=9., k_prem=.02, k_cmp=1)
    for t in ('call_method', 'bdt', 'split_call'):
        c.memos[t] = {'decision': '앱 판정에 동의', 'reason': '근거'}; c.memo_context[t] = memo_key(c, t)
    c2 = Case.from_dict(c.to_dict()); c2.method['pc_order'] = 1
    assert memo_status(c2, 'call_method').startswith('이전 조건')          # 동시 행사 우선순위도 콜 방법 판단에 쓰인다
    c3 = Case.from_dict(c.to_dict()); c3.market['S0'] = 25000.
    assert memo_status(c3, 'bdt').startswith('이전 조건')                  # BDT 검토는 평가 결과를 보고 판단한다
    c5 = Case.from_dict(c.to_dict()); c5.contract['mat_amt'] = 1.1
    assert memo_status(c5, 'bdt').startswith('이전 조건')
    assert memo_status(c3, 'call_method') == '현재 조건의 기록'
    assert L.split_call_separate(L.Terms(**{k: v for k, v in c.effective().items() if k in L.Terms.__dataclass_fields__}))
    assert memo_status(c3, 'split_call').startswith('이전 조건')           # 자동 출발 금액 = 100 + 별개 콜 가치
    c4 = Case.from_dict(c.to_dict()); c4.method['split_base_in'] = 101.; c4.method['split_base_why'] = '가상 배분액'
    c4.memo_context['split_call'] = memo_key(c4, 'split_call'); c4.market['S0'] = 25000.
    assert memo_status(c4, 'split_call') == '현재 조건의 기록'             # 출발 금액을 직접 넣으면 시장자료와 무관


def test_missing_year_blocker_says_where_to_fix_and_assumption_lifts_it():
    c = case(dp_rows=[{'fy': 2027, 'amt': 0.}], conv_class='equity')
    run = calculate(c)
    msg = ' '.join(export_blockers(run))
    assert '배당가능이익에 따른 상환 제약' in msg and '넣지 않은 발생연도의 재원 가정·근거' in msg
    with pytest.raises(ValueError):
        export_bundle(run)
    c2 = Case.from_dict(c.to_dict()); c2.sources['dp_missing_assumption'] = '가상 시험 근거'
    assert not export_blockers(calculate(c2))


def test_missing_year_note_sits_under_profit_table():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=120)
    app.session_state.case = case(dp_rows=[{'fy': 2027, 'amt': 1e9}])
    app.run()
    assert not app.exception
    box = next(w for w in app.text_area if w.label == '넣지 않은 발생연도의 재원 가정·근거')
    box.set_value('가상 시험: 사업계획 뒤 이익 충분').run()
    next(b for b in app.button if b.label == '재원 가정 저장').click().run()
    assert app.session_state.case.sources['dp_missing_assumption'] == '가상 시험: 사업계획 뒤 이익 충분'


def test_formula_workbook_opens_verified_inputs_only():
    wb = _book(calculate(case()), formula=True)
    ws = wb['가정']
    cell = {ws.cell(r, 2).value: ws.cell(r, 3) for r in range(1, ws.max_row + 1)}
    for lab in ('평가기준일 주가', '변동성 σ', '우선배당률 (계약)', '상환청구 보장수익률'):
        assert not cell[lab].protection.locked, lab
        assert cell[lab].font.color.rgb[-6:] == '0000FF'
    for lab in ('발행일', '평가기준일', '노드 수 n', '전환 시작 (스텝)'):
        assert cell[lab].protection.locked, lab
    assert all(c.protection.locked for c in wb['IR 입력곡선']._cells.values())   # 기준금리는 엑셀에서 고치지 않는다


def test_workbook_styles_are_shared_not_per_cell():
    # 칸마다 글꼴을 새로 만들면 조서 생성이 몇 배 느려진다 (11.5초 → 2.8초로 되돌림).
    wb = _book(calculate(case()), formula=True)
    assert len(wb._fonts) < 400


@pytest.mark.parametrize('model', ['GS'])
def test_grid_shows_engine_decision(model):
    from valuation.calculation_view import node_values
    r = calculate(case(model=model, issuer_call=1, k_w=1., k_s=12., k_e=48., k_f=12., k_prem=.05, k_cmp=1))
    nodes = node_values(r)['nodes']
    memo = r.raw['full']['memo']
    for key, n in nodes.items():
        o = memo[n['i'], n['j']]
        if 'up' in o:
            assert n['probability_kind'] == KINDS[o['gkind']] == n['kind'], key


def test_run_id_reads_as_date_time():
    rid = calculate(case()).summary['run_id']
    assert re.fullmatch(r'\d{8}-\d{6}-[0-9a-f]{4}', rid), rid


def test_input_helpers():
    from workspace_app import _decimals, won_words
    assert _decimals(2.1250000000000004) == 3 and _decimals(60000.) == 2 and _decimals(12.016427) == 6
    assert won_words(6e9) == '6,000,000,000원 · 60억 원'


def test_new_blank_case_lists_missing_inputs_in_one_line():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=60).run()
    next(w for w in app.text_input if w.label == '평가 건명').set_value('가상 신규')
    next(b for b in app.button if b.label == '빈 입력안 만들기').click().run()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    assert not app.exception
    lines = [e.value for e in app.error if e.value.startswith('필수 입력')]
    assert len(lines) == 1 and '평가기준일' in lines[0]


def test_case_file_name_has_case_and_date():
    from v2_workspace import case_filename
    c = case(); c.name = '가상/회차:1'
    assert case_filename(c) == '가상_회차_1_2026-01-01_평가입력.json'


def test_topic_and_split_memos_see_schedules_and_profit_table():
    c = case(dp_rows=[{'fy': 2027, 'amt': 1e9}])
    for t in ('redemption_constraint', 'split_put'):
        c.memos[t] = {'decision': '앱 판정에 동의', 'reason': '근거'}; c.memo_context[t] = memo_key(c, t)
    c2 = Case.from_dict(c.to_dict()); c2.market['dp_rows'] = [{'fy': 2027, 'amt': 5e8}]
    assert memo_status(c2, 'redemption_constraint').startswith('이전 조건')   # 배당가능이익 표는 상환제약 판단의 입력
    c3 = Case.from_dict(c.to_dict()); c3.contract['p_sched'] = '2027-01-01 104'
    assert memo_status(c3, 'split_put').startswith('이전 조건')               # 직접 적은 행사금액 일정표
    c4 = Case.from_dict(c.to_dict()); c4.market['S0'] = 30000.
    assert all(memo_status(c4, t) == '현재 조건의 기록' for t in c.memos)


def test_sha_call_quantity_stays_locked():
    from valuation.workpaper_v2 import EXCEL_INPUTS
    assert '콜 대상 주식수' not in EXCEL_INPUTS and '풋 가격 가산율 (연)' in EXCEL_INPUTS


def test_day1_memos_follow_any_pricing_input():
    c = case(); c.memos['day1_price'] = {'decision': '해당 없음', 'reason': '근거'}; c.memo_context['day1_price'] = memo_key(c, 'day1_price')
    c.contract['K0'] = 55000.
    assert memo_status(c, 'day1_price').startswith('이전 조건')          # 최초 인식 차이는 모형값 전체로 정해진다
