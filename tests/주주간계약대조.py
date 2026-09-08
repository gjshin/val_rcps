#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""주주간계약 조서가 엔진과 같은 답을 내는지 확인한다.

값 조서와 수식 조서를 `build_xlsx_sha` 한 함수에서 만들므로 자리와 차례는
갈라질 수 없다. 남는 것은 **수식이 엔진과 같은 계산을 하는가** 하나라서,
수식 조서를 실제로 풀어 엔진 값과 맞춰 본다. 값 조서 쪽도 같은 셀을 함께 본다.

    pip install formulas
    python3 tests/주주간계약대조.py            # 전부
    python3 tests/주주간계약대조.py 적격상장     # 이름으로 거르기

시트 이름에 한글이 있으면 formulas 가 깨져서 ASCII 로 바꾼 사본을 만들어 푼다.
"""
import sys, os, types, tempfile, warnings
warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

CASES = [
    ("기본", {}),
    ("콜 없음", dict(sha_call_s=0., sha_call_e=0.)),
    ("단리 보장", dict(sha_put_cmp=0, sha_call_cmp=0)),
    ("적격상장 36개월", dict(ipo_on=1, ipo_m=36., ipo_min=1200.)),
    ("적격상장 · 콜 존속",
     dict(ipo_on=1, ipo_m=36., ipo_min=1200., sha_qipo_kill=0)),
    ("무위험 할인", dict(sha_disc=0)),
    ("무위험 + 스프레드", dict(sha_disc=2, sha_spread=.04)),
    ("중간평가", dict(d_base="2026-03-31")),
    ("풋 주기 12개월", dict(sha_put_f=12.)),
    ("주가 60%", dict(S0=600.)),
    ("풋 의무자 = 발행회사", dict(sha_writer=1)),
    ("월 노드", dict(gap_m=1.0, sha_put_f=3., sha_call_f=3.)),
    # 상호소멸 — 한쪽이 행사하면 다른 쪽이 그 자리에서 소멸한다. 엔진에만 넣고
    # 수식 조서에 안 넣었던 적이 있어(조서를 풀면 다른 값이 나왔다) 여기 심는다.
    ("상호소멸", dict(sha_kill=1)),
    ("상호소멸 · 적격상장",
     dict(sha_kill=1, ipo_on=1, ipo_m=36., ipo_min=1200.)),
    ("상호소멸 · 콜 우선", dict(sha_kill=1, pc_order=1)),
]
BASE = dict(inst="SHA", S0=1000., K0=1000., d_issue="2025-03-31",
            d_base="2025-03-31", d_mat="2030-03-31", gap_m=6.0, sig=0.40,
            face_total=10_000_000_000.,
            sha_put_s=36., sha_put_e=60., sha_put_f=6., sha_put_yield=.08,
            sha_put_cmp=1, sha_call_s=12., sha_call_e=36., sha_call_f=6.,
            sha_call_prem=.15, sha_call_cmp=1,
            rf_curve=[(1, .0226), (3, .0240), (5, .0252)],
            cr_curve=[(1, .0500), (3, .0550), (5, .0600)])


def load_app():
    stub = types.ModuleType("streamlit"); stub.cache_data = lambda **k: (lambda f: f)
    sys.modules["streamlit"] = stub
    m = types.ModuleType("cbapp"); sys.modules["cbapp"] = m
    src = open(os.path.join(ROOT, "app.py"), encoding="utf-8").read()
    exec(compile(src.split("st.set_page_config")[0], "app.py", "exec"), m.__dict__)
    return m.__dict__


def solve(G, t, R, path):
    """수식 조서를 만들어 풀고 결과 시트의 셀을 돌려준다."""
    import openpyxl, formulas
    open(path, "wb").write(G["build_xlsx_sha"](t, R, formula=True))
    wb = openpyxl.load_workbook(path)
    mp = {nm: f"S{i:02d}" for i, nm in enumerate(wb.sheetnames)}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    f = c.value
                    for o, nn in sorted(mp.items(), key=lambda x: -len(x[0])):
                        f = f.replace(f"'{o}'!", f"{nn}!").replace(f"{o}!", f"{nn}!")
                    c.value = f
    for o, nn in mp.items(): wb[o].title = nn
    wb.save(path)
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    base = os.path.basename(path).upper()
    res = mp["결과"]
    out = {}
    for k, v in sol.items():
        ku = k.upper()
        if ku.startswith(f"'[{base}]{res}'!"):
            try: out[ku.split("!")[-1]] = float(v.value[0, 0])
            except Exception: pass
    return out


def main():
    G = load_app()
    import io, openpyxl
    bad = 0
    only = sys.argv[1] if len(sys.argv) > 1 else ""
    for lbl, over in CASES:
        if only and only not in lbl: continue
        t = G["Terms"](**{**BASE, **over}); G["derive"](t)
        R = G["sha_engine"](t)
        with tempfile.TemporaryDirectory() as d:
            got = solve(G, t, R, os.path.join(d, "wb.xlsx"))
        # 값 조서는 같은 자리에 숫자를 담는다. 함께 본다.
        vw = openpyxl.load_workbook(
            io.BytesIO(G["build_xlsx_sha"](t, R, formula=False)))["결과"]
        gp = R["gross"]
        rows = [("지분가치", "C6", 6, 100*t.S0/t.K0),
                ("풋", "C7", 7, R["put"]),
                ("콜", "C8", 8, R["call"]),
                ("지분+풋−콜", "C9", 9, 100*t.S0/t.K0 + R["put"] - R["call"]),
                ("총액 부채", "C15", 15, gp["pv"] if gp else 0.0)]
        print(f"\n[{lbl}]")
        for nm, cell, vr, want in rows:
            h = got.get(cell)
            v = vw.cell(vr, 3).value
            ok = (h is not None and abs(h - want) < 1e-4
                  and v is not None and abs(float(v) - want) < 1e-4)
            if not ok: bad += 1
            print("   %-12s 수식 %11s · 값 %11s · 엔진 %11.4f  %s"
                  % (nm, f"{h:.4f}" if h is not None else "없음",
                     f"{float(v):.4f}" if v is not None else "없음", want,
                     "" if ok else "★"))
    print("\n" + ("모든 항목 일치" if not bad else f"★ {bad}건 불일치"))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
