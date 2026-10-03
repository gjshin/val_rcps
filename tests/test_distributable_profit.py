"""배당가능이익 상환 제약 (상환전환우선주) — 가상 수치.

· 넣지 않으면 종전과 같다 (배당이 가능하다는 전제).
· 발생연도 Y 의 이익은 Y+1 년 재원 · 우선배당 먼저 · 갚지 못한 금액은 다음 해로 · 동순위는 상환금 비율.
· 손계산 대조는 앱 계산 함수를 쓰지 않고 시험 안에서 따로 센다.
· 발행자 상환권은 재원이 있어야 행사, 제3자 매도청구권은 직접 영향 없음, 다른 상품(전환사채)에는 영향 없음.
· 수식 조서를 리브레오피스로 다시 계산해 엔진과 같은지 본다.
"""
import io
import math
import shutil
import subprocess
import zipfile
from dataclasses import asdict

import pytest
from openpyxl import load_workbook

from valuation import legacy
from valuation.case import import_legacy, inspect_case
from valuation.service import calculate, export_bundle
from valuation.xlsx_validation import compare_cells

FACE = 10_000_000_000.0          # 평가대상 발행총액 100억


def rcps(**o):
    t = legacy.Terms(inst="RCPS", d_issue="2026-01-01", d_base="2026-01-01", d_mat="2031-01-01",
                     S0=20000., K0=60000., issue_px=60000., face_total=FACE, sig=.35, par=500.,
                     cpn=.02, div_basis=0, div_mode=0, ipay=12., mat_mode=0, ytm=.05, ytm_cmp=1,
                     p_mode="accrue", p_yield=.05, p_cmp=1, p_s=12., p_e=59., p_f=0.,
                     rfx_mode=0, issuer_call=0, k_w=0., cv_s=1., cv_e=60., view="issuer", gap_m=1.,
                     rf_curve=[[1, .03], [10, .03]], cr_curve=[[1, .08], [10, .08]])
    extra = {k: o.pop(k) for k in list(o) if k in ("dp_rows", "dp_others", "dp_delay", "dp_from")}
    for k, v in o.items(): setattr(t, k, v)
    c = import_legacy(asdict(t), "가상 우선주")
    c.exercise_styles = {"p_f": "any", "cv": "any"}
    if "dp_rows" in extra: c.market["dp_rows"] = extra.pop("dp_rows")
    c.contract.update(extra)
    return c


def ea(run):
    t = run.terms
    return legacy.exercise_amounts(t, t.n, t.T/t.n), t


def test_default_unchanged_without_profits():
    a = calculate(rcps()); b = calculate(rcps(dp_rows=[]))
    assert a.summary["amounts_100"] == b.summary["amounts_100"]
    E, t = ea(a)
    assert all(E["put_val"](i) == E["put"](i) for i in E["p_dates"]) and E["red_val"] == E["red"]


def test_hand_calculation_with_pari_product():
    # 2028 년 청구 → 발생연도 2027 이익 30억 · 2028 이익 60억. 평가대상 우선배당 2% (청구한 해만).
    # 동순위 상품: 발행 2026-01-01 · 50억 · 보장수익률 0 → 상환금 50 (100 기준) · 우선배당 1% (청구한 해).
    oth = [dict(name="가상 2회차", rank="pari", issue="2026-01-01", face=5e9, yld=0.0, cmp=1,
                start="2027-01-01", end="2030-12-31", div=.01)]
    run = calculate(rcps(dp_rows=[{"fy": 2027, "amt": 3e9}, {"fy": 2028, "amt": 6e9}], dp_others=oth))
    E, t = ea(run)
    dt_ = t.T/t.n
    i = next(i for i in sorted(E["p_dates"]) if legacy.dp_step_dt(t, dt_, i).year == 2028)
    A = E["put"](i)
    RF, CR = legacy.curves(t); f = legacy.forward_rate(CR, 0.0, dt_)
    # 손계산 — 1년차: 재원 30 − 우선배당(평가대상 2 + 동순위 0.5) = 27.5 를 A : 50 으로 나눈다
    cap1 = 30.0 - 2.0 - 0.5
    p1 = min(A, cap1*A/(A + 50.0)); q1 = min(50.0, cap1*50.0/(A + 50.0))
    # 2년차: 재원 60 (우선배당 없음 — 둘 다 청구함), 남은 금액 비율
    be, bo = A - p1, 50.0 - q1
    cap2 = 60.0
    p2 = min(be, cap2*be/(be + bo))
    # 3년차: 발생연도 2029 는 넣지 않았으므로 제한 없음 — 남은 것 전부
    p3 = be - p2
    m1 = int(math.floor(1/dt_ + .5)); m2 = int(math.floor(2/dt_ + .5))
    pv = p1 + p2*math.exp(-f*m1*dt_) + p3*math.exp(-f*m2*dt_)
    assert E["put_val"](i) == pytest.approx(pv, rel=1e-12)
    assert E["put_val"](i) < A


def test_dividends_first_and_zero_capacity_carries_forward():
    # 2028 년 재원(발생연도 2027) 이 우선배당(2)보다 작으면 그 해 상환은 0 이고 전액이 다음 해로 넘어간다.
    run = calculate(rcps(dp_rows=[{"fy": 2027, "amt": 1e8}], dp_delay=.04))
    E, t = ea(run); dt_ = t.T/t.n
    i = next(i for i in sorted(E["p_dates"]) if legacy.dp_step_dt(t, dt_, i).year == 2028)
    sc = E["dp"].schedule(i, E["put"](i))
    assert sc["rows"][0][3] == 0.0 and sc["rows"][0][6] == 0.0          # 재원 0 · 지급 0
    assert sc["rows"][1][6] == pytest.approx(E["put"](i)*1.04, rel=1e-12)  # 다음 해 가산율 붙여 전액
    assert any("우선배당보다 작습니다" in m.message for m in run.issues if m.code == "input_check")


def test_fiscal_year_mapping_and_unlisted_years_unlimited():
    # 발생연도 2030 만 넣으면 2031 년 청구분만 제한된다 — 이 평가의 청구는 2027~2030 년이라 모두 제한 없음.
    run = calculate(rcps(dp_rows=[{"fy": 2030, "amt": 0.0}]))
    E, t = ea(run)
    assert all(E["put_val"](i) == pytest.approx(E["put"](i), rel=1e-15) for i in E["p_dates"])
    assert any("넣지 않아" in m.message for m in run.issues if m.code == "input_check")


def test_fund_start_date_moves_year_boundary():
    # 재원 사용 시작일 4월 1일 — 2028년 2월 청구는 아직 2027년 재원(발생연도 2026 이익)을 쓴다.
    # 발생연도 2026 이익 0 → 그 해 지급 0. 1월 1일(기본)이면 2028년 재원(발생연도 2027, 넣지 않음 → 제한 없음).
    rows = [{"fy": 2026, "amt": 0.0}]
    a = calculate(rcps(dp_rows=rows, dp_from="04-01")); b = calculate(rcps(dp_rows=rows))
    Ea, t = ea(a); Eb, _ = ea(b); dt_ = t.T/t.n
    feb = [i for i in sorted(Ea["p_dates"]) if (d := legacy.dp_step_dt(t, dt_, i)).year == 2028 and d.month < 4]
    may = [i for i in sorted(Ea["p_dates"]) if (d := legacy.dp_step_dt(t, dt_, i)).year == 2028 and d.month >= 4]
    assert feb and may
    for i in feb:
        assert legacy.dp_fund_year(t, legacy.dp_step_dt(t, dt_, i)) == 2027
        assert Ea["dp"].schedule(i, Ea["put"](i))["rows"][0][6] == 0.0
        assert Ea["put_val"](i) < Ea["put"](i) and Eb["put_val"](i) == Eb["put"](i)
    for i in may:
        assert Ea["put_val"](i) == pytest.approx(Ea["put"](i), rel=1e-15)
    bad = [x.message for x in inspect_case(rcps(dp_rows=rows, dp_from="13-40")) if x.severity == "error"]
    assert bad


def test_issuer_call_needs_capacity_third_party_does_not():
    call = dict(k_s=12., k_e=59., k_f=12., k_prem=.06, k_cmp=1, k_w=1.)
    prof = [{"fy": y, "amt": 1e9} for y in range(2026, 2031)]
    a = calculate(rcps(issuer_call=1, dp_rows=prof, **call))
    E, _ = ea(a)
    assert E["k_dates"] and not any(E["k_on"](i) for i in E["k_dates"])       # 재원 10 < 상환금 100 이상
    assert any("발행자 상환권은 재원이 부족한" in m.message for m in a.issues if m.code == "input_check")
    b = calculate(rcps(issuer_call=2, dp_rows=prof, **call))
    E2, _ = ea(b)
    assert all(E2["k_on"](i) for i in E2["k_dates"])                          # 제3자는 직접 제한 없음
    assert any("제3자 지정 매도청구권은" in m.message for m in b.issues if m.code == "input_check")


def test_other_products_unaffected():
    t = legacy.Terms(inst="CB"); t.dp_rows = [{"fy": 2026, "amt": 0.0}]
    assert not legacy.dp_active(t) and legacy.dp_issues(t) == []


def test_input_errors_block():
    bad = rcps(dp_rows=[{"fy": 2027, "amt": 1e9}, {"fy": 2027, "amt": 2e9}, {"fy": 2028, "amt": -1.0}],
               dp_others=[dict(name="날짜 없음", rank="pari", face=1e9)])
    msgs = [i.message for i in inspect_case(bad) if i.severity == "error"]
    assert any("두 번" in m for m in msgs) and any("0 이상" in m for m in msgs) and any("YYYY-MM-DD" in m for m in msgs)


def test_value_workbooks_carry_schedule_sheet():
    run = calculate(rcps(dp_rows=[{"fy": 2027, "amt": 3e9}]))
    for detail in (False, True):
        z = zipfile.ZipFile(io.BytesIO(export_bundle(run, detail=detail, accounting=True)))
        assert "00 배당가능이익 상환" in load_workbook(io.BytesIO(z.read("value_review.xlsx"))).sheetnames


@pytest.mark.skipif(not (shutil.which("libreoffice") or shutil.which("soffice")), reason="LibreOffice 없음")
@pytest.mark.parametrize("call,frm", [(0, "01-01"), (1, "01-01"), (0, "04-01")])
def test_formula_workbook_matches_engine(tmp_path, call, frm):
    kw = dict(issuer_call=1, k_s=12., k_e=59., k_f=12., k_prem=.06, k_cmp=1, k_w=1.) if call else {}
    oth = [dict(name="가상 2회차", rank="pari", issue="2026-06-01", face=5e9, yld=.04, cmp=0,
                start="2028-01-01", end="2030-06-30", div=.01)]
    run = calculate(rcps(dp_rows=[{"fy": 2027, "amt": 3e9}, {"fy": 2028, "amt": 2.5e10}, {"fy": 2029, "amt": 4e9}],
                         dp_others=oth, dp_delay=.03, dp_from=frm, **kw))
    t, R = run.terms, run.raw
    data = zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=True, detail=True, accounting=True))).read("formula_review.xlsx")
    (tmp_path/"s.xlsx").write_bytes(data); (tmp_path/"o").mkdir()
    subprocess.run([shutil.which("libreoffice") or shutil.which("soffice"),
                    "-env:UserInstallation=" + (tmp_path/"p").as_uri(), "--headless", "--convert-to", "xlsx",
                    "--outdir", str(tmp_path/"o"), str(tmp_path/"s.xlsx")], capture_output=True, timeout=1800)
    wb = load_workbook(tmp_path/"o"/"s.xlsx", data_only=True)
    exp = legacy.formula_key_cells(t, R, legacy.eir_or_none(t, R["full"], R["b0"], R["b1"], R["b2"], R["ca"]))
    rows = compare_cells(wb, exp)
    assert all(r[4] == "일치" for r in rows), rows
    E, _ = ea(run)
    ws = wb["00 격자 공통"]
    for i in range(t.n + 1):
        if E["p_on"](i):
            assert ws.cell(7, 3+i).value == pytest.approx(E["put_val"](i), rel=1e-9, abs=1e-9), i
        if call:
            assert int(ws.cell(5, 3+i).value or 0) == (1 if E["k_on"](i) else 0), i
    assert ws.cell(10, 3+t.n).value == pytest.approx(E["red_val"], rel=1e-9)


def test_screen_keeps_inputs_and_matches_engine():
    # 입력 화면을 거쳐도 연도별 이익·다른 상품 표가 그대로 남고, 화면 평가 금액이 엔진과 같다.
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    oth = [dict(name="가상 2회차", rank="pari", issue="2026-01-01", face=5e9, yld=0.0, cmp=1,
                start="2027-01-01", end="2030-12-31", div=.01)]
    c = rcps(dp_rows=[{"fy": 2027, "amt": 3e9}, {"fy": 2028, "amt": 6e9}], dp_others=oth)
    app = AppTest.from_file(str(Path(__file__).parent.parent/"app.py"), default_timeout=300)
    app.session_state.case = c; app.run()
    assert not app.exception
    app.radio(key="_workflow_stage").set_value("평가·분석").run()
    next(b for b in app.button if b.label == "현재 입력으로 평가").click().run()
    assert not app.exception
    case = app.session_state.case
    assert case.market["dp_rows"] == [{"fy": 2027, "amt": 3e9}, {"fy": 2028, "amt": 6e9}]
    assert case.contract["dp_others"][0]["rank"] == "pari"
    assert app.session_state.run.summary["amounts_total"] == calculate(case).summary["amounts_total"]
    assert any("배당가능이익 반영" in x.label for x in app.expander)
