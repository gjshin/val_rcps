#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""수식 조서의 «살아 있는 검산» 이 값 조서·엔진과 같은가.

수식 조서에는 파이썬이 센 값을 옮겨 적은 문자열(값 조서와 같은 것)과, 같은 것을 다른 시트의
셀로 다시 계산하는 **수식**이 나란히 있다. 둘이 같아야 한다 — 다르면 수식 배선이 틀린 것이다.
`formulas` 라이브러리로 수식 조서를 실제로 풀어 견준다 (조서대조와 같은 방법).

    [1] BDT 시트 — 시장 할인계수(선도이자율 합의 수식) = 엔진 bdt_parts 의 mkt, Σ Q − mkt = 0
    [2] 분리 판단 시트 — 10% 검토 수식(행사금액 · 상각후원가 · 차이 · 판정) = split_test
    [3] 결과 시트 — 우선순위별 매도청구권 두 줄 = pc_compare, 마팅게일 검산(δ>0) = 1
    [4] 검산요약 시트 — 값·판정 수식 = 값 조서의 문자열 (수식화된 줄만)

    python3 tests/검산수식대조.py            # 전체
    python3 tests/검산수식대조.py --quick    # 앞 두 케이스
"""
import sys, os, io, re, warnings
warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tests"))
from 손계산대조 import G, Terms, derive, chk, chk_bool, FAIL      # noqa: E402
from 조서대조 import _combin                                       # noqa: E402

RF = [(1, .0226), (3, .0240), (5, .0252)]
CR = [(1, .1409), (3, .1740), (5, .1905)]
CASES = [
    ("기본 CB", dict()),
    ("BDT σ20% · 배당 2%", dict(put_bdt=1, bdt_sig=.20, div_y=.02)),
    ("부채 분류 · 상실이자 아님", dict(conv_class="liability")),
    ("RCPS 발행자콜", dict(inst="RCPS", issuer_call=1, k_w=1.0, k_s=12., k_e=48., p_s=24., p_e=57.)),
    # 겹치는 노드에서 조기상환금액 > 매도청구금액 — 우선순위가 값을 가른다
    ("우선순위 갈림", dict(p_mode="accrue", p_yield=.05, p_cmp=1, k_prem=.01, k_cmp=1,
                       k_s=12., k_e=36., p_s=12., p_e=36., p_f=6., k_f=6., k_w=.3)),
]


def make(over):
    t = Terms(); t.rf_curve = RF; t.cr_curve = CR; t.carry = 1; t.gap_m = 6.0
    for k, v in over.items(): setattr(t, k, v)
    derive(t)
    full, b0, b1, b2, ca, conv = G["decompose"](t)
    eir = G["eir_or_none"](t, full, b0, b1, b2, ca)
    return t, full, b0, b1, b2, ca, conv, eir


def solve_all(xbytes, path):
    """수식 조서를 ASCII 시트명으로 바꿔 formulas 로 푼다. {시트명: {셀: 값}} 과 시트명→SNN 사전."""
    import openpyxl, formulas
    wb = openpyxl.load_workbook(io.BytesIO(xbytes))
    mp = {nm: f"S{i:02d}" for i, nm in enumerate(wb.sheetnames)}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    f = c.value
                    for o, nn in sorted(mp.items(), key=lambda x: -len(x[0])):
                        f = f.replace(f"'{o}'!", f"{nn}!").replace(f"{o}!", f"{nn}!")
                    c.value = _combin(f)
    for o, nn in mp.items(): wb[o].title = nn
    wb.save(path)
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    base = os.path.basename(path).upper()
    out = {}
    for k, v in sol.items():
        ku = k.upper()
        m = re.match(r"'\[" + re.escape(base) + r"\](S\d\d)'!([A-Z]+\d+)$", ku)
        if not m: continue
        try:
            val = v.value[0, 0]
        except Exception:
            continue
        out.setdefault(m.group(1), {})[m.group(2)] = val
    return out, mp


def cell(sol, mp, sheet, addr):
    return sol.get(mp[sheet], {}).get(addr.upper())


def find_row(ws, col, text, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is not None and str(v).startswith(text): return r
    return None


def main():
    quick = "--quick" in sys.argv
    import openpyxl
    from openpyxl.utils import get_column_letter as gl
    print("검산수식대조 — 수식 조서의 살아 있는 검산 대 값 조서·엔진")
    for ci, (lbl, over) in enumerate(CASES[:2] if quick else CASES):
        t, full, b0, b1, b2, ca, conv, eir = make(over)
        xb = G["build_xlsx_formula"](t, full, b0, b1, b2, ca, conv, eir)
        vb = G["build_xlsx"](t, full, b0, b1, b2, ca, conv, eir)
        wbx = openpyxl.load_workbook(io.BytesIO(xb))
        wbv = openpyxl.load_workbook(io.BytesIO(vb))
        # formulas 는 파일명에 % · 한글이 있으면 깨진다 — ASCII 이름
        path = os.path.join(ROOT, "tests", "output", f"chk_formula_{ci}.xlsx")
        os.makedirs(os.path.dirname(path), exist_ok=True)
        sol, mp = solve_all(xb, path)
        print(f"\n[{lbl}]")
        n = int(t.n)
        # ── [1] BDT 시장 할인계수 ──
        if G["put_bdt_on"](t):
            bp = G["bdt_parts"](t)
            W = wbx["BDT 단기이자율"]
            rq = find_row(W, 2, "시장 할인계수"); rd = find_row(W, 2, "차이")
            worst = max(abs(cell(sol, mp, "BDT 단기이자율", f"{gl(3+i)}{rq}") - bp["mkt"][i]) for i in range(n+1))
            chk("BDT · 시장 할인계수 수식 − 엔진 mkt (최대)", worst, 0.0, 1e-9)
            dmax = max(abs(cell(sol, mp, "BDT 단기이자율", f"{gl(3+i)}{rd}")) for i in range(n+1))
            chk("BDT · |Σ Q − 시장 할인계수| 최대 (수식)", dmax, 0.0, 1e-6)
            chk_bool("BDT · 시장 할인계수 칸이 수식이다", str(W.cell(rq, 4).value).startswith("=EXP("))
        # ── [2] 분리 판단 10% 검토 ──
        SP = G["split_test"](t, full, b0, b1, b2, ca, [] if eir is None else eir[1])
        J = wbx["분리 판단"]
        r0 = find_row(J, 2, "B4.3.5(5)(가) 10% 검토 — 수식")
        if r0 and SP["put"]["지표"]:
            pv = cell(sol, mp, "분리 판단", f"C{r0+1}"); bv = cell(sol, mp, "분리 판단", f"C{r0+2}")
            gap = cell(sol, mp, "분리 판단", f"C{r0+4}"); vd = cell(sol, mp, "분리 판단", f"C{r0+5}")
            chk("분리 판단 · 행사금액 수식 = split_test", pv, SP["put"]["지표"]["첫 조기상환일 행사금액"], 1e-6)
            chk("분리 판단 · 상각후원가 수식 = split_test", bv, SP["put"]["지표"]["같은 시점 상각후원가"], 1e-6)
            chk("분리 판단 · 차이 수식 = split_test", gap, SP["put"]["지표"]["차이"], 1e-9)
            chk_bool(f"분리 판단 · 판정 수식 «{vd}» = «{SP['put']['결론']}»", vd == SP["put"]["결론"])
        else:
            chk_bool("분리 판단 · 조기상환권 없음 → 수식 블록 없음", r0 is None and not SP["put"]["지표"])
        # ── [3] 결과 시트 — 우선순위 두 줄 · 마팅게일 ──
        pcc = G["pc_compare"](t)
        for nm, wb_ in (("수식", wbx), ("값", wbv)):
            R = wb_["결과"]
            rp = find_row(R, 2, "4. 풋·콜 우선순위별")      # RCPS 는 relabel 로 «발행자 상환권» 이 된다
            chk_bool(f"결과({nm}) · 우선순위 블록이 있다", rp is not None)
            if rp and pcc:
                got = [(R.cell(rp+2+i, 2).value, R.cell(rp+2+i, 3).value, R.cell(rp+2+i, 4).value) for i in range(2)]
                for (lb, ca2, cv2, _), (glb, gca, gcv) in zip(pcc, got):
                    chk(f"결과({nm}) · {lb} 매도청구권", gca, ca2, 1e-9)
                    chk(f"결과({nm}) · {lb} 전환권대가", gcv, cv2, 1e-9)
            elif rp:
                chk_bool(f"결과({nm}) · 겹치지 않음 문구", "겹치는 노드" in str(R.cell(rp+1, 2).value))
        mg = cell(sol, mp, "결과", "C29")
        chk("결과 · 주가 마팅게일 (δ 반영) = 1", mg, 1.0, 1e-6)
        # ── [4] 검산요약 — 수식화된 줄 = 값 조서 문자열 ──
        Cx, Cv = wbx["검산요약"], wbv["검산요약"]
        nrow = 0
        for r in range(6, 40):
            nm = Cv.cell(r, 2).value
            if not nm or nm == "모든 항목 적합" or str(nm).startswith("확인 필요 "):
                if not nm: break
                continue
            fx = Cx.cell(r, 3).value
            if isinstance(fx, str) and fx.startswith("="):
                nrow += 1
                for col, what in ((3, "값"), (4, "판정"), (5, "설명")):
                    fv = Cx.cell(r, col).value
                    if isinstance(fv, str) and fv.startswith("="):
                        got = cell(sol, mp, "검산요약", f"{gl(col)}{r}")
                        want = Cv.cell(r, col).value
                        chk_bool(f"검산요약 · {nm} · {what} «{str(got)[:28]}» = «{str(want)[:28]}»", str(got) == str(want))
        print(f"  (검산요약 수식 줄 {nrow})")
    print()
    if FAIL:
        print(f"★ 어긋남 {len(FAIL)}건")
        for f in FAIL: print("   -", f)
        return 1
    print("모두 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
