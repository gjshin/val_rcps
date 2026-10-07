"""PR49 review regressions through the same export and input paths used by the app."""
import io
import os
import shutil
import zipfile
from dataclasses import asdict
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from test_commercial_v2 import case
from valuation import legacy as L
from valuation.case import import_legacy
from valuation.service import calculate, export_bundle
from valuation.xlsx_validation import recalculate_and_compare

SOFFICE = os.environ.get('VALUATION_SOFFICE') or shutil.which('libreoffice') or shutil.which('soffice')


def workbook(run, formula=True):
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=formula, detail=True))) as z:
        return load_workbook(io.BytesIO(z.read('formula_review.xlsx' if formula else 'value_review.xlsx')))


def coupon_edit(inst, model, old, new):
    common = dict(inst=inst, model=model, gap_m=6., mat_mode=1, view='holder', cv_s=99., cv_e=0.,
                  p_s=99., p_e=0., ytm=.05)
    before = calculate(case(cpn=old, **common))
    after = calculate(case(cpn=new, **common))
    wb = workbook(before)
    ws = wb['가정']
    label = '표면이자율' if inst == 'CB' else '우선배당률 (계약)'
    cell = next(ws.cell(i, 3) for i in range(1, ws.max_row+1) if ws.cell(i, 2).value == label)
    assert not cell.protection.locked
    cell.value = new
    return wb, L.formula_key_cells(after.terms, after.raw), before


@pytest.mark.parametrize('inst', ['CB', 'RCPS'])
def test_zero_coupon_export_retains_dates_without_changing_app_payments(inst):
    wb, _, before = coupon_edit(inst, 'TF', 0., .03)
    t = before.terms
    assert L.pay_steps(t, t.n, t.T/t.n) == {}
    contract_steps = L.pay_steps(t, t.n, t.T/t.n, include_zero=True)
    assert sum(contract_steps.values()) == 5
    assert any('지급일' in str(c.value) for c in wb['00 계약일 목록']['B'])
    for step in contract_steps:
        assert 'COUNTIF' in wb['00 격자 공통'].cell(9, 3+step).value


@pytest.mark.skipif(not SOFFICE, reason='LibreOffice 없음')
@pytest.mark.parametrize('inst', ['CB', 'RCPS'])
@pytest.mark.parametrize('model', ['TF', 'GS'])
@pytest.mark.parametrize('old,new', [(0., .03), (.03, 0.)])
def test_coupon_edit_crossing_zero_recalculates_to_app(inst, model, old, new):
    wb, expected, _ = coupon_edit(inst, model, old, new)
    data = io.BytesIO(); wb.save(data)
    assert recalculate_and_compare(data.getvalue(), expected=expected)['same_grid']


def sha_case(count=2, **changes):
    rows = [dict(name='1차', start='2026-01-01', end='2026-12-31', style='any', price=900., rate=0., put_q=30000., call_q=0.),
            dict(name='2차', start='2027-01-01', end='2027-12-31', style='any', price=1600., rate=.03, put_q=30000., call_q=0.)]
    t = L.Terms(inst='SHA', S0=1000., K0=1000., d_issue='2025-03-31', d_base='2025-03-31',
                d_mat='2027-12-31', gap_m=6., sig=.4, face_total=75e6, sha_rows=rows[:count],
                sha_put_s=12., sha_put_e=24., rf_curve=[(1,.025),(5,.03)], cr_curve=[(1,.06),(5,.07)])
    for k, v in changes.items():
        setattr(t, k, v)
    return import_legacy(asdict(t), 'SHA 합성 회차 검증')


@pytest.mark.parametrize('count', [0, 1, 2])
def test_sha_all_assumption_sheets_open_only_supported_inputs(count):
    wb = workbook(calculate(sha_case(count)))
    sheets = [w for w in wb if w.title.endswith('가정')]
    assert len(sheets) == max(1, count)
    for ws in sheets:
        inputs = {ws.cell(r,2).value: ws.cell(r,3) for r in range(1,ws.max_row+1)}
        for label in ('평가기준일 주가 (원)', '변동성 σ (연)', '풋 가격 가산율 (연)'):
            assert not inputs[label].protection.locked, (ws.title,label)
        assert inputs['콜 대상 주식수'].protection.locked
        assert ws.protection.sheet
        assert all(c.protection.locked for c in ws._cells.values() if c.data_type == 'f')


def sha_edit():
    wb = workbook(calculate(sha_case()))
    for ws in wb:
        if ws.title.endswith('가정'):
            for row in range(1, ws.max_row+1):
                value = {'평가기준일 주가 (원)':900., '변동성 σ (연)':.5}.get(ws.cell(row,2).value)
                if value is not None:
                    assert not ws.cell(row,3).protection.locked
                    ws.cell(row,3).value = value
    control = workbook(calculate(sha_case(S0=900., sig=.5)), formula=False)
    expected = [(f'{ws.title} {c.coordinate}', ws.title, c.coordinate, c.value)
                for ws in control if ws.title.endswith('결과') or ws.title == '회차 합계'
                for row in ws for c in row if c.column >= 3 and type(c.value) in (int,float)]
    assert expected
    return wb, expected


@pytest.mark.skipif(not SOFFICE, reason='LibreOffice 없음')
def test_sha_multiple_tranche_edits_recalculate_to_app():
    wb, expected = sha_edit()
    data=io.BytesIO(); wb.save(data)
    recalculate_and_compare(data.getvalue(), expected=expected)


def test_precision_controls_preserve_original_input_and_run():
    run=calculate(case(S0=20000.123456789, cpn=.02125))
    app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=60)
    app.session_state.case=run.case
    app.session_state.run=run
    app.run()
    for value in (True,False):
        app.toggle(key='_input_precision').set_value(value).run()
        assert not app.exception
        assert app.session_state.case.market['S0'] == 20000.123456789
        assert app.session_state.case.contract['cpn'] == .02125
        assert app.session_state.run.summary['run_id'] == run.summary['run_id']
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    for unit in ('총액 · 원','주당 · 원','원금 100'):
        app.radio(key='result_unit').set_value(unit).run()
        assert not app.exception
        for metric in app.metric[:2]:
            assert len(metric.value.split('.')[-1]) == 2
    assert app.session_state.run.summary['run_id'] == run.summary['run_id']
