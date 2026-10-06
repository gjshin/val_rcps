import io
import math
import zipfile
from dataclasses import asdict
import pytest
from openpyxl import load_workbook
from valuation import legacy as L
from valuation.case import Case, import_legacy, inspect_case
from valuation.service import calculate, export_bundle, refresh_run
from valuation.report import judgment_rows
from valuation.explain import memo_key, memo_status, put_diagnostic, node_trace, export_blockers


def case(**changes):
    t=L.Terms(inst='RCPS',d_issue='2026-01-01',d_base='2026-01-01',d_mat='2031-01-01',
        S0=20000.,K0=60000.,issue_px=60000.,face_total=1e10,sig=.35,par=500.,
        cpn=.02,div_basis=0,div_mode=0,ipay=12.,mat_mode=0,ytm=.05,ytm_cmp=1,
        p_mode='accrue',p_yield=.05,p_cmp=1,p_s=12.,p_e=59.,p_f=12.,rfx_mode=0,
        issuer_call=0,k_w=0.,cv_s=1.,cv_e=60.,view='issuer',gap_m=1.,
        rf_curve=[[1,.03],[10,.03]],cr_curve=[[1,.08],[10,.08]])
    for k,v in changes.items():setattr(t,k,v)
    return import_legacy(asdict(t),'V2 합성 검증')


def split(run):
    return L.split_test(run.terms,*[run.raw[k] for k in ('full','b0','b1','b2','ca')],[])


def test_exact_exercise_amortization_and_payment_boundary():
    r=calculate(case(cpn=0.,ytm=.12,p_s=6.,p_e=6.,p_mode='fixed',p_rate=100.,conv_class='equity',emb_approach=2))
    sp=split(r)['put'];eir=L.eir_table(r.terms,100)[0]
    exact=100*math.sqrt(1+eir)
    assert sp['지표']['같은 시점 상각후원가']==pytest.approx(exact,abs=1e-9)
    assert sp['지표']['차이']==pytest.approx(abs(100-exact)/exact)
    assert sp['결론']=='분리하지 않을 여지'
    rows=[(1,1.,100.,10.,5.,105.),(2,2.,105.,10.5,5.,110.5)]
    assert L.amortized_at(100.,.1,rows,.5)==pytest.approx(100*math.sqrt(1.1))
    assert L.amortized_at(100.,.1,rows,1.)==105.
    assert L.amortized_at(100.,.1,rows,1.5)==pytest.approx(105*math.sqrt(1.1))


def test_every_call_exercise_and_formulas():
    r=calculate(case(cpn=0.,ytm=0.,p_s=999.,p_e=0.,conv_class='equity',emb_approach=2,
        issuer_call=1,k_third=0,k_transfer=0,k_sep=0,k_w=1.,k_s=12.,k_e=48.,k_f=12.,k_prem=.08,k_cmp=1))
    sp=split(r)['call']
    assert len(sp['회차'])==4
    assert [row[1] for row in sp['회차']]==pytest.approx([100*1.08**i for i in range(1,5)])
    assert sp['지표']['가장 큰 차이']==pytest.approx(.36048896)
    assert sp['결론']=='분리'
    with zipfile.ZipFile(io.BytesIO(export_bundle(r,formula=True,judgment=False))) as z:
        w=load_workbook(io.BytesIO(z.read('formula_review.xlsx')))
    assert '분리 판단' in w
    formulas=[c.value for c in w['분리 판단']._cells.values() if c.data_type=='f']
    assert any('33' in f and 'MATCH' in f for f in formulas)
    assert any('COUNTIF' in f and '^MAX(0,' in f for f in formulas)
    assert '판단·근거' not in w
    assumption = w['가정']
    assert assumption.protection.sheet
    inputs = {assumption.cell(row,2).value: assumption.cell(row,3) for row in range(1,assumption.max_row+1)}
    assert not inputs['평가기준일 주가'].protection.locked
    assert not inputs['변동성 σ'].protection.locked
    assert inputs['발행일'].protection.locked
    assert inputs['노드 수 n'].protection.locked
    assert inputs['평가기준일 주가'].font.color.rgb[-6:] == '0000FF'


def test_put_diagnostic_uses_delayed_value():
    r=calculate(case(dp_rows=[{'fy':y,'amt':0.} for y in range(2026,2035)],conv_class='equity'))
    rows=put_diagnostic(r.terms); first=rows[0]
    ea=L.exercise_amounts(r.terms,r.terms.n,r.terms.T/r.terms.n)
    assert first['지급 제약·이자 반영 청구가치']==pytest.approx(ea['put_val'](first['스텝']))
    assert first['비율'] < 1 < first['계약 청구액']/first['계속보유가치']
    assert export_blockers(r)
    c=Case.from_dict(r.case.to_dict());c.sources['dp_missing_assumption']='전망 종료 후 전액 상환재원 확보를 가정'
    assert not export_blockers(refresh_run(r,c))


def test_memos_export_with_conditions_and_read_old_files():
    c=case();c.memos={'rcps_equity':{'decision':'해당 없음','reason':'회계기준 검토'},
        'conv_resp':{'decision':'앱 판정에 동의','reason':'전환 대응 조건'},
        'call_method':{'decision':'앱 판정에 동의','reason':'과거 방법'}}
    assert memo_status(c,'call_method')=='작성 당시 조건 미확인'
    c.memo_context['call_method']=memo_key(c,'call_method')
    assert memo_status(c,'call_method')=='현재 조건의 기록'
    c.method['k_method']=1
    assert memo_status(c,'call_method').startswith('이전 조건')
    r=calculate(c);text=str(judgment_rows(r))
    assert all(x in text for x in ('회계기준 검토','전환 대응 조건','과거 방법','이전 조건'))
    old=c.to_dict();old.pop('memo_context')
    assert Case.from_dict(old).memos==c.memos


def test_invalid_bootstrap_detected_before_pricing_and_negative_rates_allowed():
    errors=inspect_case(case(rf_curve=[[1,.01],[2,.01],[3,.9],[10,.9]]))
    assert any(x.code=='curve_discount' and '3년' in x.message for x in errors)
    assert not any(x.severity=='error' for x in inspect_case(case(rf_curve=[[1,-.005],[10,-.005]])))


@pytest.mark.parametrize('model',['TF','GS'])
def test_actual_node_trace_no_repricing(model,monkeypatch):
    r=calculate(case(model=model))
    monkeypatch.setattr(L,'engine',lambda *a,**k: (_ for _ in ()).throw(AssertionError('repriced')))
    for i,j in [(0,0),(3,1),(15,7)]:
        tr=node_trace(r,i,j)
        assert abs(tr['reconciliation'])<1e-10
        assert tr['value']==pytest.approx(r.raw['full']['memo'][i,j]['V'] if model=='GS' else sum(r.raw['full']['memo'][i,j][k] for k in ('E','B')))


def test_node_screen_survives_shorter_grid_and_unit_switch_without_repricing(monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    r=calculate(case(d_mat='2027-01-01',p_s=6.,p_e=11.,cv_e=12.))
    app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30)
    app.session_state['case']=r.case;app.session_state['run']=r
    app.session_state['_workflow_stage']='계산내역';app.session_state['_calc_topic']='노드 계산'
    app.session_state['trace_i']=100;app.session_state['trace_j']=99
    monkeypatch.setattr(L,'engine',lambda *a,**k: (_ for _ in ()).throw(AssertionError('repriced')))
    app.run()
    assert not app.exception
    assert app.session_state['trace_i']==r.terms.n
    assert app.session_state['trace_j']==r.terms.n
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    app.radio(key='result_unit').set_value('주당 · 원').run()
    assert not app.exception
    assert app.session_state['run'].summary['run_id']==r.summary['run_id']
