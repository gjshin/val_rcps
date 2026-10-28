"""주식선택권(주식결제형) — 격자·블랙-숄즈·비용 배분·입력 검사·조서·화면.

숫자는 모두 가상의 값이다 (실제 회사 자료 아님).
"""
import copy
import io
import shutil
import zipfile

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import v2_workspace
from valuation import stock_option as SO
from valuation import legacy as L
from valuation.case import Case, inspect_case
from valuation.evidence import TOPICS, applicable
from valuation.service import calculate, export_bundle, CaseError
from test_v2_workflow import ROOT


def case(**over):
    rows = over.pop('so_rows', [dict(name='1차', vest='2028-01-01', end='', qty=1000.0, keep=1.0)])
    contract = dict(inst='SO', tranche='가상', d_issue='2026-01-01', d_mat='2031-01-01', K0=100.0, so_settle=0,
                    so_fy_end='12-31', so_rows=rows, so_mod=0, so_mod_rows=[])
    market = dict(S0=100.0, sig=0.3, div_y=0.0, rf_curve=[[1, 0.03], [5, 0.03]], y_type='spot', cmp_rf=1)
    method = dict(d_base='2026-01-01', so_ex=0, so_mult=2.0, so_exit=0.0, grid_days=7.0, gap_m=1.0, so_prev_liab=-1.0,
                  so_paid=0.0)
    for key, value in over.items():
        for group in (contract, market, method):
            if key in group:
                group[key] = value
                break
        else:
            raise KeyError(key)
    return Case(name='가상 주식선택권', contract=contract, market=market, method=method)


def fv(c):
    return calculate(c).raw['rows'][0]['fv']


def test_black_scholes_formula():
    # 교과서 값 — S = K = 100, T = 1, r = 5%, σ = 20%, 배당 없음 → 10.4506
    assert SO.black_scholes(100, 100, 1, 0.05, 0, 0.2) == pytest.approx(10.4506, abs=1e-4)
    assert SO.black_scholes(100, 120, 0, 0.05, 0, 0.2) == 0.0


def test_european_lattice_converges_to_black_scholes():
    run = calculate(case())
    r = run.raw['rows'][0]
    assert r['fv'] == pytest.approx(r['bs_end'], rel=SO.BS_TOL)
    assert abs(r['fv'] / r['bs_end'] - 1) < 0.002          # 주 단위 격자면 0.2% 안쪽
    assert run.raw['integrity'] == []
    check = next(c for c in run.summary['checks'] if c['name'].startswith('블랙-숄즈'))
    assert check['passed']


def test_no_dividend_optimal_exercise_equals_european_and_dividend_makes_it_worth_more():
    assert fv(case(so_ex=1)) == pytest.approx(fv(case(so_ex=0)), rel=1e-12)
    eu, am = fv(case(so_ex=0, div_y=0.05)), fv(case(so_ex=1, div_y=0.05))
    assert am > eu


def test_exercise_multiple_and_exit_rate_lower_the_value():
    best = fv(case(so_ex=1))
    multiple = fv(case(so_ex=2, so_mult=1.5))
    assert multiple < best
    assert fv(case(so_ex=2, so_mult=3.0)) > multiple           # 배수가 클수록 늦게 행사
    assert fv(case(so_ex=1, so_exit=0.1)) < best
    run = calculate(case(so_ex=2, so_mult=1.5, so_exit=0.1))
    r = run.raw['rows'][0]
    assert r['life'] < r['t_end']                              # 조기행사·퇴사로 기대존속기간이 짧다
    assert 'so_early_exercise' not in {i.code for i in run.issues}


def test_no_exercise_before_exercise_start():
    run = calculate(case(so_ex=1, div_y=0.08, so_rows=[dict(name='1차', vest='2027-01-01', start='2028-07-01',
                                                             end='', qty=10.0, keep=1.0)]))
    r = run.raw['rows'][0]
    assert r['ex_step'] > 0
    assert not any(any(g) for g in r['exercise'][:r['ex_step']])
    assert any(any(g) for g in r['exercise'][r['ex_step']:-1])  # 배당이 크면 행사 시작 뒤 조기행사가 생긴다


def test_node_trace_reconciles_every_node():
    run = calculate(case(so_ex=2, so_mult=1.6, so_exit=0.05))
    r = run.raw['rows'][0]
    worst = max(abs(SO.node_trace(run.terms, r, i, j)['reconciliation'])
                for i in range(0, r['n'] + 1, 5) for j in range(0, i + 1, 3))
    assert worst < 1e-9
    assert SO.node_trace(run.terms, r, r['n'], 0)['terminal']
    with pytest.raises(ValueError):
        SO.node_trace(run.terms, r, 3, 4)


def test_cost_schedule_is_day_proportional_and_stops_after_vesting():
    rows = [dict(name='1차', vest='2027-07-01', end='', qty=1000.0, keep=0.9),
            dict(name='2차', vest='2029-01-01', end='', qty=2000.0, keep=0.8)]
    run = calculate(case(so_rows=rows, so_ex=1))
    raw = run.raw
    one = raw['rows'][0]
    assert one['cost'] == pytest.approx(one['fv'] * 1000 * 0.9)
    s1 = [x for x in raw['schedule'] if x['name'] == '1차']
    assert [x['fy_end'] for x in s1] == ['2026-12-31', '2027-12-31']   # 가득된 뒤 결산일에는 줄이 없다
    days = (L.dt.date(2026, 12, 31) - L.dt.date(2026, 1, 1)).days
    span = (L.dt.date(2027, 7, 1) - L.dt.date(2026, 1, 1)).days
    assert s1[0]['share'] == pytest.approx(days / span)
    assert sum(x['expense'] for x in raw['schedule']) == pytest.approx(raw['cost'])
    assert run.summary['amounts_total']['so_cost'] == pytest.approx(raw['cost'])
    assert run.summary['amounts_per_share']['so_fair'] == pytest.approx(raw['fair'] / 3000)
    # 결산일이 3월 말인 회사
    run = calculate(case(so_rows=rows, so_fy_end='03-31'))
    assert run.raw['schedule'][0]['fy_end'] == '2026-03-31'


@pytest.mark.parametrize('over, words', [
    (dict(so_settle=2), '지원하지 않는 선택값'),
    (dict(so_settle=1, so_mod=1, d_base='2026-12-31'), '조건 변경 평가를 따로 하지 않습니다'),
    (dict(so_mod=1, d_base='2026-01-01', so_mod_rows=[dict(name='1차', grant_fv=10.0)]), '부여일보다 뒤'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', K=80.0)]), '부여일 1주당 공정가치'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[]), '조건 변경표에 없는 회차'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', grant_fv=10.0), dict(name='9차', grant_fv=1.0)]),
     '회차표에 없습니다'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', qty=900.0, grant_fv=10.0)]), '수량이 줄었습니다'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', K=0.0, grant_fv=10.0)]), '행사가격은 0보다'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', vest='2026-06-30', grant_fv=10.0)]),
     '조건 변경일보다 빠릅니다'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', price=1.0, grant_fv=10.0)]), '조건 변경표의 형식'),
    (dict(so_mod=1, d_base='2027-01-01', so_mod_rows=[dict(name='1차', end='2051-01-01', grant_fv=10.0)]),
     '1차 (변경 후): 격자가'),
    (dict(grid_days=1.0), '계산 한도'),
    (dict(so_settle=1, d_base='2026-12-31', so_paid=-1.0), '결제액은 0 이상'),
    (dict(so_paid=100.0), '현금결제형에서만'),
    (dict(so_ex=2, so_mult=1.0), '행사 배수'),
    (dict(so_exit=1.0), '퇴사율'),
    (dict(so_fy_end='12/31'), '결산일'),
    (dict(so_rows=[]), '한 줄 이상'),
    (dict(so_rows=[dict(name='1차', vest='2025-06-01', qty=1.0)]), '부여일보다 빠릅니다'),
    (dict(so_rows=[dict(name='1차', vest='2028-01-01', start='2027-01-01', qty=1.0)]), '가득일보다 빠릅니다'),
    (dict(so_rows=[dict(name='1차', vest='2028-01-01', end='2032-01-01', qty=1.0)]), '행사기간 종료일'),
    (dict(so_rows=[dict(name='1차', vest='2028-01-01', qty=0.0)]), '부여 수량'),
    (dict(so_rows=[dict(name='1차', vest='2028-01-01', qty=1.0, keep=0.0)]), '가득 예상 비율'),
    (dict(so_rows=[dict(name='A', vest='2028-01-01', qty=1.0), dict(name='A', vest='2029-01-01', qty=1.0)]), '겹칩니다'),
    (dict(so_rows=[dict(name='1차', vest='2028-01-01', qty=1.0, price=3)]), '형식'),
    (dict(so_exit=0.1, so_rows=[dict(name='1차', vest='2027-01-01', start='2028-01-01', qty=1.0)]), '행사 시작일이 가득일보다 늦은'),
])
def test_input_checks_block_calculation(over, words):
    c = case(**over)
    errors = [i for i in inspect_case(c) if i.severity == 'error']
    assert any(words in i.message for i in errors), [i.message for i in errors]
    with pytest.raises(CaseError):
        calculate(c)


def test_review_prompts_follow_the_standard():
    codes = {i.code for i in calculate(case(so_ex=0)).issues}
    assert {'so_early_exercise', 'so_vesting', 'so_dividend'} <= codes
    codes = {i.code for i in calculate(case(d_base='2026-06-30', so_ex=2, so_mult=2.0, div_y=0.01,
                                            so_rows=[dict(name='1차', vest='2028-01-01', qty=1.0, keep=0.9)])).issues}
    assert {'so_measure_date', 'so_multiple'} <= codes and 'so_vesting' not in codes
    assert 'normalized' not in codes                     # 사채 전용 칸을 끈 것은 알릴 일이 아니다


def test_only_stock_option_judgment_topics_apply():
    values = case().effective()
    ids = {t['id'] for t in TOPICS if applicable(t, values)}
    assert ids == {'so_volatility', 'so_early_exercise', 'so_vesting', 'so_modification', 'so_settlement',
                   'fair_value_inputs', 'method_change'}
    cb = {**values, 'inst': 'CB', 'view': 'issuer'}
    assert not any(applicable(t, cb) for t in TOPICS if t['trigger'] == 'so')


def _workbook(run, **kw):
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, **kw))) as z:
        return load_workbook(io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith('.xlsx')))))


def test_value_workpapers_carry_the_run_values():
    run = calculate(v2_workspace.sample_so_case())
    basic = _workbook(run, accounting=True)
    for name in ('평가요약', '회차별 공정가치', '회차표 입력', '블랙-숄즈 대조', '비용 배분', '회계처리', '금리곡선'):
        assert name in basic.sheetnames, name
    assert '행사방식' not in basic.sheetnames
    rows = list(basic['회차별 공정가치'].iter_rows(values_only=True))
    head = rows[0]
    # 엑셀 파일은 숫자를 유효숫자 16자리까지만 적는다 — 마지막 자리까지 같기를 요구하지 않는다
    for r, x in zip(rows[1:-1], run.summary['so_rows']):
        assert r[head.index('1주당 공정가치')] == pytest.approx(x['1주당 공정가치'], rel=1e-12)
    assert rows[-1][head.index('보상원가')] == pytest.approx(run.summary['amounts_total']['so_cost'])
    labels = [r[0] for r in basic['계약조건'].iter_rows(values_only=True)]
    assert '행사가격(원)' in labels and '현재 전환가액(원)' not in labels and '표면이자·우선배당률(연, %)' not in labels
    journal = [r for r in basic['회계처리'].iter_rows(values_only=True) if r and r[1] == '주식보상비용']
    assert sum(r[2] for r in journal) == pytest.approx(run.summary['amounts_total']['so_cost'])
    detail = _workbook(run, detail=True)
    lattices = [n for n in detail.sheetnames if n.startswith('격자 ')]
    assert len(lattices) == 3
    ws, r = detail[lattices[0]], run.raw['rows'][0]
    found = [c.value for row in ws.iter_rows() for c in row
             if isinstance(c.value, float) and abs(c.value - r['values'][0][0]) <= 1e-12*abs(r['values'][0][0])]
    assert found                                          # 격자 시점 0의 가치 = 회차 1주당 공정가치
    assert r['values'][0][0] == r['fv']
    inputs = [r[:2] for r in detail['평가 입력값'].iter_rows(values_only=True)]
    assert inputs[1] == ('측정기준일 주가(원)', run.terms.S0) and inputs[2] == ('변동성 σ (연)', run.terms.sig)


def test_screens_open_evaluate_and_show_calculations():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app.run()
    next(b for b in app.button if b.label == '주식선택권 합성사례로 시작하기').click().run()
    assert not app.exception, app.exception
    assert app.session_state.case.contract['inst'] == 'SO'
    assert '가득 회차' in [h.value for h in app.subheader]
    app.session_state['_workflow_stage'] = '평가·분석'
    app.run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    assert not app.exception, app.exception
    assert [h.value for h in app.subheader][:2] == ['회차별 공정가치', '회계연도별 보상원가 배분']
    assert any(m.label.startswith('1주당 공정가치') for m in app.metric)
    for topic in ('노드 계산표', '노드 따라가기', '금리·확률', '비용 배분'):
        app.session_state['_workflow_stage'] = '계산내역'
        app.session_state['_calc_topic'] = topic
        app.run()
        assert not app.exception, (topic, app.exception)
    app.session_state['_workflow_stage'] = '조서 출력'
    app.run()
    choice = next(r for r in app.radio if r.label == '조서 구성')
    assert choice.options == ['기본 값 조서', '상세 계산 값 조서', '상세 계산 수식 조서']
    app.session_state['_workflow_stage'] = '입력·시장자료'
    app.session_state['_input_area'] = '주가·변동성·금리 자료'
    app.run()
    tools = next(s for s in app.selectbox if s.label == '시장자료 도구')
    assert tools.options == ['변동성 산출', '기준일 주가 조회', '금리곡선 불러오기']


def test_new_stock_option_case_from_the_sidebar():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app.run()
    app.text_input[0].set_value('가상 부여 건')
    next(s for s in app.selectbox if s.label == '평가 상품').set_value('SO')
    next(b for b in app.button if b.label == '빈 입력안 만들기').click().run()
    assert not app.exception, app.exception
    c = app.session_state.case
    assert c.contract['inst'] == 'SO' and c.method['so_ex'] == 1
    missing = {i.field for i in inspect_case(c) if i.code == 'required'}
    assert {'S0', 'K0', 'sig'} <= missing


def test_exercise_at_measurement_date_still_exports():
    # 부여일 = 가득일 = 측정기준일이고 주가가 이미 행사 배수를 넘으면 뿌리에서 행사한다 — 기대존속기간 0년.
    # 블랙-숄즈 대조표가 0으로 나누지 않고 내재가치를 싣는다 (자동 검토 지적).
    c = case(S0=250.0, K0=100.0, so_ex=2, so_mult=2.0,
             so_rows=[dict(name='1차', vest='2026-01-01', end='', qty=10.0, keep=1.0)])
    run = calculate(c)
    r = run.raw['rows'][0]
    assert r['life'] == 0 and r['fv'] == pytest.approx(150.0)
    assert run.raw['schedule'][0]['share'] == 1.0            # 가득기간이 없으면 첫 결산일에 전액
    for detail in (False, True):
        wb = _workbook(run, detail=detail, accounting=True)
        rows = list(wb['블랙-숄즈 대조'].iter_rows(values_only=True))
        life_row = next(x for x in rows if x[1] == '기대존속기간')
        assert life_row[4] is None and life_row[8] == pytest.approx(150.0)


# ── 2단계 ① 조건 변경 · ② 현금결제형 ─────────────────────────────────────────

ONE = [dict(name='1차', vest='2028-01-01', end='', qty=1000.0, keep=0.9)]


def mod_case(mods, rows=ONE, **over):
    return case(so_mod=1, d_base='2027-01-01', so_rows=rows, so_mod_rows=mods, **over)


def test_repricing_measures_the_increment_on_the_modification_date():
    run = calculate(mod_case([dict(name='1차', K=80.0, grant_fv=20.0)]))
    r = run.raw['rows'][0]
    # 변경 전·후 값 = 같은 날·같은 시장자료로 따로 잰 부여일 평가 값 (문단 27, 부록 나43⑴)
    before = fv(case(d_base='2027-01-01', so_rows=ONE))
    after = fv(case(d_base='2027-01-01', so_rows=ONE, K0=80.0))
    assert r['fv_before'] == pytest.approx(before, rel=1e-12) and r['fv'] == pytest.approx(after, rel=1e-12)
    assert r['K'] == 80.0 and r['K_before'] == 100.0
    assert r['incr'] == pytest.approx(after - before) and r['incr'] > 0
    assert r['grant_cost'] == pytest.approx(20.0*1000*0.9)            # 당초 보상원가는 그대로 인식한다
    assert r['fair'] == pytest.approx(r['incr']*1000) and r['cost'] == pytest.approx(r['incr']*1000*0.9)
    total = run.summary['amounts_total']
    assert list(total) == ['so_incr_fair', 'so_incr_cost'] and total['so_incr_cost'] == pytest.approx(r['cost'])
    assert run.raw['integrity'] == []
    sched = run.raw['schedule']
    assert [(x['fy_end'], x['part']) for x in sched] == [
        ('2026-12-31', '당초'), ('2027-12-31', '당초'), ('2028-12-31', '당초'), ('2027-12-31', '증분'), ('2028-12-31', '증분')]
    inc = [x for x in sched if x['part'] == '증분']
    assert inc[0]['share'] == pytest.approx(364/365)                    # 증분은 변경일부터 가득일까지
    assert sum(x['expense'] for x in sched) == pytest.approx(r['grant_cost'] + r['cost'])
    assert 'so_mod_grant_fv' in {i.code for i in run.issues} and 'so_measure_date' not in {i.code for i in run.issues}
    check = next(c for c in run.summary['checks'] if c['name'].startswith('증분 공정가치'))
    assert check['passed']


def test_unfavourable_change_adds_nothing():
    run = calculate(mod_case([dict(name='1차', K=120.0, grant_fv=20.0)]))
    r = run.raw['rows'][0]
    assert r['fv'] < r['fv_before'] and r['incr'] == 0 and r['cost'] == 0
    assert all(x['expense'] == 0 for x in run.raw['schedule'] if x['part'] == '증분')   # 줄은 있고 금액은 0
    assert 'so_mod_no_increment' in {i.code for i in run.issues}


def test_added_options_are_valued_at_the_modified_value():
    run = calculate(mod_case([dict(name='1차', qty=1200.0, grant_fv=20.0)]))
    r = run.raw['rows'][0]
    assert r['incr'] == pytest.approx(0.0, abs=1e-12)                 # 다른 조건은 그대로
    assert r['qty_before'] == 1000 and r['qty_added'] == 200
    assert r['fair'] == pytest.approx(r['fv']*200) and r['cost'] == pytest.approx(r['fv']*200*0.9)
    assert run.raw['qty'] == 1000
    assert 'so_mod_no_increment' not in {i.code for i in run.issues}


def test_shorter_vesting_catches_up_without_rewriting_past_years():
    rows = [dict(name='1차', vest='2029-01-01', end='', qty=1000.0, keep=1.0)]
    run = calculate(mod_case([dict(name='1차', vest='2027-06-30', grant_fv=20.0)], rows=rows))
    first = [x for x in run.raw['schedule'] if x['part'] == '당초']
    span = (L.dt.date(2029, 1, 1) - L.dt.date(2026, 1, 1)).days
    assert first[0]['fy_end'] == '2026-12-31' and first[0]['share'] == pytest.approx(364/span)   # 지난 결산은 당초 추정
    assert first[1]['fy_end'] == '2027-12-31' and first[1]['cum'] == pytest.approx(20000.0)    # 줄인 가득일로 따라잡음
    assert len(first) == 2
    assert run.raw['rows'][0]['start'] == '2027-06-30'                # 가득일부터 행사하던 회차는 새 가득일부터
    late = [dict(name='1차', vest='2028-01-01', start='2028-07-01', end='', qty=1000.0, keep=1.0)]
    run = calculate(mod_case([dict(name='1차', vest='2027-06-30', grant_fv=20.0)], rows=late))
    assert run.raw['rows'][0]['start'] == '2028-07-01'                # 따로 정한 행사 시작일은 그대로
    assert 'so_mod_vest_shorter' in {i.code for i in run.issues}


def test_longer_vesting_keeps_the_original_cost_on_the_original_date():
    run = calculate(mod_case([dict(name='1차', K=80.0, vest='2029-01-01', grant_fv=20.0)]))
    first = [x for x in run.raw['schedule'] if x['part'] == '당초']
    inc = [x for x in run.raw['schedule'] if x['part'] == '증분']
    assert first[-1]['fy_end'] == '2028-12-31' and first[-2]['share'] == pytest.approx(729/730)   # 당초 가득일 2028-01-01
    assert inc[-1]['fy_end'] == '2029-12-31' and inc[0]['share'] == pytest.approx(364/(365 + 366))
    assert 'so_mod_vest_longer' in {i.code for i in run.issues}


def test_extending_the_exercise_period_past_the_original_expiry():
    # 회차 행사 종료일 빈칸 = 행사기간 종료일 2031-01-01. 연장한 종료일은 그보다 늦어도 되고, 변경 전은 그대로다 (자동 검토 지적)
    run = calculate(mod_case([dict(name='1차', end='2032-01-01', grant_fv=20.0)]))
    r = run.raw['rows'][0]
    assert r['end_before'] == '2031-01-01' and r['end'] == '2032-01-01'
    before = fv(case(d_base='2027-01-01', so_rows=ONE))
    after = fv(case(d_base='2027-01-01', d_mat='2032-01-01', so_rows=ONE))
    assert r['fv_before'] == pytest.approx(before, rel=1e-12) and r['fv'] == pytest.approx(after, rel=1e-12)
    assert r['incr'] > 0 and run.raw['integrity'] == []


def test_increment_after_vesting_is_recognised_at_once():
    rows = [dict(name='1차', vest='2026-06-30', end='', qty=1000.0, keep=1.0)]
    run = calculate(mod_case([dict(name='1차', K=80.0, grant_fv=20.0)], rows=rows))
    inc = [x for x in run.raw['schedule'] if x['part'] == '증분']
    assert len(inc) == 1 and inc[0]['fy_end'] == '2027-12-31' and inc[0]['share'] == 1.0
    assert inc[0]['expense'] == pytest.approx(run.raw['rows'][0]['cost'])
    first = [x for x in run.raw['schedule'] if x['part'] == '당초']
    assert len(first) == 1 and first[0]['fy_end'] == '2026-12-31'


def cash_case(d_base='2026-12-31', **over):
    rows = [dict(name='1차', vest='2028-01-01', end='', qty=1000.0, keep=0.9),
            dict(name='2차', vest='2026-06-30', end='', qty=500.0, keep=1.0)]
    return case(so_settle=1, d_base=d_base, so_rows=rows, **over)


def test_cash_settled_liability_is_remeasured_at_year_end():
    run = calculate(cash_case(so_prev_liab=5000.0))
    one, two = run.raw['rows']
    assert one['service'] == pytest.approx(364/730) and two['service'] == 1.0
    assert one['fv'] == pytest.approx(fv(case(d_base='2026-12-31', so_rows=ONE)), rel=1e-12)
    assert one['liab'] == pytest.approx(one['fv']*1000*0.9*364/730)
    assert run.raw['liab'] == pytest.approx(one['liab'] + two['liab'])
    assert run.raw['expense'] == pytest.approx(run.raw['liab'] - 5000.0)
    assert run.summary['so_cash']['당기 비용'] == pytest.approx(run.raw['expense'])
    assert list(run.summary['amounts_total']) == ['so_fair', 'so_liab']
    assert run.raw['integrity'] == []
    codes = {i.code for i in run.issues}
    assert 'so_cash_prev' not in codes and 'so_cash_date' not in codes and 'so_measure_date' not in codes
    codes = {i.code for i in calculate(cash_case(d_base='2026-11-30')).issues}
    assert {'so_cash_prev', 'so_cash_date'} <= codes
    # 당기 중 결제 — 부채 증감만 보면 결제한 4,000원만큼 비용이 줄어든다 (자동 검토 지적)
    paid = calculate(cash_case(so_prev_liab=5000.0, so_paid=4000.0))
    assert paid.raw['expense'] == pytest.approx(run.raw['expense'] + 4000.0)
    assert paid.summary['so_cash']['당기 현금 결제액'] == 4000.0
    wb = _workbook(paid, accounting=True)
    rows = list(wb['회계처리'].iter_rows(values_only=True))
    assert ('당기 중', '주식보상부채', 4000.0, '현금', 4000.0) in [r[:5] for r in rows]
    # 당기 중 모두 결제 — 남은 수량 0, 부채 0, 비용 = 결제액 − 전기말 부채 (자동 검토 지적)
    done = cash_case(so_prev_liab=10000.0, so_paid=12000.0)
    for r in done.contract['so_rows']:
        r['qty'] = 0.0
    gone = calculate(done)
    assert gone.raw['liab'] == 0 and gone.raw['expense'] == pytest.approx(2000.0)
    assert gone.summary['amounts_per_share'] == {'so_fair': 0.0, 'so_liab': 0.0}
    assert '부채 재측정' in _workbook(gone, detail=True, accounting=True).sheetnames
    assert any('부여 수량' in i.message for i in inspect_case(case(so_rows=[dict(name='1차', vest='2028-01-01', qty=0.0)])))


def test_mode_workpapers():
    wb = _workbook(calculate(mod_case([dict(name='1차', K=80.0, grant_fv=20.0)])), detail=True, accounting=True)
    for name in ('조건 변경', '조건 변경표 입력', '격자 1차 변경 후', '격자 1차 변경 전', '비용 배분', '회계처리'):
        assert name in wb.sheetnames, name
    assert '회차별 공정가치' not in wb.sheetnames
    run = calculate(cash_case(so_prev_liab=5000.0))
    wb = _workbook(run, accounting=True)
    assert '부채 재측정' in wb.sheetnames and '조건 변경' not in wb.sheetnames
    journal = [r for r in wb['회계처리'].iter_rows(values_only=True) if r and r[1] == '주식보상비용']
    assert sum(r[2] for r in journal) == pytest.approx(run.raw['expense'])


def test_screens_switch_between_grant_modification_and_cash():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app.run()
    app.session_state['case'] = mod_case([dict(name='1차', K=80.0, grant_fv=20.0)])
    app.run()
    assert not app.exception, app.exception
    assert '조건 변경' in [h.value for h in app.subheader]
    app.session_state['_workflow_stage'] = '평가·분석'
    app.run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    assert not app.exception, app.exception
    assert [h.value for h in app.subheader][:2] == ['회차별 조건 변경 전·후 공정가치', '회계연도별 보상원가 배분']
    stale = cash_case(so_prev_liab=5000.0)
    stale.contract['so_mod'] = 1                 # 조건 변경을 켠 채 현금결제형으로 바꾼 입력 (자동 검토 지적)
    app.session_state['case'] = stale
    app.session_state['_workflow_stage'] = '입력·시장자료'
    app.run()
    assert not app.exception, app.exception
    assert '조건 변경' not in [h.value for h in app.subheader]
    labels = [n.label for n in app.number_input]
    assert '당기 현금 결제액(원)' in labels and '전기말 주식보상부채(원, -1은 없음)' in labels
    app.session_state['_workflow_stage'] = '평가·분석'
    app.run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    assert not app.exception, app.exception
    assert app.session_state.case.contract['so_mod'] == 0
    assert [h.value for h in app.subheader][:2] == ['회차별 공정가치와 부채', '결산일 부채와 당기 비용']
    app.session_state['_workflow_stage'] = '조서 출력'
    app.run()
    next(r for r in app.radio if r.label == '조서 구성').set_value('상세 계산 수식 조서').run()
    next(b for b in app.button if b.label == '조서 생성').click().run()
    assert not app.exception, app.exception
    with zipfile.ZipFile(io.BytesIO(app.session_state.bundle)) as z:
        assert 'formula_review.xlsx' in z.namelist()

    # 결제액을 넣은 현금결제형을 주식결제형으로 바꾼 입력 (자동 검토 지적) — 숨은 결제액이 평가를 막지 않는다
    back = case(so_rows=ONE)
    back.method['so_paid'] = 4000.0
    app.session_state['case'] = back
    app.session_state['_workflow_stage'] = '입력·시장자료'
    app.run()
    assert '당기 현금 결제액(원)' not in [n.label for n in app.number_input]
    app.session_state['_workflow_stage'] = '평가·분석'
    app.run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    assert not app.exception, app.exception
    assert app.session_state.case.method['so_paid'] == 0.0
    assert [h.value for h in app.subheader][:2] == ['회차별 공정가치', '회계연도별 보상원가 배분']

# ── 2단계 ③ 수식 조서 — 상세 값 조서와 같은 시트·같은 칸, 다시 계산하면 같은 값 ─────────────────────

LIBRE = shutil.which('libreoffice') or shutil.which('soffice')
THREE = [dict(name='1차', vest='2027-01-01', end='2029-01-01', qty=1000.0, keep=0.9),
         dict(name='2차', vest='2028-01-01', end='', qty=2000.0, keep=0.8)]
FORMULA_CASES = {
    '부여일 · 행사 배수 · 퇴사율': lambda: case(so_ex=2, so_mult=1.8, so_exit=0.05, div_y=0.01, d_mat='2029-07-01',
                                       so_rows=THREE, grid_days=14.0),
    '부여일 · 최적 행사 · 배당': lambda: case(so_ex=1, div_y=0.03, d_mat='2029-07-01', so_rows=THREE, grid_days=14.0),
    '조건 변경': lambda: mod_case([dict(name='1차', K=80.0, qty=1200.0, vest='2028-07-01', grant_fv=20.0)],
                              grid_days=14.0, so_ex=1, div_y=0.02),
    '현금결제형': lambda: cash_case(so_prev_liab=5000.0, so_paid=1500.0, grid_days=14.0, so_ex=2, so_mult=2.0),
    '현금결제형 · 모두 결제': lambda: _settled(cash_case(so_prev_liab=10000.0, so_paid=12000.0, grid_days=14.0)),
    # 가득된 선택권의 행사기간 연장 — 주가가 행사 배수를 넘어 변경 전·후 모두 바로 행사하므로 처음 증분 0
    '조건 변경 · 증분 0': lambda: mod_case([dict(name='1차', end='2032-01-01', grant_fv=20.0)], grid_days=14.0, so_ex=2,
                                       so_mult=2.0, S0=250.0,
                                       rows=[dict(name='1차', vest='2026-06-30', end='', qty=1000.0, keep=1.0)]),
}


def _settled(c):
    for r in c.contract['so_rows']:
        r['qty'] = 0.0
    return c


def _xlsx(run, **kw):
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, **kw))) as z:
        return z.read(next(n for n in z.namelist() if n.endswith('.xlsx')))


def _recalc(data):
    from valuation.xlsx_validation import recalculate
    return load_workbook(io.BytesIO(recalculate(data)), data_only=True)


def test_formula_workbook_has_live_lattices_and_only_estimates_open():
    from valuation.xlsx_validation import inspect_workbook
    run = calculate(FORMULA_CASES['조건 변경']())
    data = _xlsx(run, detail=True, formula=True, accounting=True)
    assert not inspect_workbook(data)['errors']
    wb = load_workbook(io.BytesIO(data))
    value = load_workbook(io.BytesIO(_xlsx(run, detail=True, accounting=True)))
    assert wb.sheetnames == value.sheetnames                       # 값 조서와 같은 시트
    ws = wb['격자 1차 변경 후']
    assert ws['C14'].value == '=C24' and ws['C9'].value.startswith('=EXP(')
    assert all(isinstance(c.value, str) and c.value.startswith('=') for c in ws[24][2:5])
    assert isinstance(ws['C18'].value, float)                     # 금리는 앱이 정한 값 (엑셀에서 고치지 않는다)
    editable = {(s.title, c.coordinate) for s in wb for row in s.iter_rows() for c in row
                if c.protection.locked is False}
    assert editable == ({('평가 입력값', f'B{r}') for r in range(2, 7)} | {('회차표 입력', 'F2')} |
                        {('조건 변경표 입력', 'G2')})
    assert all(s.protection.sheet for s in wb)
    assert wb['계산대사']['C2'].value.startswith("='평가요약'!")


@pytest.mark.skipif(not LIBRE, reason='리브레오피스 없음')
@pytest.mark.parametrize('name', list(FORMULA_CASES))
def test_formula_workbook_recalculates_to_the_value_workbook(name):
    run = calculate(FORMULA_CASES[name]())
    value = load_workbook(io.BytesIO(_xlsx(run, detail=True, accounting=True)))
    live = _recalc(_xlsx(run, detail=True, formula=True, accounting=True))
    bad, count = [], 0
    for ws in value:
        if ws.title in ('00_안내', '계산대사'):                    # 출력 형식 안내 · 수식 대사 표는 조서마다 다르다
            continue
        other = live[ws.title]
        for row in ws.iter_rows():
            for c in row:
                a, b = c.value, other[c.coordinate].value
                if isinstance(a, (int, float)) and not isinstance(a, bool):
                    count += 1
                    if not (isinstance(b, (int, float)) and abs(a - b) <= 1e-9*max(1.0, abs(a))):
                        bad.append((ws.title, c.coordinate, a, b))
                elif a != b and not (a is None and b == ''):
                    bad.append((ws.title, c.coordinate, a, b))
    assert count > 10000 and not bad, bad[:10]
    check = list(live['계산대사'].iter_rows(min_row=2, values_only=True))
    assert check and all(abs(r[3]) <= 1e-9*max(1.0, abs(r[1])) for r in check), check


@pytest.mark.skipif(not LIBRE, reason='리브레오피스 없음')
def test_formula_workbook_follows_changed_estimates_like_the_app():
    c = FORMULA_CASES['부여일 · 행사 배수 · 퇴사율']()
    wb = load_workbook(io.BytesIO(_xlsx(calculate(c), detail=True, formula=True)))
    edits = {'S0': 110.0, 'sig': 0.35, 'so_mult': 2.0, 'so_exit': 0.03}
    rows = {'S0': 2, 'sig': 3, 'so_mult': 5, 'so_exit': 6}
    for k, v in edits.items():
        wb['평가 입력값'].cell(rows[k], 2).value = v
    wb['회차표 입력']['F3'].value = 0.7
    out = io.BytesIO()
    wb.save(out)
    live = _recalc(out.getvalue())
    c2 = copy.deepcopy(c)
    c2.market.update(S0=110.0, sig=0.35)
    c2.method.update(so_mult=2.0, so_exit=0.03)
    c2.contract['so_rows'][1]['keep'] = 0.7
    app = calculate(c2)
    from valuation.service import AMOUNT_LABELS
    lab = {r[0]: r for r in live['평가요약'].iter_rows(values_only=True)}
    for key, want in app.summary['amounts_total'].items():
        assert lab[AMOUNT_LABELS[key]][1] == pytest.approx(want, rel=1e-10), key
        assert lab[AMOUNT_LABELS[key]][3] == pytest.approx(app.summary['amounts_100'][key], rel=1e-10), key
    sched = [r for r in live['비용 배분'].iter_rows(min_row=2, values_only=True)]
    assert [x[3] for x in sched] == pytest.approx([x['당기 비용'] for x in app.summary['so_schedule']], rel=1e-10)


@pytest.mark.skipif(not LIBRE, reason='리브레오피스 없음')
def test_formula_workbook_keeps_rows_for_amounts_that_start_at_zero():
    # 처음에 0인 증분·결제액도 비용 배분·분개에 줄이 있어, 엑셀에서 입력을 바꾸면 따라온다 (자동 검토 지적)
    c = FORMULA_CASES['조건 변경 · 증분 0']()
    assert calculate(c).raw['rows'][0]['incr'] == 0
    wb = load_workbook(io.BytesIO(_xlsx(calculate(c), detail=True, formula=True, accounting=True)))
    wb['평가 입력값']['B2'].value = 150.0                  # 주가를 낮추면 바로 행사하지 않아 연장의 증분이 생긴다
    wb['조건 변경표 입력']['G2'].value = 10.0
    out = io.BytesIO()
    wb.save(out)
    live = _recalc(out.getvalue())
    c2 = copy.deepcopy(c)
    c2.market['S0'] = 150.0
    c2.contract['so_mod_rows'][0]['grant_fv'] = 10.0
    app = calculate(c2)
    assert app.raw['cost'] > 1000
    sched = [r for r in live['비용 배분'].iter_rows(min_row=2, values_only=True)]
    assert [x[4] for x in sched] == pytest.approx([x['당기 비용'] for x in app.summary['so_schedule']], rel=1e-10)
    journal = {r[0]: r[2] for r in live['회계처리'].iter_rows(min_row=3, values_only=True)
               if r and r[1] == '주식보상비용'}
    want = {}
    for x in app.summary['so_schedule']:
        want[x['결산일']] = want.get(x['결산일'], 0.0) + x['당기 비용']
    assert journal == pytest.approx(want, rel=1e-10)
    # 현금결제형 — 처음 결제액 0, 엑셀에서 3,000원을 넣는다
    c = cash_case(grid_days=14.0, so_prev_liab=5000.0)
    wb = load_workbook(io.BytesIO(_xlsx(calculate(c), detail=True, formula=True, accounting=True)))
    wb['평가 입력값']['B8'].value = 3000.0
    out = io.BytesIO()
    wb.save(out)
    rows = {r[0]: r for r in _recalc(out.getvalue())['회계처리'].iter_rows(values_only=True) if r and r[0]}
    assert rows['당기 중'][1:5] == ('주식보상부채', 3000.0, '현금', 3000.0)
    c.method['so_paid'] = 3000.0
    assert rows[c.method['d_base']][2] == pytest.approx(calculate(c).raw['expense'], rel=1e-10)
