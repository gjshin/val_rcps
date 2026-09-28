"""Independent payoffs and stale-evidence/data-transfer regression coverage."""
import copy
import io
import json
import math
import zipfile
from dataclasses import asdict
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from test_v2_workflow import synthetic, ROOT
from valuation import legacy
from valuation.case import Case, import_legacy
from valuation.evidence import evidence_cards, judgment_record, TOPICS
from valuation.service import calculate, refresh_run, calculation_key, export_bundle
from valuation.bridge import apply_changes, apply_volatility
from valuation.engagement import Engagement, calculate_engagement, engagement_bundle, value_bridge
from valuation.contract_analysis import calculate_schedule, scenario_key


def sample():
    tm=legacy.Terms(inst='RCPS', model='TF', d_issue='2026-01-01', d_base='2026-01-01', d_mat='2028-01-01',
                    S0=100., K0=100., issue_px=100., sig=.3, div_y=0., gap_m=12.,
                    cpn=0., div_mode=0, issuer_call=0, k_w=0., p_s=99.,p_e=0.,
                    rfx_mode=0, put_bdt=0, ipo_on=0, mat_mode=1, cv_s=99.,cv_e=0.,
                    y_type='spot', cmp_rf=1, cmp_cr=1, rf_curve=[[1,.05],[5,.05]], cr_curve=[[1,.08],[5,.08]])
    return import_legacy(asdict(tm), 'Independent example')


def row(date, **kwargs):
    result=dict(date=date, dividend_due=0., dividend_paid=0., redemption_allowed=False,
                redemption_amount=100., profit_limit=1000., cash_limit=1000., reset_price=0.)
    result.update(kwargs);return result


def scenario(**kwargs):
    result=dict(name='Explicit path', rationale='Synthetic hand calculation', end_date='2028-01-01',terminal='redeem',
                opening_unpaid=0., arrears_rate=0.,conversion_arrears='forfeit',
                schedule=[row('2028-01-01',redemption_allowed=True)])
    result.update(kwargs);return result


def test_unpaid_dividends_interest_and_paid_amount_hand_pv():
    case=sample()
    sc=scenario(opening_unpaid=5.,arrears_rate=.1,schedule=[
        row('2027-01-01',dividend_due=4.,dividend_paid=3.),
        row('2028-01-01',dividend_due=4.,dividend_paid=2.,redemption_allowed=True)])
    result=calculate_schedule(case,sc,intervals=2)
    unpaid1=5*1.1+4-3
    unpaid2=unpaid1*1.1+4-2
    assert result['value_100']==pytest.approx(3/1.08+(2+100+unpaid2)/1.08**2)
    assert result['timeline'][-1]['원금100당 미지급누적']==pytest.approx(unpaid2)


def test_terminal_stock_option_independent_binomial_sum():
    case=sample();case.contract.update(cv_s=24.,cv_e=24.)
    sc=scenario()
    u=math.exp(.3);d=1/u;q=(1.05-d)/(u-d)
    expected=0
    for ups, probability in [(0,(1-q)**2),(1,2*q*(1-q)),(2,q*q)]:
        stock=100*u**ups*d**(2-ups)
        expected+=probability*(stock/1.05**2 if stock>100 else 100/1.08**2)
    assert calculate_schedule(case,sc,intervals=2)['value_100']==pytest.approx(expected)


def test_event_reset_applies_on_date_and_terminal_conversion():
    case=sample()
    base=scenario(terminal='convert')
    reset=scenario(terminal='convert',schedule=[row('2027-01-01',reset_price=50.),row('2028-01-01')])
    assert calculate_schedule(case,base,intervals=2)['value_100']==pytest.approx(100)
    result=calculate_schedule(case,reset,intervals=2)
    assert result['value_100']==pytest.approx(200)
    assert result['timeline'][0]['조정 후 전환가액(원)']==50


def test_redemption_requires_both_profit_and_cash():
    case=sample()
    for limits in [{'profit_limit':50.},{'cash_limit':50.}]:
        sc=scenario(schedule=[row('2027-01-01',redemption_allowed=True,**limits),row('2028-01-01',redemption_allowed=True)])
        assert calculate_schedule(case,sc,intervals=2)['value_100']==pytest.approx(100/1.08**2)
    open_sc=scenario(schedule=[row('2027-01-01',redemption_allowed=True),row('2028-01-01',redemption_allowed=True)])
    assert calculate_schedule(case,open_sc,intervals=2)['value_100']==pytest.approx(100/1.08)


def test_extension_and_conversion_arrears_policy():
    case=sample()
    end='2029-01-01';T=(__import__('datetime').date.fromisoformat(end)-__import__('datetime').date(2026,1,1)).days/365
    sc=scenario(end_date=end,opening_unpaid=5.,arrears_rate=.1,schedule=[row(end,redemption_allowed=True)])
    assert calculate_schedule(case,sc,intervals=3)['value_100']==pytest.approx((100+5*1.1**T)/1.08**T)
    for policy, extra in [('pay',5/1.08**2),('forfeit',0)]:
        sc=scenario(terminal='convert',opening_unpaid=5.,conversion_arrears=policy)
        assert calculate_schedule(case,sc,intervals=2)['value_100']==pytest.approx(100+extra)


@pytest.mark.parametrize('mutation',[
    lambda s:s.update(opening_unpaid=float('nan')),
    lambda s:s['schedule'][0].update(cash_limit=10.),
    lambda s:s['schedule'][0].update(dividend_paid=5.),
    lambda s:s['schedule'].insert(0,row('2028-01-01')),
    lambda s:s.update(end_date='2027-01-01'),
])
def test_incomplete_or_infeasible_schedule_is_blocked(mutation):
    sc=scenario();mutation(sc)
    with pytest.raises(ValueError):calculate_schedule(sample(),sc)


def test_malformed_saved_schedule_rejected_before_ui_render():
    case=sample();case.contract_scenarios=[scenario()]
    payload=case.to_dict();payload['contract_scenarios'][0]['end_date']=123
    with pytest.raises(ValueError):Case.from_dict(payload)


@pytest.mark.parametrize('field,value',[('k_w',.2),('rfx_mode',1),('ipo_on',1)])
def test_unsupported_interactions_never_silently_disappear(field,value):
    case=sample();case.contract[field]=value
    with pytest.raises(ValueError):calculate_schedule(case,scenario())


def test_judgment_stale_after_input_change_and_refresh_without_pricing(monkeypatch):
    case=synthetic()
    case.judgments['classification']=judgment_record(case,'검토 완료','Example conclusion','Contract reviewed','Reviewer','§3')
    run=calculate(case)
    assert next(c for c in evidence_cards(case) if c['id']=='classification')['status']=='검토 완료'
    case.market['S0']+=1
    assert next(c for c in evidence_cards(case) if c['id']=='classification')['review_stale']
    case.market['S0']-=1
    case.judgments['classification']['rationale']='Updated review'
    monkeypatch.setattr(legacy,'decompose',lambda *_: (_ for _ in ()).throw(AssertionError('Unexpected pricing')))
    updated=refresh_run(run,case)
    with zipfile.ZipFile(io.BytesIO(export_bundle(updated))) as z:
        assert 'judgment_evidence.json' in z.namelist()
        wb=load_workbook(io.BytesIO(z.read('value_review.xlsx')))
        assert '판단근거' in wb.sheetnames
        assert any(c.value=='Updated review' for row in wb['판단근거'] for c in row)


def test_source_cards_have_verifiable_locations_and_small_quotes():
    unique={}
    for t in TOPICS:
        assert t['questions'] and t['limitation']
        for s in t['sources']:
            assert s['location'] and s['checked_on']
            if 'pdf_page' in s:assert s['pdf_page']==s['printed_page']+20 and s['sha256']
            if s.get('url') and s['quote']:
                unique.setdefault(s['url'],set()).add(s['quote'])
    assert all(sum(len(q.split()) for q in quotes)<=25 for quotes in unique.values())


def test_changed_source_or_additional_clause_requires_rereview():
    case=synthetic()
    case.judgments['classification']=judgment_record(case,'검토 완료','Conclusion','Evidence','Reviewer','§3')
    case.sources['contract']='Updated contract amendment'
    assert next(c for c in evidence_cards(case) if c['id']=='classification')['review_stale']
    assert Case.from_dict(case.to_dict()).judgments == case.judgments


def test_legacy_export_sources_leave_numerical_formulas_untouched():
    from openpyxl import Workbook
    from valuation.evidence import attach_evidence
    wb=Workbook();wb.active['A1']='=2+3';wb.active['A2']=5
    raw=io.BytesIO();wb.save(raw)
    case=synthetic();case.name='=NO_EXECUTION()'
    updated=load_workbook(io.BytesIO(attach_evidence(raw.getvalue(),case)))
    assert updated['Sheet']['A1'].value=='=2+3'
    assert updated['Sheet']['A1'].data_type=='f'
    assert updated['Sheet']['A2'].value==5
    assert updated['판단근거']['B2'].value=='=NO_EXECUTION()'
    assert updated['판단근거']['B2'].data_type=='s'


def test_legacy_cli_refuses_implicit_formula_approximation_before_writing(tmp_path):
    import subprocess, sys
    case=synthetic();case.contract['rfx_mode']=1;case.method['carry']=0
    input_path=tmp_path/'case.json';input_path.write_text(json.dumps(case.effective()))
    target=tmp_path/'output'
    run=subprocess.run([sys.executable,str(ROOT/'tools/run_valuation.py'),str(input_path),'--out',str(target)],capture_output=True,text=True)
    assert run.returncode != 0
    assert '--no-formula' in run.stderr
    assert not target.exists()


def test_legacy_import_preserves_values_and_rejects_case_envelopes():
    from valuation.bridge import legacy_term_values
    from tools.run_valuation import make_terms
    case = synthetic()
    values = case.effective()
    values['_schema'] = 4
    assert legacy_term_values(values)['S0'] == case.market['S0']
    assert legacy_term_values(values)['rf_curve'] == case.market['rf_curve']
    for payload in (case.to_dict(), {'schema':'valuation-engagement/1','cases':[]}, [], {'notes':'memo'}):
        with pytest.raises(ValueError):
            legacy_term_values(payload)
        with pytest.raises(ValueError):
            make_terms(vars(legacy), payload, {})


def test_bridge_changes_override_without_destroying_contract_facts():
    case=synthetic();original=case.market['S0']
    case.assumptions=[dict(field='S0',value=original+5,rationale='Conditional evidence')]
    updated=apply_changes(case,{'S0':original+10,'sig':.4})
    assert updated.market['S0']==original
    assert updated.effective()['S0']==original+10
    assert updated.market['sig']==.4
    assert case.effective()['S0']==original+5


def test_volatility_wrong_date_blocked():
    with pytest.raises(ValueError):apply_volatility(synthetic(),dict(kind='vol_pack',sigma=.3,opt={'asof':'1999-01-01'}))


def test_multi_tranche_reuse_total_and_safe_zip_paths(monkeypatch):
    c1=synthetic();c2=synthetic();c2.contract['face_total']=2*c1.contract['face_total']
    eng=Engagement('Example','Issuer',{'../outside':c1,'Two':c2})
    result=calculate_engagement(eng)
    assert result['totals']['순포지션_원']==pytest.approx(result['rows'][0]['순포지션_원']*3)
    monkeypatch.setattr(legacy,'decompose',lambda *_: (_ for _ in ()).throw(AssertionError('Unexpected pricing')))
    second=calculate_engagement(eng,result['runs'])
    assert second['totals']==result['totals']
    with zipfile.ZipFile(io.BytesIO(engagement_bundle(eng,second))) as z:
        assert all('..' not in name for name in z.namelist())
        assert len([n for n in z.namelist() if n.endswith('.zip')])==2
    eng.cases['Two'].market['S0']+=10
    with pytest.raises(ValueError):engagement_bundle(eng,second)


def test_tranche_date_view_and_duplicate_id_protection():
    c1=synthetic();c2=synthetic();c2.method['d_base']='2025-02-01'
    with pytest.raises(ValueError):calculate_engagement(Engagement('Example','Issuer',{'A':c1,'B':c2}))
    obj=Engagement('Example','Issuer',{'A':c1}).to_dict();obj['cases']*=2
    with pytest.raises(ValueError):Engagement.from_dict(obj)


def test_value_bridge_reconciles_independently_linear_common_stock_case():
    previous=synthetic();current=copy.deepcopy(previous)
    current.market['S0']=60.
    bridge=value_bridge(previous,current)
    assert sum(r['변동액_원'] for r in bridge['rows'])==pytest.approx(.2*current.contract['face_total'])
    assert bridge['closing']-bridge['opening']==pytest.approx(sum(r['변동액_원'] for r in bridge['rows']))
    assert abs(bridge['reconciliation_difference'])<1e-6


@pytest.mark.parametrize('page',['평가 작업','판단 근거','계약조건 분석','여러 회차·변동 분석','상세 기능','변동성 산출'])
def test_all_routes_open_without_implicit_valuation(page,monkeypatch):
    def forbidden(*args,**kwargs):raise AssertionError('Rendering invoked pricing')
    monkeypatch.setattr(legacy,'decompose',forbidden)
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=15)
    app.session_state.case=synthetic()
    app.run()
    app.sidebar.radio[0].set_value(page).run()
    assert not app.exception


def test_detailed_tool_all_sections_and_return_preserve_current_case():
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=40)
    original=synthetic();app.session_state.case=original
    app.run();app.sidebar.radio[0].set_value('상세 기능').run()
    next(b for b in app.button if b.label=='현재 평가를 상세 기능에 가져오기').click().run()
    assert not app.exception
    assert app.session_state.tm.gap_m==original.method['gap_m']
    next(b for b in app.button if b.label=='상세 입력으로 계산').click().run()
    assert not app.exception
    section=next(s for s in app.selectbox if s.label=='상세 분석 항목')
    for name in section.options:
        next(s for s in app.selectbox if s.label=='상세 분석 항목').set_value(name).run()
        assert not app.exception, name
    app.sidebar.radio[0].set_value('평가 작업').run()
    assert not app.exception
    assert app.session_state.case.to_dict()==original.to_dict()
