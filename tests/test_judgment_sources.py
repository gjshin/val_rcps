"""판단·근거 탭 · 근거 원문 · Day 1 (보정 기록·이어 적용·이연/당기손익)."""
import copy
import io
import re
import zipfile

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from synthetic_rcps_case import sample
from test_v2_workflow import ROOT
from valuation import legacy, sources
from valuation.case import Case, inspect_case
from valuation.evidence import TOPICS
from valuation.service import calculate, calculation_key, export_bundle, refresh_run


def call_case():
    case = sample()
    case.contract.update(p_s=0., p_e=12., issuer_call=2, k_w=.3, k_s=3., k_e=9., k_prem=.02)
    case.method.update(conv_class='equity', k_method=0)
    return case


def open_tab(case, tab='판단·근거'):
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app.session_state.case = case
    app.run()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    next(w for w in app.selectbox if w.label == '분석 도구').set_value('상세 계산·회계 참고표').run()
    next(w for w in app.selectbox if w.label == '상세 분석 항목').set_value(tab).run()
    return app


def test_every_reference_resolves_to_text():
    for topic in sources.REFS:
        for key in sources.refs(topic):
            rows = sources.lookup(key)
            assert rows and all(r['text'].strip() for r in rows), key


def test_book_pages_contain_the_short_quotes_used_in_evidence_cards():
    norm = lambda x: re.sub(r'\s+', '', x)
    pages = sources._book()[0]
    for tp in TOPICS:
        for src in tp['sources']:
            if src['kind'] != '실무사례·해설':
                continue
            p = src['printed_page']
            text = ''.join(norm(pages[q]['text']) for q in range(p - 1, p + 2) if q in pages)
            assert norm(src['quote'])[:30] in text, (tp['id'], p)


def test_judgment_tab_is_light_until_buttons(monkeypatch):
    for name in ('call_compare', 'pc_compare', 'rate_signals'):
        monkeypatch.setattr(legacy, name, lambda *a, **k: (_ for _ in ()).throw(AssertionError(name)))
    app = open_tab(call_case())
    assert not app.exception, app.exception
    text = ' '.join(m.value for m in app.markdown)
    for head in ('분리 판정', '조기상환 행사 진단', '매도청구권 평가방법', '풋·콜 우선순위', '이자율모형(BDT) 검토', '추가 검토 항목'):
        assert head in text, head
    boxes = ' '.join(str(w.value) for kind in ('success', 'info') for w in app.get(kind))
    assert '상각후원가' in boxes and '기준 10%' in boxes


def test_memo_saved_to_case_and_workpaper():
    app = open_tab(call_case())
    box = next(w for w in app.text_input if w.key and w.key.startswith('memo_r_split_put'))
    box.set_value('행사금액이 상각후원가와 40% 차이 — 분리').run()
    next(b for b in app.button if b.key and b.key.startswith('memo_s_split_put')).click().run()
    assert not app.exception
    case = app.session_state.case
    assert case.memos['split_put']['reason'].startswith('행사금액')
    # 판단 기록은 계산키를 바꾸지 않는다 — 재평가 없이 조서에 실린다.
    assert app.session_state.run.summary['calculation_key'] == calculation_key(case)
    with zipfile.ZipFile(io.BytesIO(export_bundle(app.session_state.run))) as z:
        wb = load_workbook(io.BytesIO(z.read('value_review.xlsx')))
    rows = [[c.value for c in r] for r in wb['판단·근거']]
    put = next(r for r in rows if r[1] == '조기상환청구권')
    assert put[4] == '앱 판정에 동의' and put[5].startswith('행사금액') and '1109 B4.3.5' in put[6]
    # 원문(기준서 본문·책 발췌)은 조서에 싣지 않는다 — 출처 표기만 남는다.
    assert not any(r[0] == '근거 원문 발췌' for r in rows)
    assert not any(r[0] == '1109 B4.3.5' and '상각후원가' in (r[3] or '') for r in rows)
    assert '실무사례' in put[6]


def test_judgment_sheets_can_be_left_out():
    run = calculate(call_case())
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, judgment=False))) as z:
        wb = load_workbook(io.BytesIO(z.read('value_review.xlsx')))
    assert '판단·근거' not in wb.sheetnames


def day1_case():
    case = sample()
    case.method['d_base'] = case.contract['d_issue']
    return case


def test_day1_gap_choice_and_journal():
    run = calculate(day1_case())
    d1 = run.summary['day1']
    assert d1['diff'] == pytest.approx(run.raw['b2'] - run.raw['ca'] - 100)
    assert any(i.code == 'day1_gap' for i in run.issues)
    h = legacy.holder_rows(run.terms, *[run.raw[k] for k in ('full', 'b0', 'b1', 'b2', 'ca')])
    assert any(a.startswith('최초 인식 차이 — 이연') for _, a, _ in h['journal'])
    pl = Case.from_dict(run.case.to_dict()); pl.method['d1_pl'] = 1
    assert any(i.code == 'day1_reason' for i in inspect_case(pl))
    pl.method['d1_reason'] = '관측 가능한 시장자료만 사용'
    assert calculation_key(pl) == run.summary['calculation_key']
    again = refresh_run(run, pl)
    h = legacy.holder_rows(again.terms, *[again.raw[k] for k in ('full', 'b0', 'b1', 'b2', 'ca')])
    side, acct, v = next(j for j in h['journal'] if '최초 인식 차이' in j[1])
    assert acct.startswith('금융자산평가손실' if d1['diff'] < 0 else '금융자산평가이익')
    assert v == pytest.approx(abs(d1['diff']))
    for kw in (dict(accounting=True), dict(formula=True, accounting=True)):
        with zipfile.ZipFile(io.BytesIO(export_bundle(again, **kw))) as z:
            name = next(n for n in z.namelist() if n.endswith('.xlsx'))
            wb = load_workbook(io.BytesIO(z.read(name)))
        texts = [c.value for r in wb['회계처리'] for c in r if isinstance(c.value, str)]
        assert any('금융자산평가' in x and '최초 인식 차이' in x for x in texts)
        checks = [r[1].value for r in wb['판단·근거'].iter_rows() if r[0].value == '최초 인식 · 원인 점검']
        assert len(checks) == 3


def test_calibration_record_and_carry_forward_in_ui():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app.session_state.case = day1_case(); app.run()
    app.radio(key='_input_area').set_value('주가·변동성·금리 자료').run()
    next(w for w in app.selectbox if w.label == '시장자료 도구').set_value('주가 역산').run()
    next(w for w in app.text_area if w.label == '거래가격을 목표로 사용하는 근거').set_value('정상 거래 · 동일 권리 범위').run()
    next(b for b in app.button if b.label == '주가 역산 실행').click().run()
    assert not app.exception, app.exception
    fitted = app.session_state._backsolve_result[1]
    next(b for b in app.button if b.label == '역산 주당가치 적용').click().run()
    cal = app.session_state.case.calibration
    assert cal['after'] == pytest.approx(fitted) and cal['equity_ps'] > 0 and cal['target'] == 'S0'
    run = calculate(app.session_state.case)
    assert abs(run.summary['day1']['diff']) < 0.005
    with zipfile.ZipFile(io.BytesIO(export_bundle(run))) as z:
        wb = load_workbook(io.BytesIO(z.read('value_review.xlsx')))
    assert '보정기록' in wb.sheetnames
    # 다음 분기 — 같은 파일을 복사해 평가기준일만 옮긴다.
    nxt = Case.from_dict(app.session_state.case.to_dict()); nxt.method['d_base'] = '2026-06-30'
    app2 = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app2.session_state.case = nxt; app2.run()
    app2.radio(key='_input_area').set_value('주가·변동성·금리 자료').run()
    next(w for w in app2.selectbox if w.label == '시장자료 도구').set_value('보정 이어 적용').run()
    now = cal['equity_ps'] * 1.2
    next(w for w in app2.number_input if w.label == '이번 평가기준일 지분평가 주당가치(원)').set_value(now).run()
    next(b for b in app2.button if b.label == '이어 적용').click().run()
    assert not app2.exception
    assert app2.session_state.case.market['S0'] == pytest.approx(cal['after'] * 1.2)
    assert '보정 이어 적용' in app2.session_state.case.market['s0_src']


def test_day1_panel_in_results():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180)
    app.session_state.case = day1_case(); app.run()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    assert any('최초 인식 — 모형값' in str(w.value) for w in app.warning)
    save = next(b for b in app.button if b.key and b.key.startswith('d1_save'))
    next(w for w in app.selectbox if w.label == '최초 인식 차이 처리').set_value(1).run()
    save = next(b for b in app.button if b.key and b.key.startswith('d1_save'))
    assert save.disabled                       # 근거 없이는 당기손익을 고를 수 없다
    next(w for w in app.text_input if w.label.startswith('당기손익 근거')).set_value('관측 가능한 시장자료만 사용').run()
    next(b for b in app.button if b.key and b.key.startswith('d1_save')).click().run()
    assert not app.exception
    assert app.session_state.case.method['d1_pl'] == 1
    assert app.session_state.run.terms.d1_pl == 1   # 재평가 없이 반영


def test_sources_show_citations_only_unless_internal(monkeypatch):
    # 배포본은 출처만 — 원문 파일이 없어도 자료명·문단·쪽을 풀어 쓴다.
    monkeypatch.delenv('VAL_INTERNAL_SOURCES', raising=False)
    assert not sources.internal()
    assert sources.citation('1109:B4.3.5') == 'K-IFRS 제1109호 문단 B4.3.5'
    assert sources.citation('book:26-32').endswith('26~32쪽')
    assert sources.citation('KGAAP15:15.20') == '일반기업회계기준 제15장 문단 15.20'
    monkeypatch.setenv('VAL_INTERNAL_SOURCES', '1')
    assert sources.internal()
