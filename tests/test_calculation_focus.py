"""Synthetic-only acceptance of the calculation-focused workflow."""
import copy
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from valuation import legacy
from valuation.case import Case, inspect_case
from valuation.market_data import query_spec, make_pack, validate_pack, parse_price_table, export_volatility_workbook
from valuation.service import calculate, export_bundle
from valuation.xlsx_validation import inspect_workbook, compare_cells
from test_v2_workflow import synthetic
from test_unified_workflow import prices

OMITTED = {'검토기록','V2_검토기록','V2_계약과가정','V2_추가권리','판단근거','확정상태',
           '계약반영표','독립검산대사','계약검토안','추가확인자료','별도계약조건'}


def workbook(run, formula=False, detail=False, accounting=False):
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=formula, detail=detail, accounting=accounting))) as z:
        return load_workbook(io.BytesIO(z.read('formula_review.xlsx' if formula else 'value_review.xlsx')))


def test_recalc_check_compares_app_values_not_a_removed_sheet():
    """재계산 검사가 조서에서 뺀 «수식대사» 시트 대신 앱 값을 들고 와 견준다."""
    run = calculate(synthetic())
    r = run.raw
    eir = legacy.eir_or_none(run.terms, r['full'], r['b0'], r['b1'], r['b2'], r['ca'])
    keys = legacy.formula_key_cells(run.terms, r, eir)
    assert {k[0] for k in keys} >= {'전체 (적용 물량)', '주계약', '매도청구권 (적용값)', '상각표 유효이자율'}
    wb = workbook(run, formula=True, accounting=True)
    for label, sheet, cell, _ in keys:           # 모두 수식 칸을 가리킨다 — 값으로 굳은 칸이 아니다
        v = wb[sheet][cell].value
        assert isinstance(v, str) and v.startswith('='), (label, sheet, cell, v)
    from openpyxl import Workbook
    fake = Workbook(); ws = fake.active; ws.title = '결과'
    ws['C10'] = 100.0; ws['C16'] = 37.5
    rows = compare_cells(fake, [('a', '결과', 'C10', 100.0 + 1e-9), ('b', '결과', 'C16', 37.6),
                                ('c', '상각표', 'C10', 0.1)])
    assert [x[4] for x in rows] == ['일치', '불일치', '불일치']


@pytest.mark.parametrize('days', [0,7,14])
def test_grid_preserves_legacy_month_and_allows_explicit_week(days):
    case = synthetic(); case.method['grid_days'] = float(days)
    run = calculate(case)
    direct = legacy.Terms(**case.effective()); expected = legacy.decompose(direct)
    assert run.raw['b2'] == expected[3]
    assert run.terms.n == (max(4,round(run.terms.T*365/days)) if days else max(4,round(run.terms.T*12/run.terms.gap_m)))
    assert not [i for i in inspect_case(case) if i.severity=='error']


@pytest.mark.parametrize('formula,detail', [(False,False),(False,True),(True,True)])
@pytest.mark.parametrize('holder', [False,True])
def test_optional_accounting_and_no_review_sheets(formula, detail, holder):
    case=synthetic();case.method['view']='holder' if holder else 'issuer'
    run=calculate(case)
    basic=workbook(run,formula,detail)
    assert not OMITTED & set(basic.sheetnames)
    assert not {'회계처리','상각표'} & set(basic.sheetnames)
    with_acc=workbook(run,formula,detail,True)
    assert '회계처리' in with_acc
    assert any('초안' in str(c.value) for row in with_acc['회계처리'].iter_rows(min_row=1,max_row=2) for c in row)
    assert not OMITTED & set(with_acc.sheetnames)
    # Adding accounting never changes the original calculation.
    again=calculate(case)
    assert again.summary['amounts_100']==run.summary['amounts_100']


def pack_with_exceptions():
    series=[['A',prices()[:-1]]]
    q=query_spec(False,['A,A','B,B'],30,'2025-01-01',allow_failed=True,allow_missing=True)
    exclusions=[dict(name='A',date=prices()[-1][0],reason='주가 누락'),dict(name='B',date='',reason='조회 실패')]
    return make_pack(series,q,tdays=250,drop=False,pick='median',source='Synthetic',exclusions=exclusions)


def test_market_exceptions_are_explicit_and_exported():
    pack=pack_with_exceptions();assert validate_pack(pack)
    assert any('B: 조회 실패' in w for w in pack['warnings'])
    assert any('29개' in w for w in pack['warnings'])
    strict=query_spec(False,['A,A','B,B'],30,'2025-01-01')
    with pytest.raises(ValueError):make_pack(pack['series'],strict,tdays=250,drop=False,pick='median',source='test')
    with pytest.raises(ValueError):
        changed=copy.deepcopy(pack);changed['exclusions'][0]['reason']='tampered';validate_pack(changed)
    wb=load_workbook(io.BytesIO(export_volatility_workbook(pack)))
    assert '조회·제외내역' in wb
    assert any('B: 조회 실패' in str(c.value) for row in wb['조회·제외내역'] for c in row)


def test_missing_cells_do_not_silently_disappear_or_allow_zero():
    source='date,A,B\n2024-01-01,100,\n2024-01-02,101,200\n'
    with pytest.raises(ValueError):parse_price_table(source)
    series,exclusions=parse_price_table(source,allow_missing=True,return_exclusions=True)
    assert len(series[1][1])==1 and exclusions==[dict(name='B',date='2024-01-01',reason='주가 누락')]
    with pytest.raises(ValueError):parse_price_table(source.replace('100','0'),allow_missing=True)
    with pytest.raises(ValueError):parse_price_table(source.replace('2024-01-02','2024-01-01'),allow_missing=True)


def test_numerical_warnings_and_listed_raw_prices_are_in_export():
    case=synthetic();case.contract.update(base_shares=1630868.,dil_shares=473729.)
    series=[['Listed',prices()]]
    q=query_spec(True,['A,Listed'],30,case.effective()['d_base'])
    pack=make_pack(series,q,tdays=250,drop=False,pick='median',source='Synthetic')
    case.market_evidence['sig']=pack;case.market['sig']=pack['sigma']
    run=calculate(case)
    assert any('29.05%' in i.message for i in run.issues)
    for formula in (False,True):
        wb=workbook(run,formula,True)
        # 변동성 산출내역(σ 시트)이 있으면 원자료도 그 안에 있다 — 같은 주가를 두 번 싣지 않는다.
        if '변동성원자료' in wb.sheetnames:
            raw = list(wb['변동성원자료'].values)[1:]
            assert [(r[0],r[1]) for r in raw] == [('Listed',day) for day,price in series[0][1]]
            assert [r[2] for r in raw] == pytest.approx([price for day,price in series[0][1]],rel=1e-14)
        else:
            sig = [n for n in wb.sheetnames if n.startswith('σ ')]
            vals = [c.value for n in sig for row in wb[n] for c in row if isinstance(c.value, (int, float))]
            assert all(any(abs(v - price) < 1e-9 for v in vals) for day, price in series[0][1])
        assert any('29.05%' in str(c.value) for row in wb['조서 정보'] for c in row)


@pytest.mark.parametrize('days',[0,7])
def test_screen_to_excel_without_review_or_final_gate(days):
    case=synthetic()
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=120)
    app.session_state.case=case;app.run()
    next(w for w in app.selectbox if w.label=='계산 간격 단위').set_value(days).run()
    app.radio(key='_input_area').set_value('주가·변동성·금리 자료').run()
    assert not app.exception
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    assert not app.exception
    actual=app.session_state.run
    assert actual.summary['amounts_100']==calculate(app.session_state.case).summary['amounts_100']
    next(w for w in app.selectbox if w.label=='분석 도구').set_value('상세 계산·회계 참고표').run()
    for section in ['판단·근거','이자율곡선','주가·변동성','검산']:
        selector=next(w for w in app.selectbox if w.label=='상세 분석 항목')
        if section in selector.options:
            selector.set_value(section).run()
            assert not app.exception
    app.radio(key='_workflow_stage').set_value('조서 출력').run()
    for option in ['기본 값 조서','상세 계산 값 조서','상세 계산 수식 조서']:
        next(w for w in app.radio if w.label=='조서 구성').set_value(option).run()
        next(b for b in app.button if b.label=='조서 생성').click().run()
        assert not app.exception and not app.error
        with zipfile.ZipFile(io.BytesIO(app.session_state.bundle)) as z:
            name=next(n for n in z.namelist() if n.endswith('.xlsx'))
            assert not inspect_workbook(z.read(name))['errors']
            wb=load_workbook(io.BytesIO(z.read(name)))
            assert not OMITTED & set(wb.sheetnames)
        assert app.session_state.run.terms.grid_days==days


def test_unified_cli_overrides_sensitivity_and_compatibility(tmp_path):
    case=synthetic();inp=tmp_path/'case.json';inp.write_text(json.dumps(case.to_dict()))
    sens=tmp_path/'sensitivity.json';sens.write_text(json.dumps([{'label':'Up','set':{'S0':44.}}]))
    out=tmp_path/'result.zip'
    done=subprocess.run([sys.executable,str(ROOT/'tools/run_case.py'),str(inp),'--out',str(out),
                         '--set','S0=42','--sens',str(sens)],capture_output=True,text=True)
    assert done.returncode==0,done.stderr
    with zipfile.ZipFile(out) as z:
        assert json.loads(z.read('case.json'))['market']['S0']==42
        assert '민감도.csv' in z.namelist()


def test_numeric_comparison_buttons_run_on_request():
    case=synthetic();case.contract.update(p_s=0.,p_e=12.,issuer_call=2,k_w=.3,k_s=3.,k_e=9.,k_prem=.02)
    case.method.update(conv_class='equity',k_method=0)
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=120)
    app.session_state.case=case;app.run()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    next(w for w in app.selectbox if w.label=='분석 도구').set_value('상세 계산·회계 참고표').run()
    next(w for w in app.selectbox if w.label=='상세 분석 항목').set_value('판단·근거').run()
    assert not app.exception
    # 격자를 다시 도는 비교는 누르기 전에는 계산하지 않는다.
    assert all(k not in app.session_state for k in ('jd_call', 'jd_pc', 'jd_bdt'))
    assert any('분리 판정' in m.value for m in app.markdown)
    for title in ['콜 평가방법 비교','우선순위 비교','BDT 적용 검토']:
        next(b for b in app.button if b.label==title).click().run()
        assert not app.exception, app.exception
    assert len(app.session_state.jd_call[1][0]) >= 3
    assert app.session_state.jd_bdt[1][1]['관문']


def test_conversion_response_comparison_runs_on_request():
    """매도청구 통지 뒤 전환 대응 — 값이 갈릴 자리(전환청구기간과 겹치고 의무보유가 없음)에서만 보이고,
    누르기 전에는 계산하지 않는다. 두 읽기의 값이 나란히 나오고 지금 설정에 «적용» 표시가 붙는다."""
    # 주가 150 · 전환가 100 — 전환이 유리한 자리에서 콜이 열려 두 읽기가 크게 갈린다(0 대 약 15)
    case=synthetic();case.contract.update(p_s=0.,p_e=12.,issuer_call=2,k_w=.3,k_s=3.,k_e=9.,k_prem=.02,k_hold=0)
    case.market.update(S0=150.)
    case.method.update(conv_class='equity',k_method=2,k_split=1)
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=120)
    app.session_state.case=case;app.run()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label=='현재 입력으로 평가').click().run()
    next(w for w in app.selectbox if w.label=='분석 도구').set_value('상세 계산·회계 참고표').run()
    next(w for w in app.selectbox if w.label=='상세 분석 항목').set_value('판단·근거').run()
    assert not app.exception
    assert 'jd_cr' not in app.session_state
    next(b for b in app.button if b.label=='전환 대응 비교').click().run()
    assert not app.exception, app.exception
    rows = app.session_state.jd_cr[1]
    assert [r[0] for r in rows] == ['통지 뒤 전환할 수 있다', '통지 뒤 전환할 수 없다']
    assert [r[3] for r in rows] == [True, False]          # 기본 설정은 «전환할 수 있다»
    assert rows[1][1] > rows[0][1]                        # 전환으로 피하지 못하면 콜 값이 크다


def test_explicit_failed_peer_exception_in_ui(monkeypatch):
    import streamlit as st
    st.cache_data.clear()
    def fetch(code,*args,**kwargs):
        if code=='BAD':raise RuntimeError('Synthetic failure')
        return prices(),'Synthetic'
    monkeypatch.setattr(legacy,'fetch_prices',fetch)
    app=AppTest.from_file(str(ROOT/'vol_app.py'),default_timeout=60)
    app.session_state.case=synthetic();app.session_state['_app_embedded']=True;app.run()
    next(w for w in app.text_area if w.label.startswith('피어 목록')).set_value('GOOD,A\nBAD,B')
    next(w for w in app.number_input if w.label=='조회 일수 (거래일)').set_value(30)
    app.run()
    next(b for b in app.button if b.label=='주가 받기').click().run()
    assert app.error and 'vol_snapshot' not in app.session_state
    next(w for w in app.checkbox if w.label=='못 받은 종목 빼고 계산').check().run()
    next(b for b in app.button if b.label=='주가 받기').click().run()
    assert not app.exception and not app.error
    assert len(app.session_state.vol_snapshot['series'])==1
    next(b for b in app.button if b.label=='산출내역 엑셀 생성').click().run()
    assert not app.exception and app.session_state.vol_xlsx


def test_invalid_saved_grid_is_reported_without_silent_coercion():
    case=synthetic();case.method['grid_days']=7.1
    app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=60)
    app.session_state.case=case;app.run()
    assert not app.exception and app.error
    assert app.session_state.case.method['grid_days']==7.1
    next(w for w in app.selectbox if w.label=='계산 간격 단위').set_value(7).run()
    assert not app.exception and app.session_state.case.method['grid_days']==7.


def test_file_missing_prices_do_not_backfill_outside_requested_window():
    from valuation.market_data import select_price_window
    data = 'date,A\n'+'\n'.join(f'2025-01-{i:02d},{100+i if i not in (2,14) else ""}' for i in range(1,16))
    series, omitted = parse_price_table(data, allow_missing=True, return_exclusions=True)
    selected, notes = select_price_window(series, omitted, 12, '2025-01-15')
    assert len(selected[0][1]) == 11
    assert selected[0][1][0][0] == '2025-01-04'
    assert [r['date'] for r in notes] == ['2025-01-14']


def test_future_date_with_all_prices_missing_is_still_rejected():
    from valuation.market_data import select_price_window
    series, notes = parse_price_table('date,A\n2025-01-01,100\n2025-01-02,\n', allow_missing=True, return_exclusions=True)
    with pytest.raises(ValueError, match='기준일 이후'):
        select_price_window(series, notes, 10, '2025-01-01')



def test_sha_tranche_table_screen_to_excel():
    """주주간계약 회차별 표 — 입력 화면 · 평가 · 회차별 결과 · 상세 화면 전 항목 · 세 조서가 오류 없이 돈다.

    회차 둘(고정 가격 · 가산 2% 정기 행사, 풋·콜 수량이 다름). 가상 수치다."""
    rows = [dict(name='1차', start='2026-01-01', end='2026-12-31', style='any', price=900., rate=0.,
                 put_q=30000., call_q=30000., rf=.021, pdisc=.061),
            dict(name='2차', start='2027-01-01', end='2027-12-31', style='periodic', freq=3., price=1600.,
                 rate=.02, put_q=30000., call_q=20000.)]
    case = Case(name='회차 시험', contract=dict(inst='SHA', d_issue='2021-11-15', d_mat='2027-12-31', K0=900.,
                face_total=900.*30000+1600.*30000, rfx_mode=0, cpn=0., cv_s=99., cv_e=0., p_s=99., p_e=0., k_w=0.,
                sha_put_s=99., sha_put_e=0., sha_call_s=99., sha_call_e=0., sha_rows=rows),
                market=dict(S0=1000., sig=.6, rf_curve=[[1, .0226], [3, .024]], cr_curve=[[1, .05], [3, .055]]),
                method=dict(d_base='2025-09-30', model='TF', view='issuer', gap_m=1., grid_days=14.))
    app = AppTest.from_file(str(ROOT/'app.py'), default_timeout=300)
    app.session_state.case = case; app.run()
    assert not app.exception and not app.error
    assert len(app.session_state.case.contract['sha_rows']) == 2
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    next(b for b in app.button if b.label == '현재 입력으로 평가').click().run()
    assert not app.exception
    run = app.session_state.run
    ref = calculate(app.session_state.case)
    assert run.summary['amounts_total'] == ref.summary['amounts_total']
    assert [r['회차'] for r in run.summary['sha_rows']] == ['1차', '2차']
    assert abs(sum(r['풋 전액'] for r in run.summary['sha_rows']) - run.summary['amounts_total']['put']) < 1e-6
    assert any('순액' in m.label for m in app.metric)
    next(w for w in app.selectbox if w.label == '분석 도구').set_value('상세 계산·회계 참고표').run()
    for section in next(w for w in app.selectbox if w.label == '상세 분석 항목').options:
        next(w for w in app.selectbox if w.label == '상세 분석 항목').set_value(section).run()
        assert not app.exception, section
    app.radio(key='_workflow_stage').set_value('조서 출력').run()
    for option, want in [('기본 값 조서', {'회차별 결과', '회차별 입력'}),
                         ('상세 계산 값 조서', {'1·결과', '2·결과', '회차 합계', '회계처리'}),
                         ('상세 계산 수식 조서', {'1·결과', '2·결과', '회차 합계', '회계처리'})]:
        next(w for w in app.radio if w.label == '조서 구성').set_value(option).run()
        next(w for w in app.checkbox if w.label.startswith('회계처리')).check().run()
        next(b for b in app.button if b.label == '조서 생성').click().run()
        assert not app.exception and not app.error
        with zipfile.ZipFile(io.BytesIO(app.session_state.bundle)) as z:
            name = next(n for n in z.namelist() if n.endswith('.xlsx'))
            wb = load_workbook(io.BytesIO(z.read(name)))
            assert want <= set(wb.sheetnames), (option, wb.sheetnames)
