"""Business scenarios: evidence reuse, stale outputs, units and optional work."""
import copy
import datetime as dt
import io
import json
import zipfile
from dataclasses import asdict

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from test_v2_workflow import synthetic, ROOT
from valuation import legacy
from valuation.case import Case, FIELDS, import_legacy, inspect_case
from valuation.presentation import LABELS, event_months
from valuation.service import calculate, refresh_run, calculation_key, export_bundle, CaseError
from valuation.analysis import sensitivity


def forbidden(*args, **kwargs):
    raise AssertionError('Unexpected extra pricing/legacy judgement call')


def button(app, text):
    return next(b for b in app.button if b.label == text)


def test_metadata_edit_updates_evidence_without_pricing(monkeypatch):
    case = synthetic()
    run = calculate(case)
    for name in ['engine', 'decompose', 'sha_engine', 'validate', 'sha_validate', 'eir_or_none']:
        monkeypatch.setattr(legacy, name, forbidden)
    edited = copy.deepcopy(case)
    edited.name = 'Reviewed case'
    edited.sources['S0'] = 'Reviewer evidence, page 17'
    edited.notes = 'Updated after review'
    edited.assumptions = [dict(field='S0', value=case.market['S0'], rationale='Same input, documented reason')]
    updated = refresh_run(run, edited)
    assert updated.raw is run.raw
    assert updated.summary['amounts_total'] == run.summary['amounts_total']
    assert updated.summary['case_sha256'] != run.summary['case_sha256']
    assert updated.summary['calculated_at'] == run.summary['calculated_at']
    assert run.case.sources.get('S0') != edited.sources['S0']
    with zipfile.ZipFile(io.BytesIO(export_bundle(updated))) as z:
        assert json.loads(z.read('case.json'))['notes'] == edited.notes
        wb = load_workbook(io.BytesIO(z.read('value_review.xlsx')))
        assert any(c.value == edited.sources['S0'] for row in wb['출처기록'] for c in row)


@pytest.mark.parametrize('section,key,value', [
    ('market', 'S0', 60.), ('market', 'sig', .3), ('contract', 'issue_px', 120.),
    ('contract', 'face_total', 1000.), ('method', 'gap_m', 1.),
    ('method', 'd_base', '2025-03-31'), ('contract', 'p_s', 0.),
])
def test_changed_calculation_inputs_reject_reuse(section, key, value):
    case = synthetic(); run = calculate(case)
    getattr(case, section)[key] = value
    assert calculation_key(case) != run.summary['calculation_key']
    with pytest.raises(ValueError, match='다시 평가'):
        refresh_run(run, case)


def test_all_numerical_input_fields_participate_in_reuse_key():
    from valuation.service import ACCOUNTING_ONLY
    case = synthetic(); initial = calculation_key(case)
    for key, value in case.facts().items():
        if isinstance(value, (float, int)) and key not in ACCOUNTING_ONLY:
            changed = copy.deepcopy(case)
            from valuation.case import section_for
            getattr(changed, section_for(key))[key] = value + 1
            assert calculation_key(changed) != initial, key


@pytest.mark.parametrize('offset', [0., 1., 12., 12.123456, 59.9999])
def test_date_display_preserves_original_fractional_months(offset):
    issue = '2024-02-29'
    visible = legacy.months_to_date(issue, offset)
    assert event_months(issue, visible, offset, issue) == offset
    later = visible + dt.timedelta(days=1)
    assert legacy.months_to_date(issue, event_months(issue, later, offset, issue)) == later


def test_all_inputs_have_human_labels():
    assert not FIELDS - set(LABELS)


def test_input_warnings_are_short_and_computed_once(monkeypatch):
    run = calculate(synthetic())
    assert not any(i.code == 'engine_review' for i in run.issues)
    assert not any(i.code == 'judgement_scope' for i in run.issues)
    # 옛 입력 경고는 첫 문장만 '확인 내용'으로 싣고, 증빙 갱신 때는 다시 만들지 않는다.
    assert all('**' not in i.message and len(i.message) < 240 for i in run.issues if i.code == 'input_check')
    monkeypatch.setattr(legacy, 'validate', forbidden)
    again = refresh_run(run, copy.deepcopy(run.case))
    assert [i for i in again.issues if i.code == 'input_check'] == [i for i in run.issues if i.code == 'input_check']


def test_basic_workpaper_ties_to_result_without_extra_calculation(monkeypatch):
    run = calculate(synthetic())
    monkeypatch.setattr(legacy, 'engine', forbidden)
    monkeypatch.setattr(legacy, 'eir_or_none', forbidden)
    monkeypatch.setattr(legacy, 'build_xlsx', forbidden)
    with zipfile.ZipFile(io.BytesIO(export_bundle(run))) as z:
        wb = load_workbook(io.BytesIO(z.read('value_review.xlsx')))
        row = next(row for row in wb['평가요약'].values if row[0] == '순포지션 가치')
        assert row[1] == run.summary['amounts_total']['net']
        assert row[2] == run.summary['amounts_per_share']['net']
        assert row[3] == run.summary['amounts_100']['net']
        assert not {'분리 판단', '회계처리', '99_모형검증'} & set(wb.sheetnames)
        assert not any(c.data_type == 'f' for ws in wb for row in ws for c in row)


def test_sensitivity_prices_only_two_requested_variants(monkeypatch):
    run = calculate(synthetic())
    calls = []
    original = legacy.decompose
    def spy(tm):
        calls.append(tm.S0)
        return original(tm)
    monkeypatch.setattr(legacy, 'decompose', spy)
    analysis = sensitivity(run, 'S0', 10.)
    assert calls == [36., 44.]
    assert [r['net'] for r in analysis['rows']] == pytest.approx([9e9, 1e10, 11e9])
    assert run.terms.S0 == 40.


def test_rollforward_retains_old_market_date_review():
    case = synthetic(); case.sources['market_date'] = case.method['d_base']
    case.method['d_base'] = '2025-03-31'
    run = calculate(case)
    assert any(i.code == 'market_date' and '2025-01-01' in i.message for i in run.issues)


@pytest.mark.parametrize('schedule', ['invalid row', '2025-04-01 abc'])
def test_malformed_payment_schedule_blocks_calculation(schedule):
    case = synthetic(); case.contract['p_sched'] = schedule
    with pytest.raises(CaseError):
        calculate(case)


@pytest.mark.parametrize('inst', ['CB', 'BW', 'RCPS', 'SHA'])
def test_instrument_screens_preserve_values_and_only_render_relevant_rights(inst):
    case = synthetic(); case.contract['inst'] = inst
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=30)
    app.session_state['case'] = case
    app.run()
    assert not app.exception
    assert next(r for r in app.radio if r.label == '평가 진행').options == ['입력·시장자료', '평가·분석', '조서 출력']
    assert not any('JSON' in w.label for w in app.text_area)
    assert not any(w.label == '상환청구 주기(개월)' for w in app.number_input)
    assert app.session_state['case'].to_dict() == case.to_dict()
    if inst == 'SHA':
        assert not any(w.label == '전환·신주인수권 행사 가능' for w in app.checkbox)
    # Read-only viewing must neither run the engine nor change the saved input.
    assert 'run' not in app.session_state


def test_ui_rerender_evidence_edit_and_basic_export_have_zero_extra_pricing(monkeypatch):
    case = synthetic()
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=30)
    app.session_state['case'] = case
    app.run()
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    button(app, '현재 입력으로 평가').click().run()
    assert not app.exception
    old = app.session_state['run']
    for name in ['engine', 'decompose', 'sha_engine', 'validate', 'eir_or_none']:
        monkeypatch.setattr(legacy, name, forbidden)
    app.run()
    app.radio(key='_workflow_stage').set_value('입력·시장자료').run()
    app.radio(key='_input_area').set_value('출처·평가가정').run()
    next(w for w in app.text_area if w.label == '검토메모').set_value('Reviewer update')
    button(app, '출처·메모 저장').click().run()
    assert not app.exception
    assert app.session_state['run'].case.notes == 'Reviewer update'
    assert app.session_state['run'].raw is old.raw
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    button(app, '현재 입력으로 평가').click().run()
    app.radio(key='_workflow_stage').set_value('조서 출력').run()
    button(app, '조서 생성').click().run()
    assert not app.exception
    assert 'bundle' in app.session_state


def test_percent_edit_uses_percent_units_and_blocks_old_export():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=30)
    app.session_state['case'] = synthetic()
    app.run(); app.radio(key='_workflow_stage').set_value('평가·분석').run(); button(app, '현재 입력으로 평가').click().run()
    app.radio(key='_workflow_stage').set_value('입력·시장자료').run()
    widget = next(w for w in app.number_input if w.label == '주가 변동성(연, %)')
    assert widget.value == 25.
    widget.set_value(30.).run()
    assert app.session_state['case'].market['sig'] == .3
    app.radio(key='_workflow_stage').set_value('조서 출력').run()
    assert button(app, '조서 생성').disabled
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    assert any('변경 전 입력' in w.value for w in app.warning)


@pytest.mark.parametrize('inst', ['CB', 'BW', 'RCPS', 'SHA'])
@pytest.mark.parametrize('formula', [False, True])
def test_detailed_exports_judgment_sheets_are_optional(inst, formula):
    case = synthetic(); case.contract['inst'] = inst
    run = calculate(case)
    name = 'formula_review.xlsx' if formula else 'value_review.xlsx'
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, detail=True, formula=formula, judgment=False))) as z:
        wb = load_workbook(io.BytesIO(z.read(name)))
        assert not {'해설', '분리 판단', '검산요약', '99_모형검증', '판단·근거', '회계처리', '상각표'} & set(wb.sheetnames)
        assert '결과' in wb.sheetnames
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, detail=True, formula=formula))) as z:
        wb = load_workbook(io.BytesIO(z.read(name)))
        assert '판단·근거' in wb.sheetnames
        if inst != 'SHA':
            assert {'분리 판단', '검산요약', '99_모형검증'} <= set(wb.sheetnames)


def test_reversed_exercise_dates_cannot_silently_remove_a_right():
    app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=30)
    app.session_state['case'] = synthetic()
    app.run()
    next(w for w in app.date_input if w.label == '전환·신주인수권 행사 시작일').set_value(dt.date(2026, 6, 1)).run()
    assert any('시작일이 종료일보다' in w.value for w in app.error)
    app.radio(key='_workflow_stage').set_value('평가·분석').run()
    assert button(app, '현재 입력으로 평가').disabled
