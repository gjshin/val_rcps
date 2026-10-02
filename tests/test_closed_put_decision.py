"""열리지 않은 상환청구는 결정 후보가 아니다 — 가상 수치.

만기 자동전환 우선주에서 상환청구권이 없을 때, 주가가 0 에 가까운 자리(전환가치·보유가치가 허용오차
1e-9 안)에서 금액 0 짜리 «상환청구» 가 «풋은 동점이면 이긴다» 규칙으로 골라져 «행사 불가능한 자리의
결정» 점검에 걸리고 조서가 거부되던 문제. 매도청구가 열리지 않으면 무한대로 넘기듯 상환청구도 −무한대로
넘긴다. 금액은 그 자리 값이 1e-9 수준이라 바뀌지 않는다.
"""
import io
import math
import shutil
import zipfile
from collections import Counter
from dataclasses import asdict

import pytest
from openpyxl import load_workbook

from valuation import legacy
from valuation.case import import_legacy
from valuation.service import calculate, export_bundle

LAB = {"conv": "전환", "auto": "자동전환", "ipo": "상장전환", "put": "상환P", "call": "상환C",
       "hold": "보유", "mat": "만기상환"}


def case(sig=.9, years=10, **o):
    t = legacy.Terms(inst="RCPS", d_issue="2026-01-01", d_base="2026-01-01", d_mat=f"{2026+years}-01-01",
                     S0=20000., K0=80000., issue_px=80000., face_total=8e9, sig=sig, par=500.,
                     cpn=.03, div_basis=1, div_mode=0, ipay=12., mat_mode=0, ytm=.05, ytm_cmp=1,
                     p_mode="accrue", p_yield=.05, p_cmp=1, rfx_mode=0, issuer_call=0, k_w=0.,
                     p_s=12.*years, p_e=0., cv_s=1., cv_e=12.*years, view="issuer", gap_m=1.,
                     rf_curve=[[1, .025], [5, .03], [30, .035]], cr_curve=[[1, .07], [5, .08], [30, .09]])
    for k, v in o.items(): setattr(t, k, v)
    return import_legacy(asdict(t), "가상 자동전환 우선주")


def test_no_decision_on_closed_put_and_workpaper_builds():
    run = calculate(case())
    assert run.raw["integrity"] == []
    bad = [k for k, o in run.raw["full"]["memo"].items() if o["kind"] == "put" and o.get("pv", 0.0) <= 0]
    assert bad == []
    with zipfile.ZipFile(io.BytesIO(export_bundle(run, accounting=True))) as z:
        assert "value_review.xlsx" in z.namelist()


def test_value_unchanged_by_decision_fix():
    # 그 자리의 값은 1e-9 수준이다 — 금액은 소수 10자리까지 같다 (고치기 전 25.116012547111072).
    run = calculate(case())
    assert run.summary["amounts_100"]["net"] == pytest.approx(25.116012547111072, abs=1e-10)


def test_open_put_still_wins_ties():
    # 상환청구가 열려 있는 자리의 동점 규칙(풋은 동점이면 이긴다)은 그대로다.
    assert legacy.node_decide(5.0, 5.0, math.inf, 5.0, False) == "put"
    assert legacy.node_decide(1e-10, -math.inf, math.inf, 1e-10, False) == "hold"


@pytest.mark.skipif(not (shutil.which("libreoffice") or shutil.which("soffice")), reason="LibreOffice 없음")
def test_formula_workbook_decides_same_as_engine(tmp_path):
    # 수식 조서를 다시 계산해 «09 의사결정» 시트의 결정 표시가 엔진과 계산 시점마다 같은지 본다.
    import subprocess
    run = calculate(case())
    data = zipfile.ZipFile(io.BytesIO(export_bundle(run, formula=True, detail=True))).read("formula_review.xlsx")
    (tmp_path/"s.xlsx").write_bytes(data); (tmp_path/"o").mkdir()
    subprocess.run([shutil.which("libreoffice") or shutil.which("soffice"),
                    "-env:UserInstallation=" + (tmp_path/"p").as_uri(), "--headless", "--convert-to", "xlsx",
                    "--outdir", str(tmp_path/"o"), str(tmp_path/"s.xlsx")], capture_output=True, timeout=1800)
    ws = load_workbook(tmp_path/"o"/"s.xlsx", data_only=True)["09 의사결정"]
    n = run.terms.n
    eng = Counter(); xl = Counter()
    for (i, j), o in run.raw["full"]["memo"].items():
        eng[(i, LAB[o["kind"]])] += 1
    names = set(LAB.values())
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value in names and c.column >= 3:
                xl[(c.column - 3, c.value)] += 1
    assert {k: v for k, v in xl.items() if k[0] <= n} == eng
