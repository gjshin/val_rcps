"""Reusable multi-tranche work and an explicitly ordered value-change bridge."""
import copy
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass, field
from .case import Case, section_for
from .service import calculate, calculation_key, refresh_run, export_bundle

SCHEMA = 'valuation-engagement/1'


@dataclass
class Engagement:
    name: str
    entity: str
    cases: dict[str, Case] = field(default_factory=dict)

    def to_dict(self):
        return dict(schema=SCHEMA, name=self.name, entity=self.entity,
                    cases=[dict(id=k, case=v.to_dict()) for k,v in self.cases.items()])

    @classmethod
    def from_dict(cls, data):
        if not isinstance(data, dict) or set(data) != {'schema', 'name', 'entity', 'cases'} or data['schema'] != SCHEMA:
            raise ValueError('용역 묶음 파일 형식을 확인하십시오.')
        if not all(isinstance(data[k], str) and data[k].strip() for k in ['name','entity']) or not isinstance(data['cases'], list):
            raise ValueError('용역명·평가대상회사·회차 목록이 필요합니다.')
        result = cls(data['name'], data['entity'])
        for row in data['cases']:
            if not isinstance(row, dict) or set(row) != {'id','case'} or not isinstance(row['id'], str) or not row['id'].strip() or row['id'] in result.cases:
                raise ValueError('회차 식별값이 비어 있거나 중복되었습니다.')
            result.cases[row['id']] = Case.from_dict(row['case'])
        return result

    def fingerprint(self):
        return hashlib.sha256(json.dumps(self.to_dict(), sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def calculate_engagement(engagement, previous_runs=None):
    if not engagement.cases:
        raise ValueError('평가할 회차를 추가하십시오.')
    dates = {c.effective().get('d_base') for c in engagement.cases.values()}
    views = {c.effective().get('view') for c in engagement.cases.values() if c.effective().get('inst') != 'SHA'}
    if len(dates) != 1 or len(views) > 1:
        raise ValueError('합산할 회차의 평가기준일과 발행자·투자자 관점을 맞추십시오.')
    runs, rows = {}, []
    for key, case in engagement.cases.items():
        old = (previous_runs or {}).get(key)
        run = refresh_run(old, case) if old and old.summary['calculation_key'] == calculation_key(case) else calculate(case)
        runs[key] = run
        values = run.summary['amounts_total']
        sha = run.terms.inst == 'SHA'
        rows.append(dict(회차식별값=key, 건명=case.name, 상품=run.terms.inst, 기준일=run.terms.d_base,
                         평가대상금액_원=run.terms.face_total, 본체_원=None if sha else values['whole_before_call'],
                         콜차감_원=None if sha else values['call_deduction'], 순포지션_원=None if sha else values['net'],
                         별도풋_원=values['put'] if sha else None, 별도콜_원=values['call'] if sha else None,
                         검토필요=len(run.issues), 입력식별값=case.fingerprint()))
    totals = {key: sum(r[key] or 0 for r in rows) for key in ['본체_원','콜차감_원','순포지션_원','별도풋_원','별도콜_원']}
    return dict(key=engagement.fingerprint(), runs=runs, rows=rows, totals=totals, currency='KRW',
                scope='회차별 독립 평가의 산술 합계입니다. 회차 간 우선순위·상호 희석·교차 조건을 공동 평가하지 않습니다. 주주간계약 풋·콜은 별도 총액이며 순포지션에 합산하지 않습니다.')


def engagement_bundle(engagement, result):
    if result['key'] != engagement.fingerprint():
        raise ValueError('용역 입력이 변경되었습니다. 다시 평가하십시오.')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = '회차별총괄'
    ws.append([engagement.name, engagement.entity, 'KRW']); ws.append([result['scope']]); ws.append([])
    ws.append(list(result['rows'][0]))
    for row in result['rows']:
        ws.append(list(row.values()))
    ws.append([])
    for k,v in result['totals'].items(): ws.append([k, v])
    for row in ws:
        for c in row:
            if isinstance(c.value, str): c.data_type = 's'
            elif isinstance(c.value, (int, float)): c.number_format = '#,##0.00'
    for c in ws[4]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='17365D')
    ws.freeze_panes = 'A5'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 25
    out = io.BytesIO(); wb.save(out)
    files = {'engagement.json': json.dumps(engagement.to_dict(), ensure_ascii=False, indent=2).encode(), '회차별총괄.xlsx': out.getvalue(),
             'aggregate.json': json.dumps({k:v for k,v in result.items() if k!='runs'}, ensure_ascii=False, indent=2).encode()}
    for i,(key,run) in enumerate(result['runs'].items(), 1):
        # User labels never become filesystem/ZIP paths.
        files[f'tranche_{i:03d}.zip'] = export_bundle(run)
    files['manifest.json'] = json.dumps({k: hashlib.sha256(v).hexdigest() for k,v in files.items()}, indent=2).encode()
    zipped = io.BytesIO()
    with zipfile.ZipFile(zipped, 'w', zipfile.ZIP_DEFLATED) as z:
        for k,v in files.items(): z.writestr(k,v)
    return zipped.getvalue()


def value_bridge(previous, current):
    """Sequential replacement. No Shapley/causal attribution is claimed."""
    if previous.effective().get('inst') != current.effective().get('inst') or previous.effective().get('view') != current.effective().get('view'):
        raise ValueError('같은 상품·관점의 전기와 당기를 비교하십시오.')
    if previous.effective().get('inst') == 'SHA':
        raise ValueError('주주간계약은 풋·콜을 따로 대사해야 하므로 이 순포지션 변동 분석에서 제외합니다.')
    start = calculate(previous)
    working = copy.deepcopy(previous)
    before, target = previous.effective(), current.effective()
    groups = [('기준일 경과', {'d_base'}),
              ('계약·물량 변경', set(current.contract)|set(previous.contract)),
              ('주당가치 변경', {'S0'}), ('변동성 변경', {'sig'}),
              ('금리곡선·기타 시장자료 변경', set(current.market)|set(previous.market)-{'S0','sig'}),
              ('평가방법·나머지 입력 변경', set(before)|set(target))]
    rows = [dict(단계='전기', 평가액_원=start.summary['amounts_total']['net'], 변동액_원=0., 변경항목=[])]
    seen = set()
    applied = dict(before)
    for title, keys in groups:
        changes = sorted(k for k in keys-seen if applied.get(k) != target.get(k))
        seen |= keys
        if not changes: continue
        for key in changes:
            if key in target: applied[key] = target[key]
            else: applied.pop(key, None)
        working.assumptions = []
        for group in ['contract','market','method']:
            setattr(working, group, {k: v for k,v in applied.items() if section_for(k) == group})
        try:
            run = calculate(working)
        except ValueError as exc:
            raise ValueError(f'{title} 단계의 중간 조건이 유효하지 않아 분해할 수 없습니다. 계약 변경과 시점 변경을 함께 검토하십시오. {exc}') from exc
        value = run.summary['amounts_total']['net']
        rows.append(dict(단계=title, 평가액_원=value, 변동액_원=value-rows[-1]['평가액_원'], 변경항목=changes))
    end = calculate(current)
    difference = end.summary['amounts_total']['net']-rows[-1]['평가액_원']
    tolerance = max(1e-6, abs(end.summary['amounts_total']['net'])*1e-10)
    if abs(difference) > tolerance:
        raise ArithmeticError('변동 원인 합계와 당기 평가액이 일치하지 않습니다.')
    return dict(rows=rows, opening=start.summary['amounts_total']['net'], closing=end.summary['amounts_total']['net'],
                reconciliation_difference=difference, previous_key=calculation_key(previous), current_key=calculation_key(current),
                scope='기준일→계약→주가→변동성→금리·시장→방법 순서의 대체 분석. 원인별 금액은 대체 순서에 영향을 받으며 독립적 인과효과 또는 회계상 평가손익을 뜻하지 않습니다. 중도 현금수취·상환·신규취득은 별도 대사하십시오.')
