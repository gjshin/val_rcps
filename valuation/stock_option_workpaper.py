"""주식선택권 상세 조서 — 전환사채·상환전환우선주 상세 조서(legacy.build_xlsx)와 같은 서식으로 만든다.

시트: 해설 → 가정 → 격자(회차마다 한 장) → 결과 → 비용 배분 → 회계처리 → 금리곡선.
격자 시트는 전환사채 조서처럼 1~17행에 시점별 공통조건을 두고, 그 아래 ① 주가 · ② 옵션 가치
(· ③ 목표 주가 미도달 상태) 트리를 싣는다. 행 번호 r 은 하락 횟수다.

값 조서는 평가 당시의 값을, 수식 조서는 같은 시트·같은 칸에 엔진과 같은 순서의 수식을 넣는다.
수식 조서를 다시 계산하면 값 조서와 같은 값이 나와야 한다 (tests/test_stock_option.py 가 리브레오피스로 대조).
이미 평가한 실행 결과(run)로만 만든다 — 새로 계산하지 않는다.
"""
import datetime as dt
import math

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter as gl

from . import legacy as L
from . import stock_option as SO
from . import stock_option_report as SOR

# 전환사채 상세 조서와 같은 글꼴·색·숫자 형식 (legacy.build_xlsx)
FONT = "KoPub돋움체 Medium"
NAVY, SUB, LIGHT, BAND, RFXC = "1F3864", "44618C", "DCE6F1", "F2F5F8", "FCE4D6"
RED, GREEN, GREY, AMB = "C00000", "006100", "6B7480", "BF8F00"
N2, N0, N4, N6, P2, P4 = '#,##0.00', '#,##0', '0.0000', '0.000000', '0.00%', '0.0000%'
DATE = 'yyyy-mm-dd'
# 입력 기록은 넣은 자릿수를 그대로 보인다 (2.125% 를 2.13% 로 적지 않는다)
IN_NUM, IN_PCT = '#,##0.00####', '0.00####%'
_THIN = Side(style="thin", color="BFC7D0")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

ASSUME = '가정'
RESULT = '결과'
SCHED = '비용 배분'
JOURNAL = '회계처리'

# 격자 시트 1~17행 — 시점별 공통조건 (전환사채 조서의 1~17행과 같은 자리)
HEAD = ('날짜', '스텝(노드 번호)', '행사 가능 (1=예)', '가득 전 (1=예)', '시장조건 판정 (1=예)', '행사가격',
        '시장조건 목표 주가', '무위험 선도이자율', '구간 길이 Δt (년)', '주가변동성 σ', '상승계수 u', '하락계수 d',
        '위험중립 상승확률 q', '하락확률 1−q', '할인계수 exp(−r × Δt)', '구간 퇴사 확률', '구간 비가득조건 중단 확률')
HR = {h: i + 1 for i, h in enumerate(HEAD)}


def layout(n):
    """격자 시트의 자리 — 값 조서와 수식 조서가 같은 자리를 쓴다. 트리마다 제목 줄 · 스텝 머리 줄 · 노드 줄."""
    stock = 20                                 # 18행 제목, 19행 스텝 — 전환사채 조서의 R0 와 같다
    value = stock + n + 4
    wait = value + n + 4
    return dict(head=HR, stock=stock, value=value, wait=wait)


class _Writer:
    """legacy.build_xlsx 의 put · title · sec 과 같은 모양. 같은 모양의 서식은 한 번만 만든다."""

    def __init__(self, formula):
        self.formula = formula
        self._st = {}
        self.inputs = set()                    # 수식 조서에서 고칠 수 있는 칸 {(시트, 칸)}

    def put(self, ws, r, c, v, *, bold=False, color="000000", fill=None, fmt=None, size=10, align=None,
            border=False, wrap=False):
        cl = ws.cell(row=r, column=c, value=v)
        if isinstance(v, str):
            cl.data_type = 's'                 # 입력·출처 문자열은 엑셀 수식으로 읽히지 않는다
        self._style(cl, (bold, color, fill, fmt, size, align, border, wrap))
        return cl

    def fx(self, ws, r, c, formula, **kw):
        """칸의 값을 수식으로 바꾼다 — 서식은 값 조서와 같다."""
        cl = ws.cell(row=r, column=c)
        cl.value = formula
        if kw:
            self._style(cl, tuple(kw.get(k, d) for k, d in (('bold', False), ('color', '000000'), ('fill', None),
                                                              ('fmt', None), ('size', 10), ('align', None),
                                                              ('border', False), ('wrap', False))))
        return cl

    def _style(self, cl, key):
        if key in self._st:
            cl._style = self._st[key]          # 같은 모양은 한 객체를 함께 쓴다 (legacy.build_xlsx 와 같다 — 마무리에서 칸마다 떼어 낸다)
            return
        bold, color, fill, fmt, size, align, border, wrap = key
        cl._style = StyleArray()               # 다른 칸과 함께 쓰는 모양을 고치지 않게 새로 만든다
        cl.font = Font(name=FONT, size=size, bold=bold, color=color)
        if fill:
            cl.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            cl.number_format = fmt
        cl.alignment = Alignment(horizontal=align or "general", vertical="center", wrap_text=wrap)
        if border:
            cl.border = BOX
        self._st[key] = cl._style

    def title(self, ws, r, t, span=8):
        self.put(ws, r, 2, t, bold=True, color="FFFFFF", fill=NAVY, size=13)
        for c in range(3, 2 + span):
            self.put(ws, r, c, None, fill=NAVY)
        ws.row_dimensions[r].height = 22

    def sec(self, ws, r, t, span=8):
        self.put(ws, r, 2, t, bold=True, color="FFFFFF", fill=SUB, size=10)
        for c in range(3, 2 + span):
            self.put(ws, r, c, None, fill=SUB)

    def note(self, ws, r, t, color=GREY):
        self.put(ws, r, 2, t, color=color, size=9)

    def head(self, ws, r, c0, heads):
        for k, h in enumerate(heads):
            self.put(ws, r, c0 + k, h, bold=True, fill=LIGHT, align="center", border=True, size=9, wrap=True)

    def value_in(self, ws, r, c, v, fmt):
        """입력값 — 값 조서는 빨강 글씨, 수식 조서는 노란 칸 (파란 글씨로 열린다)."""
        if self.formula:
            self.put(ws, r, c, v, fill=L.INPUT_FILL, fmt=fmt, align="right", border=True)
            self.inputs.add((ws.title, f"{gl(c)}{r}"))
        else:
            self.put(ws, r, c, v, color=RED, fmt=fmt, align="right", border=True)


def _q(name):
    return "'" + str(name).replace("'", "''") + "'"


def _abs(sheet, col, row):
    return f"{_q(sheet)}!${col}${row}"


# ── 해설 ─────────────────────────────────────────────────────────────────

def guide_sheet(w, wb, run):
    t, md = run.terms, run.raw['mode']
    mkt = int(getattr(t, 'so_mkt', 0))
    H = wb.create_sheet("해설")
    H.sheet_view.showGridLines = False
    H.column_dimensions["B"].width = 24
    H.column_dimensions["C"].width = 100
    w.title(H, 2, "이 조서를 읽는 법", span=2)
    ex_text = {0: "만기에만 행사 (유럽형) — 행사 종료일 전에는 H 를 그대로 씁니다.",
               1: "가득 뒤 최적 행사 — 큰 값(주가 − 행사가격, H).",
               2: f"행사 배수에 도달하면 행사 — 주가 ≥ {float(t.so_mult):g} × 행사가격이고 내가격이면 주가 − 행사가격, 아니면 H."
               }[int(t.so_ex)]
    mkt_text = {0: "이 평가에는 시장조건이 없습니다.",
                1: "가득일 판정 — 가득 시점 노드에서 주가가 목표 미만이면 0 입니다. 목표가 걸친 노드는 그 노드가 대표하는 로그 주가 "
                   "구간(위아래 반 칸씩) 가운데 목표 이상인 비율 w = 최솟값(1, 큰 값(0, (ln S + ln u − ln 목표) ÷ (2 × ln u))) "
                   "만큼 남깁니다.",
                2: "한 번이라도 판정 — ③ 목표 주가 미도달 상태 가치트리: 목표에 닿은 노드는 ② 옵션 가치, 가득 시점까지 못 닿으면 0, "
                   "그 사이는 다음 시점의 미도달 상태 가치를 할인합니다. 격자 노드(계산 간격)마다 관찰하며, 시점 0의 값이 "
                   "1주당 공정가치입니다."}[mkt]
    flow = "가정 → 격자 (회차마다 한 장) → 결과 → 비용 배분 → 회계처리 → 금리곡선"
    kind = {'grant': "부여일 평가입니다.",
            'mod': "조건 변경 평가입니다 — 회차마다 같은 날·같은 시장자료로 잰 변경 후·변경 전 격자가 따로 있습니다.",
            'cash': "현금결제형 결산일 재측정입니다 — 결산일의 시장자료로 다시 잰 격자입니다."}[md]
    ex = [("시트 순서", ""),
          ("구조", flow + ". " + kind),
          ("따라가기", "시트 탭을 왼쪽부터 차례로 보면 계산이 쌓이는 순서와 같습니다."),
          ("", ""),
          ("격자 시트의 공통조건 1~17행", ""),
          ("1행 날짜", "측정기준일 + 스텝 × 구간 길이(Δt). 가득일·행사 시작일과 대조하는 자리입니다."),
          ("2행 스텝", "측정기준일부터 센 노드 번호입니다."),
          ("3~5행 표시", "행사 가능(행사 시작일 이후 첫 노드부터) · 가득 전(비가득조건을 곱하는 구간) · 시장조건 판정 시점에 1 을 "
                       "표시합니다."),
          ("6~7행 금액", "행사가격 · 시장조건 목표 주가 (원)."),
          ("8~9행 금리·기간", "무위험 선도이자율(한 구간, 연속복리)과 구간 길이 Δt(년)."),
          ("10~15행 격자 변수", "주가변동성 σ · 상승계수 u = exp(σ × √Δt) · 하락계수 d = 1 ÷ u · 위험중립 상승확률 "
                           "q = (exp((r − 배당수익률) × Δt) − d) ÷ (u − d) · 하락확률 1 − q · 할인계수 exp(−r × Δt)."),
          ("16~17행 확률", "구간 퇴사 확률 = 1 − (1 − 연 퇴사율)^Δt (행사 시작 뒤) · 구간 비가득조건 중단 확률 = "
                        "1 − (1 − 연 중단율)^Δt (가득 전)."),
          ("", ""),
          ("행 번호 r 은 하락 횟수", ""),
          ("r = 0", "한 번도 하락하지 않은 경로 — 맨 위, 주가가 가장 높습니다."),
          ("r = 스텝", "매번 하락한 경로 — 맨 아래입니다."),
          ("빈칸", "그 시점에는 존재하지 않는 노드입니다."),
          ("", ""),
          ("계산 순서", ""),
          ("행사 종료일부터", "행사 종료일에는 내가격이면 행사, 아니면 소멸합니다: 큰 값(주가 − 행사가격, 0)."),
          ("한 시점씩 앞으로", "계속 보유 C = 할인계수 × [q × 위 노드 가치 + (1 − q) × 아래 노드 가치]."),
          ("가득 전", "비가득조건이 있으면 C × (1 − 구간 중단 확률). 행사 시작 전에는 행사할 수 없어 이 값이 노드 가치입니다."),
          ("행사 시작 뒤", "퇴사 반영 H = (1 − 구간 퇴사 확률) × C + 구간 퇴사 확률 × 큰 값(주가 − 행사가격, 0). "
                        "조기행사 가정: " + ex_text),
          ("시장조건", mkt_text),
          ("순환참조가 없는 이유", "각 칸은 같은 시점의 값과 다음 시점의 계산 결과만 참조하므로 순환참조가 발생하지 않습니다."),
          ("", ""),
          ("근무조건과 비시장 성과조건", ""),
          ("수량에 반영", "근무조건·비시장 성과조건은 1주당 공정가치가 아니라 가득 예상 비율(수량)에 반영합니다 (문단 19·20). "
                       "보상원가 = 1주당 공정가치 × 부여 수량 × 가득 예상 비율."),
          ("", ""),
          ("이 파일의 성격", ""),
          ("값 조서", "앱이 계산한 결과를 값으로 담았습니다. 수식이 아니므로 칸을 바꿔도 다시 계산되지 않습니다."),
          ("수식 조서", "같은 시트·같은 칸에 엔진과 같은 순서의 수식이 들어 있습니다. 고치지 않고 열면 값 조서와 같은 값입니다. "
                     "가정 시트의 노란 칸만 고칠 수 있고, 날짜·행사가격·수량·구간 수·금리는 앱에서 바꾼 뒤 조서를 다시 만듭니다.")]
    r = 4
    for a, b in ex:
        if a and not b:
            w.sec(H, r, a, span=2)
        elif a:
            w.put(H, r, 2, a, bold=True, size=9)
            w.put(H, r, 3, b, size=9, wrap=True)
        r += 1
    return H


# ── 가정 ─────────────────────────────────────────────────────────────────

def assumption_sheet(w, wb, run):
    """평가 입력·부여 조건·평가방법·가득 회차·조건 변경표·출처. 돌려주는 refs 는 수식이 가리키는 칸이다."""
    t, s, raw, md = run.terms, run.summary, run.raw, run.raw['mode']
    A = wb.create_sheet(ASSUME)
    A.sheet_view.showGridLines = False
    for cc, wd in (("B", 34), ("C", 24), ("D", 16), ("E", 16), ("F", 16), ("G", 16), ("H", 20)):
        A.column_dimensions[cc].width = wd
    span = 7
    w.title(A, 2, "주식선택권 평가 조서", span=span)
    w.note(A, 3, f"평가 {s['calculated_at'][:16].replace('T', ' ')} (UTC)   ·   금액은 원 — 1주당 · 총액")
    w.note(A, 4, "적용범위 — 주식결제형 부여일 평가 · 조건 변경 증분 공정가치 · 현금결제형 결산일 재측정. 회차마다 이항격자로 "
                 "1주당 공정가치를 재고, 근무조건·비시장 성과조건은 가득 예상 비율(수량)에 반영한다.")
    w.note(A, 5, "반영하지 않은 것 — 희석 효과 조정, 행사가격·수량 조정 조항, 행사하면 새 선택권을 다시 주는 조건. "
                 "해당하는 계약이면 따로 검토하십시오.", color=AMB)
    refs = {}
    r = 7

    def block(ttl, items):
        nonlocal r
        w.sec(A, r, ttl, span=span)
        r += 1
        for nm, v, fm in items:
            w.put(A, r, 2, nm, border=True)
            num = isinstance(v, (int, float)) and not isinstance(v, bool)
            w.put(A, r, 3, v, fmt=fm, align="right" if num else "left", border=True)
            r += 1
        r += 1

    stamp = L.run_stamp(t)
    block("0. 재현 기록", [
        ("평가 실행번호", s['run_id'], None), ("평가시각 (UTC)", s['calculated_at'], None),
        ("앱 버전 식별값 (app.py 앞 12자리)", stamp['app_sha12'] or "확인 못 함", None),
        ("계산 코드 버전 (git)", stamp['git_head'] or "확인 못 함", None),
        ("평가체계 버전", f"{stamp['schema']}", None),
        ("입력 식별값", s['case_sha256'], None), ("계산 코드 식별값", s['code_sha256'], None)])
    # 1. 평가 입력값 — 수식 조서에서 고칠 수 있는 칸
    w.sec(A, r, "1. 평가 입력값", span=span)
    r += 1
    items = [('S0', '측정기준일 주가(원)', t.S0, IN_NUM, '격자의 출발 주가 · 블랙-숄즈 대조의 주가'),
             ('sig', '변동성 σ (연)', t.sig, IN_PCT, '상승계수 u = exp(σ × √Δt), 하락계수 d = 1 ÷ u'),
             ('div_y', '배당수익률 (연, 연속복리)', t.div_y, IN_PCT, '상승확률 q = (exp((r − 배당수익률) × Δt) − d) ÷ (u − d)'),
             ('so_mult', '행사 배수 (주가 ÷ 행사가격)', t.so_mult, IN_NUM, '«행사 배수에 도달하면 행사» 가정에서만 씁니다'),
             ('so_exit', '가득 뒤 연간 퇴사율', t.so_exit, IN_PCT, '구간 퇴사 확률 = 1 − (1 − 퇴사율)^Δt'),
             ('so_nv', '비가득조건 연간 중단율 (가득일까지)', getattr(t, 'so_nv', 0.0), IN_PCT,
              '가득 전 구간마다 (1 − 구간 중단 확률) 을 곱함 (문단 21A)')]
    if md == 'cash':
        items += [('so_prev_liab', '전기말 주식보상부채(원, -1은 없음)', t.so_prev_liab, IN_NUM,
                   '당기 비용 = 당기말 부채 − 전기말 부채(없으면 0) + 당기 결제액'),
                  ('so_paid', '당기 현금 결제액(원)', t.so_paid, IN_NUM, '당기 중 행사되어 현금으로 지급한 금액')]
    for key, nm, v, fm, desc in items:
        w.put(A, r, 2, nm, border=True)
        w.value_in(A, r, 3, float(v), fm)
        w.put(A, r, 4, desc, color=GREY, size=9)
        refs[key] = _abs(ASSUME, 'C', r)
        r += 1
    r += 1
    K = t.K0 if md != 'mod' else f"{t.K0:,.2f} (변경 전 — 변경 후 값은 조건 변경표)"
    block("2. 부여 조건", [
        ("평가 종류", SOR.MODE_TEXT[md], None), ("결제 방식", SO.SETTLE_LABELS[int(t.so_settle)], None),
        ("차수", t.tranche or "—", None), ("부여일", t.d_issue, None),
        ("측정기준일" + {'grant': '', 'mod': ' (조건 변경일)', 'cash': ' (결산일)'}[md], t.d_base, None),
        ("행사가격(원)", K, N2 if md != 'mod' else None), ("행사기간 종료일 (마지막 회차)", t.d_mat, None),
        ("결산일 (월-일)", t.so_fy_end, None), ("회차 수", len(raw['rows']), N0),
        ("부여 수량 합계(주)", raw['qty'], N0)])
    sources = run.case.sources or {}
    block("3. 평가방법", [
        ("평가모형", "이항격자 (CRR) — 회차마다 따로 · 블랙-숄즈 대조", None),
        ("계산 간격", f"{t.grid_days:g}일" if t.grid_days else f"{t.gap_m:g}개월", None),
        ("조기행사 가정", SO.EX_LABELS[int(t.so_ex)] + (f" · 행사 배수 {t.so_mult:g}배" if int(t.so_ex) == 2 else ""), None),
        ("시장조건 판정", SO.MKT_LABELS[int(getattr(t, 'so_mkt', 0))], None),
        ("무위험 금리곡선", ("만기수익률 곡선" if t.y_type == "par" else "현물이자율 곡선") + f" · 연 {t.cmp_rf}회", None),
        ("금리곡선 출처", sources.get('rf_curve', '미기록'), None),
        ("주가 출처", t.s0_src or sources.get('S0', '직접 입력'), None),
        ("신용위험", "반영하지 않음 — 신용위험이 없는 지분상품", None)])
    # 4. 가득 회차 — 가득 예상 비율은 평가자 추정이라 수식 조서에서 고칠 수 있다
    w.sec(A, r, "4. 가득 회차", span=span)
    r += 1
    heads = ['회차', '가득일 (비용 배분 종료)', '행사 시작일', '행사 종료일', '부여 수량(주)', '가득 예상 비율',
             '시장조건 목표 주가(원)']
    w.head(A, r, 2, heads)
    A.row_dimensions[r].height = 30
    r += 1
    refs['keep'] = {}
    for x, out in zip(SO.rows(t), raw['rows']):
        vals = [x['name'], x['vest'], x['start'], x['end'], x['qty'], None, SO.target(x)]
        fmts = [None, None, None, None, N0, None, N2]
        for k, (v, fm) in enumerate(zip(vals, fmts)):
            if k == 5:
                w.value_in(A, r, 7, float(x['keep']), IN_PCT)
            else:
                w.put(A, r, 2 + k, str(v) if k == 0 else v, fmt=fm, align="right" if k else None, border=True)
        refs['keep'][out['name']] = _abs(ASSUME, 'G', r)
        r += 1
    r += 1
    refs['grant_fv'] = {}
    if md == 'mod':
        w.sec(A, r, "5. 조건 변경표 — 바뀐 칸만 · 부여일 1주당 공정가치는 부여일 평가 조서의 값", span=span)
        r += 1
        table = SOR.mod_input_table(run)
        w.head(A, r, 2, table[0])
        A.row_dimensions[r].height = 30
        r += 1
        for row in table[1:]:
            for k, v in enumerate(row):
                if k == 6:
                    w.value_in(A, r, 8, float(v), IN_NUM)
                else:
                    num = isinstance(v, (int, float)) and not isinstance(v, bool)
                    w.put(A, r, 2 + k, str(v) if k == 0 else v, fmt=N2 if num else None,
                          align="right" if k else None, border=True)
            refs['grant_fv'][str(row[0]).strip()] = _abs(ASSUME, 'H', r)
            r += 1
        r += 1
    if sources:
        from .presentation import label
        w.sec(A, r, "6. 출처·평가가정", span=span)
        r += 1
        for k, v in sources.items():
            w.put(A, r, 2, label(k), border=True)
            w.put(A, r, 3, str(v), border=True)
            r += 1
        r += 1
    w.note(A, r, "빨강 글씨는 입력값입니다. 수식 조서에서는 노란 칸이 고칠 수 있는 입력이며, 바꾸면 격자·결과·비용 배분·회계처리가 "
                 "다시 계산됩니다.")
    return refs


# ── 격자 ─────────────────────────────────────────────────────────────────

def lattice_sheet(w, wb, run, row, name, title, refs):
    """한 회차의 격자 — 1~17행 공통조건, ① 주가, ② 옵션 가치, (③ 목표 주가 미도달 상태). 돌려주는 값은 뿌리 칸."""
    t = run.terms
    n, dt_ = row['n'], row['dt']
    W = wb.create_sheet(name)
    W.sheet_view.showGridLines = False
    W.column_dimensions["B"].width = 26
    for i in range(n + 1):
        W.column_dimensions[gl(3 + i)].width = 10
    pos = layout(n)
    K = float(row.get('K') or t.K0)
    H = row.get('target')
    mkt, vest_step, ex_step = int(row.get('mkt') or 0), int(row['vest_step']), int(row['ex_step'])
    e_step = 1.0 - (1.0 - float(t.so_exit))**dt_ if t.so_exit else 0.0
    db = dt.date.fromisoformat(t.d_base)
    for h, r in HR.items():
        w.put(W, r, 2, h, bold=True, size=8, fill=LIGHT, border=True)
    for i in range(n + 1):
        g = lambda r, v, fm=None, col="000000", i=i: w.put(W, r, 3 + i, v, fmt=fm, align="center", size=8, color=col)
        g(HR['날짜'], db + dt.timedelta(days=round(i*dt_*365)), DATE, GREY)
        g(HR['스텝(노드 번호)'], i, N0)
        g(HR['행사 가능 (1=예)'], 1 if i >= ex_step else 0, N0)
        g(HR['가득 전 (1=예)'], 1 if i < vest_step else 0, N0)
        g(HR['시장조건 판정 (1=예)'], 1 if (mkt == 1 and i == vest_step) or (mkt == 2 and i <= vest_step) else 0, N0, RED)
        g(HR['행사가격'], K, N2)
        if H is not None:
            g(HR['시장조건 목표 주가'], float(H), N2)
        g(HR['구간 길이 Δt (년)'], dt_, N6)
        g(HR['주가변동성 σ'], float(t.sig), P2)
        g(HR['상승계수 u'], row['u'], N6)
        g(HR['하락계수 d'], row['d'], N6)
        if i < n:
            g(HR['무위험 선도이자율'], row['fwd'][i], P4)
            g(HR['위험중립 상승확률 q'], row['p'][i], N6)
            g(HR['하락확률 1−q'], 1 - row['p'][i], N6)
            g(HR['할인계수 exp(−r × Δt)'], math.exp(-row['fwd'][i]*dt_), N6)
            g(HR['구간 퇴사 확률'], e_step if i >= ex_step else 0.0, P4)
            g(HR['구간 비가득조건 중단 확률'], row['nv_step'] if i < vest_step else 0.0, P4)
    trees = [(pos['stock'], "① 주가트리  S = S0 × u^(스텝−r) × d^r", "위로 갈수록 주가 상승 횟수가 많은 경로다.",
              row['prices'], n),
             (pos['value'], "② 옵션가치트리  V = 행사(주가 − 행사가격)와 계속 보유 가운데 조기행사 가정이 고른 값",
              "주황 칸은 그 노드에서 행사한다.", row['values'], n)]
    if row.get('wait') is not None:
        trees.append((pos['wait'], "③ 목표 주가 미도달 상태 가치트리  W — 목표에 닿으면 ② 의 값, 가득 시점까지 못 닿으면 0",
                      f"가득 시점 {vest_step}까지만 있다. 시점 0의 값이 1주당 공정가치다.", row['wait'], vest_step))
    for r0, ttl, note, grid, last in trees:
        w.title(W, r0 - 2, ttl, span=min(n + 1, 14))
        w.put(W, r0 - 2, 3 + min(n + 1, 14), note, color=GREY, size=9)
        w.put(W, r0 - 1, 2, "하락 횟수 r ＼ 스텝", bold=True, size=8, fill=LIGHT, border=True, align="center")
        for i in range(n + 1):
            w.put(W, r0 - 1, 3 + i, i, bold=True, size=8, fmt=N0, align="center",
                  fill=(RFXC if i == vest_step else LIGHT), border=True)
        for r in range(last + 1):
            w.put(W, r0 + r, 2, r, bold=True, size=8, fmt=N0, align="center", fill=LIGHT, border=True)
        mark = grid is row['values']
        for i in range(last + 1):
            for r in range(i + 1):
                go = mark and row['exercise'][i][r] and not w.formula
                w.put(W, r0 + r, 3 + i, grid[i][r], fmt=N2, size=8, align="right", fill=(RFXC if go else None))
    end = (pos['wait'] + vest_step) if row.get('wait') is not None else (pos['value'] + n)
    w.note(W, end + 2, "r 은 하락 횟수. 위로 갈수록 주가가 높다. 주황 열은 가득 시점이다.")
    w.put(W, end + 3, 2, "이 시트가 참조하는 시트: 가정", color=GREEN, size=9)
    W.freeze_panes = "C20"
    if w.formula:
        formula_lattice(w, W, run, row, refs)
    root = pos['wait'] if row.get('wait') is not None else pos['value']
    return f"C{root}"


def formula_lattice(w, W, run, row, refs):
    """격자 칸마다 엔진(value_row)과 같은 순서의 수식."""
    t, n = run.terms, row['n']
    pos = layout(n)
    S0, SIG, Q, MULT, EXIT, NV = (refs[k] for k in ('S0', 'sig', 'div_y', 'so_mult', 'so_exit', 'so_nv'))
    mkt, vest_step, ex_step, mode = int(row.get('mkt') or 0), int(row['vest_step']), int(row['ex_step']), int(t.so_ex)
    R0, R1, R2 = pos['stock'], pos['value'], pos['wait']
    h = HR
    for i in range(n + 1):
        c = gl(3 + i)
        w.fx(W, h['주가변동성 σ'], 3 + i, f"={SIG}")
        w.fx(W, h['상승계수 u'], 3 + i, f"=EXP({c}{h['주가변동성 σ']}*SQRT({c}{h['구간 길이 Δt (년)']}))")
        w.fx(W, h['하락계수 d'], 3 + i, f"=1/{c}{h['상승계수 u']}")
        if i < n:
            r_, dt_ = f"{c}{h['무위험 선도이자율']}", f"{c}{h['구간 길이 Δt (년)']}"
            u_, d_ = f"{c}{h['상승계수 u']}", f"{c}{h['하락계수 d']}"
            w.fx(W, h['위험중립 상승확률 q'], 3 + i, f"=(EXP(({r_}-{Q})*{dt_})-{d_})/({u_}-{d_})")
            w.fx(W, h['하락확률 1−q'], 3 + i, f"=1-{c}{h['위험중립 상승확률 q']}")
            w.fx(W, h['할인계수 exp(−r × Δt)'], 3 + i, f"=EXP(-{r_}*{dt_})")
            w.fx(W, h['구간 퇴사 확률'], 3 + i, f"=IF({c}{h['행사 가능 (1=예)']}=1,1-(1-{EXIT})^{dt_},0)")
            w.fx(W, h['구간 비가득조건 중단 확률'], 3 + i, f"=IF({c}{h['가득 전 (1=예)']}=1,1-(1-{NV})^{dt_},0)")
    U = f"$C${h['상승계수 u']}"
    for i in range(n + 1):
        c = gl(3 + i)
        K, TGT, Uc = f"{c}${h['행사가격']}", f"{c}${h['시장조건 목표 주가']}", f"{c}${h['상승계수 u']}"
        P, NP, DF = f"{c}${h['위험중립 상승확률 q']}", f"{c}${h['하락확률 1−q']}", f"{c}${h['할인계수 exp(−r × Δt)']}"
        E, NVS = f"{c}${h['구간 퇴사 확률']}", f"{c}${h['구간 비가득조건 중단 확률']}"
        for r in range(i + 1):
            w.fx(W, R0 + r, 3 + i, f"={S0}*{U}^({i - 2*r})")
            S = f"{c}{R0 + r}"
            if i == n:
                f = f"=MAX({S}-{K},0)"
            else:
                c1 = gl(4 + i)
                cont = f"{DF}*({P}*{c1}{R1 + r}+{NP}*{c1}{R1 + r + 1})"
                if i < ex_step:
                    f = f"=(1-{NVS})*({cont})"          # 가득 전 비가득조건 (가득 뒤·중단율 0이면 그대로)
                else:
                    Hh = f"((1-{E})*({cont})+{E}*MAX({S}-{K},0))"
                    if mode == 1:
                        f = f"=IF({S}-{K}>{Hh}+1E-12*MAX(1,ABS({Hh})),{S}-{K},{Hh})"
                    elif mode == 2:
                        f = f"=IF(AND({S}>={MULT}*{K}-1E-12*{K},{S}-{K}>0),{S}-{K},{Hh})"
                    else:
                        f = f"={Hh}"
            if mkt == 1 and i == vest_step:
                # 가득일 주가가 목표 미만이면 소멸 — 목표가 걸친 노드는 구간 가운데 목표 이상인 비율만큼 (엔진 pass_share)
                f = f"=MIN(1,MAX(0,(LN({S})+LN({Uc})-LN({TGT}))/(2*LN({Uc}))))*({f[1:]})"
            w.fx(W, R1 + r, 3 + i, f)
    # 행사 노드 색 — 행사할 수 있는 시점에서 옵션 가치가 지금 행사한 값(주가 − 행사가격 > 0)과 같은 칸
    last = gl(3 + n)
    rule = FormulaRule(formula=[f"AND(C${h['행사 가능 (1=예)']}=1,C{R0}-C${h['행사가격']}>0,"
                                f"ABS(C{R1}-(C{R0}-C${h['행사가격']}))<=1E-9*MAX(1,ABS(C{R1})))"],
                       fill=PatternFill("solid", fgColor=RFXC))
    W.conditional_formatting.add(f"C{R1}:{last}{R1 + n}", rule)
    if mkt == 2:
        for i in range(vest_step + 1):
            c = gl(3 + i)
            TGT = f"{c}${h['시장조건 목표 주가']}"
            P, NP, DF = f"{c}${h['위험중립 상승확률 q']}", f"{c}${h['하락확률 1−q']}", f"{c}${h['할인계수 exp(−r × Δt)']}"
            NVS = f"{c}${h['구간 비가득조건 중단 확률']}"
            for r in range(i + 1):
                S, V = f"{c}{R0 + r}", f"{c}{R1 + r}"
                hit = f"{S}>={TGT}*(1-{SO.MKT_TOL:g})"
                if i == vest_step:
                    f = f"=IF({hit},{V},0)"
                else:
                    c1 = gl(4 + i)
                    f = f"=IF({hit},{V},(1-{NVS})*({DF}*({P}*{c1}{R2 + r}+{NP}*{c1}{R2 + r + 1})))"
                w.fx(W, R2 + r, 3 + i, f)


# ── 결과 ─────────────────────────────────────────────────────────────────

FMT = {'가득 예상 비율': IN_PCT, '근무 경과율': '0.0000%', '누적 배분율': '0.0000%', '구간 수': N0, '부여 수량': N0,
       '변경 전 수량': N0, '늘어난 수량': N0, '기대존속기간(년)': N4, '기간(년)': N4, '무위험 현물금리(연속)': P4,
       'd1': N6, 'd2': N6, 'N(d1)': N6, 'N(d2)': N6}


def _table(w, ws, r, heads, rows, *, total=False, label=True, fmts=None):
    """머리줄(하늘색) + 테두리 표. 돌려주는 값: ({머리: 열 글자}, 첫 자료 줄). label 이면 첫 열은 글자(회차·항목)."""
    w.head(ws, r, 2, heads)
    ws.row_dimensions[r].height = 30
    cols = {h: gl(2 + k) for k, h in enumerate(heads)}
    first = r + 1
    for k, row in enumerate(rows):
        last = total and k == len(rows) - 1
        for j, v in enumerate(row):
            num = isinstance(v, (int, float)) and not isinstance(v, bool)
            w.put(ws, first + k, 2 + j, (str(v) if label and j == 0 and v is not None else v),
                  fmt=((fmts or {}).get(heads[j]) or FMT.get(heads[j], N2) if num else None),
                  align=("right" if num else None), border=True,
                  bold=last, fill=(BAND if last else None))
    return cols, first


def result_sheet(w, wb, run):
    from .service import AMOUNT_LABELS
    t, s, raw, md = run.terms, run.summary, run.raw, run.raw['mode']
    R = wb.create_sheet(RESULT)
    R.sheet_view.showGridLines = False
    R.column_dimensions["B"].width = 40
    for k in range(3, 20):
        R.column_dimensions[gl(k)].width = 19
    rows_t = SOR.row_table(run)
    bs_t = SOR.bs_table(run)
    span = max(5, len(rows_t[0]), len(bs_t[0]))
    w.title(R, 2, "평가결과", span=span)
    out = dict(amount={})
    w.sec(R, 4, "1. 평가 금액", span=span)
    amounts = []
    for key in SO.AMOUNT_KEYS[md]:
        amounts.append((key, [AMOUNT_LABELS[key], s['amounts_total'][key], s['amounts_per_share'][key],
                              s['amounts_100'][key]]))
    if md == 'mod':
        amounts.append(('grant_cost', ['당초 보상원가(원) — 부여일 공정가치 기준', raw['grant_cost'], None, None]))
    if md == 'cash':
        amounts += [('prev_liab', ['전기말 주식보상부채(원)', raw['prev_liab'], None, None]),
                    ('paid', ['당기 현금 결제액(원)', raw['paid'], None, None]),
                    ('expense', [SOR.EXPENSE_LABEL, raw['expense'], None, None])]
    _, first = _table(w, R, 5, ['항목', '총액(원)', '1주당(원)', '측정기준일 주가 100 기준'], [v for _, v in amounts])
    out['amount'] = {key: first + k for k, (key, _) in enumerate(amounts)}
    r = first + len(amounts) + 1
    w.sec(R, r, "2. " + {'grant': '회차별 공정가치와 보상원가', 'mod': '회차별 조건 변경 전·후 공정가치와 증분',
                         'cash': '회차별 공정가치와 주식보상부채'}[md], span=span)
    out['rows_cols'], out['rows_first'] = _table(w, R, r + 1, rows_t[0], rows_t[1:], total=True)
    out['rows_total'] = out['rows_first'] + len(rows_t) - 2
    r = out['rows_total'] + 2
    w.sec(R, r, "3. 블랙-숄즈 대조 — 조기행사·퇴사·조건이 없을 때의 값", span=span)
    out['bs_cols'], out['bs_first'] = _table(w, R, r + 1, bs_t[0], bs_t[1:])
    out['bs_rows'] = {(x[0], x[1]): out['bs_first'] + k for k, x in enumerate(bs_t[1:])}
    r = out['bs_first'] + len(bs_t) + 0
    notes = {'grant': "보상원가 = 1주당 공정가치 × 부여 수량 × 가득 예상 비율. 공정가치는 부여 수량 전체 기준이다.",
             'mod': "증분 = 큰 값(변경 후 − 변경 전, 0) × 변경 전 수량 + 변경 후 1주당 × 늘어난 수량. 줄어든 공정가치는 인식하지 "
                    "않는다 (문단 27).",
             'cash': "부채 = 1주당 공정가치 × 부여 수량 × 가득 예상 비율 × 근무 경과율. 결산일마다 다시 잰다 (문단 30)."}
    w.note(R, r + 1, notes[md])
    w.note(R, r + 2, "블랙-숄즈 값은 조기행사·퇴사·조건이 없을 때의 대조값이다. 만기에만 행사하고 조건이 없으면 격자값과 같아야 한다.")
    R.freeze_panes = "C4"
    return out


def schedule_sheet(w, wb, run):
    md = run.raw['mode']
    C = wb.create_sheet(SCHED)
    C.sheet_view.showGridLines = False
    C.column_dimensions["B"].width = 16
    for k in range(3, 9):
        C.column_dimensions[gl(k)].width = 18
    table = SOR.schedule_table(run)
    w.title(C, 2, "결산일 주식보상부채" if md == 'cash' else "회계연도별 보상원가 배분", span=max(5, len(table[0])))
    w.sec(C, 4, "1. 결산일별 " + ("부채" if md == 'cash' else "배분"), span=max(5, len(table[0])))
    cols, first = _table(w, C, 5, table[0], table[1:])
    w.note(C, first + len(table), '근무 경과율 = (결산일 − 부여일) ÷ (가득일 − 부여일), 1을 넘지 않음.' if md == 'cash' else
           '누적 배분율 = (결산일 − 배분 시작일) ÷ (가득일 − 배분 시작일), 1을 넘지 않음. 배분 시작일은 부여일, 조건 변경의 증분은 '
           '조건 변경일. 당기 비용 = 누적 비용 − 전기 누적 비용.')
    C.freeze_panes = "C6"
    return cols, first


def journal_sheet(w, wb, run):
    """회계처리 — 1행은 «초안» 문구 자리 (report.calculation_sheets_only 가 넣는다)."""
    J = wb.create_sheet(JOURNAL)
    J.sheet_view.showGridLines = False
    J.column_dimensions["B"].width = 22
    for cc, wd in (("C", 34), ("D", 18), ("E", 34), ("F", 18)):
        J.column_dimensions[cc].width = wd
    rows = SOR.journal_rows(run)[1:]              # 첫 줄의 초안 문구는 1행에 따로 들어간다
    w.title(J, 2, "회계처리 (초안)", span=5)
    w.head(J, 4, 2, rows[0])
    at, r = {}, 5
    body = True
    for row in rows[1:]:
        if not row:
            body = False
            r += 1
            continue
        for j, v in enumerate(row):
            num = isinstance(v, (int, float)) and not isinstance(v, bool)
            w.put(J, r, 2 + j, v, fmt=N2 if num else None, align="right" if num else None, border=body,
                  bold=(not body and j == 0))
        at.setdefault(row[0], r)
        r += 1
    return at


def rate_sheet(w, wb, run):
    t = run.terms
    G = wb.create_sheet('금리곡선')
    G.sheet_view.showGridLines = False
    G.column_dimensions["B"].width = 26
    for k in range(3, 8):
        G.column_dimensions[gl(k)].width = 18
    w.title(G, 2, "무위험 금리곡선", span=5)
    w.sec(G, 4, "1. 입력 곡선 — " + ("만기수익률" if t.y_type == "par" else "현물이자율") + f" · 연 {t.cmp_rf}회", span=5)
    # 서식은 표를 그릴 때 정한다 — 칸의 서식을 나중에 고치면 같은 모양을 함께 쓰는 칸이 모두 바뀐다
    _, first = _table(w, G, 5, ['만기(년)', '수익률'], [[float(a), float(b)] for a, b in t.rf_curve], label=False,
                      fmts={'만기(년)': N4, '수익률': IN_PCT})
    r = first + len(t.rf_curve) + 1
    w.sec(G, r, "2. 회차별 적용 금리 (연속복리 현물)", span=5)
    rows = []
    for name, row in SOR.lattices(run).items():
        rows.append([str(name), row['t_end'], row['r_end'], row['life'], SOR._rate(run, row['life'])])
    _table(w, G, r + 1, ['격자', '행사 종료일까지(년)', '무위험 현물금리', '기대존속기간(년)', '그 기간 무위험 현물금리'], rows,
           fmts={'행사 종료일까지(년)': N4, '무위험 현물금리': P4, '기대존속기간(년)': N4, '그 기간 무위험 현물금리': P4})
    w.note(G, r + 3 + len(rows), "격자의 구간별 선도이자율은 이 곡선에서 구한다 (격자 시트 8행). 금리는 앱에서만 바꾼다 — "
                                 f"출처: {(run.case.sources or {}).get('rf_curve', '미기록')}")


# ── 수식 조서의 표 연결 ───────────────────────────────────────────────────

def formula_tables(w, wb, run, refs, roots, res, sched, jat):
    from .service import AMOUNT_LABELS
    t, raw, md = run.terms, run.raw, run.raw['mode']
    S0, SIG, Q = refs['S0'], refs['sig'], refs['div_y']
    R = wb[RESULT]
    names = SOR.lattices(run)

    def root(name):
        return f"{_q(roots[name][0])}!${roots[name][1][0]}${roots[name][1][1:]}"

    def strike(name):
        return f"{_q(roots[name][0])}!$C${HR['행사가격']}"

    # 블랙-숄즈 대조
    bc = res['bs_cols']
    for (name, kind), r in res['bs_rows'].items():
        K = strike(name)
        T, rate = f"{bc['기간(년)']}{r}", f"{bc['무위험 현물금리(연속)']}{r}"
        d1, d2, n1, n2 = (f"{bc[k]}{r}" for k in ('d1', 'd2', 'N(d1)', 'N(d2)'))
        if R[f"{bc['d1']}{r}"].value is None:            # 남은 기간이 없다 — 내재가치
            w.fx(R, r, _col(bc['블랙-숄즈 1주당(원)']), f"=MAX({S0}-{K},0)")
        else:
            w.fx(R, r, _col(bc['d1']), f"=(LN({S0}/{K})+({rate}-{Q}+0.5*{SIG}*{SIG})*{T})/({SIG}*SQRT({T}))")
            w.fx(R, r, _col(bc['d2']), f"={d1}-{SIG}*SQRT({T})")
            w.fx(R, r, _col(bc['N(d1)']), f"=NORMSDIST({d1})")
            w.fx(R, r, _col(bc['N(d2)']), f"=NORMSDIST({d2})")
            w.fx(R, r, _col(bc['블랙-숄즈 1주당(원)']), f"={S0}*EXP(-{Q}*{T})*{n1}-{K}*EXP(-{rate}*{T})*{n2}")
        w.fx(R, r, _col(bc['격자 1주당(원)']), f"={root(name)}")
        w.fx(R, r, _col(bc['격자 − 블랙-숄즈']), f"={bc['격자 1주당(원)']}{r}-{bc['블랙-숄즈 1주당(원)']}{r}")
    # 회차표
    cols, first = res['rows_cols'], res['rows_first']
    at = {}
    for k, x in enumerate(raw['rows']):
        rr = first + k
        name = x['name']
        cell = (lambda h, rr=rr: f"{cols[h]}{rr}")
        at[name] = cell
        put = lambda h, f: w.fx(R, rr, _col(cols[h]), f)
        put('가득 예상 비율', f"={refs['keep'][name]}")
        if md == 'mod':
            put('부여일 1주당 공정가치', f"={refs['grant_fv'][str(name).strip()]}")   # 조건 변경표는 이름 앞뒤 빈칸을 떼고 맞춘다
            put('변경 전 1주당 (변경일)', f"={root(f'{name} 변경 전')}")             # lattices() 와 같은 이름 (숫자 회차명도)
            put('변경 후 1주당 (변경일)', f"={root(f'{name} 변경 후')}")
            put('1주당 증분', f"=MAX({cell('변경 후 1주당 (변경일)')}-{cell('변경 전 1주당 (변경일)')},0)")
            put('당초 보상원가', f"={cell('부여일 1주당 공정가치')}*{cell('변경 전 수량')}*{cell('가득 예상 비율')}")
            put('증분 공정가치', f"={cell('1주당 증분')}*{cell('변경 전 수량')}+{cell('변경 후 1주당 (변경일)')}*{cell('늘어난 수량')}")
            put('증분 보상원가', f"={cell('증분 공정가치')}*{cell('가득 예상 비율')}")
            continue
        put('1주당 공정가치', f"={root(name)}")
        put('블랙-숄즈(행사 종료일)', f"={bc['블랙-숄즈 1주당(원)']}{res['bs_rows'][name, '행사 종료일']}")
        put('공정가치', f"={cell('1주당 공정가치')}*{cell('부여 수량')}")
        if md == 'cash':
            put('가득 예상 수량의 공정가치', f"={cell('1주당 공정가치')}*{cell('부여 수량')}*{cell('가득 예상 비율')}")
            put('주식보상부채', f"={cell('가득 예상 수량의 공정가치')}*{cell('근무 경과율')}")
        else:
            put('보상원가', f"={cell('1주당 공정가치')}*{cell('부여 수량')}*{cell('가득 예상 비율')}")
    tot_r = res['rows_total']
    total = {}
    for h in SOR.SUM_COLS[md]:
        w.fx(R, tot_r, _col(cols[h]), f"=SUM({cols[h]}{first}:{cols[h]}{tot_r - 1})")
        total[h] = f"{cols[h]}{tot_r}"
    # 평가 금액
    am = res['amount']
    sources = {'grant': ('공정가치', '보상원가'), 'mod': ('증분 공정가치', '증분 보상원가'),
               'cash': ('공정가치', '주식보상부채')}[md]
    qty = float(raw['qty'])
    for key, h in zip(SO.AMOUNT_KEYS[md], sources):
        r = am[key]
        w.fx(R, r, 3, f"={total[h]}")
        w.fx(R, r, 4, f"=C{r}/{qty!r}" if qty else "=0")      # 현금결제형에서 모두 결제했으면 남은 수량 0
        w.fx(R, r, 5, f"=D{r}/{S0}*100")
    if md == 'mod':
        w.fx(R, am['grant_cost'], 3, f"={total['당초 보상원가']}")
    if md == 'cash':
        w.fx(R, am['prev_liab'], 3, f"=MAX(0,{refs['so_prev_liab']})")
        w.fx(R, am['paid'], 3, f"={refs['so_paid']}")
        w.fx(R, am['expense'], 3, f"={total['주식보상부채']}-C{am['prev_liab']}+C{am['paid']}")
    # 비용 배분 — 누적 = 총액 × 누적 배분율, 당기 = 누적 − 직전 누적 (회차·구분마다)
    C = wb[SCHED]
    sh, sfirst = sched
    prev, by_end = {}, {}
    rr_ = f"{_q(RESULT)}!"
    for k, x in enumerate(raw['schedule']):
        r = sfirst + k
        if md == 'cash':
            w.fx(C, r, _col(sh['주식보상부채']), f"={rr_}{at[x['name']]('주식보상부채')}")
            continue
        src = {'grant': '보상원가', 'mod': {'당초': '당초 보상원가', '증분': '증분 보상원가'}.get(x.get('part'))}[md]
        w.fx(C, r, _col(sh['누적 비용']), f"={rr_}{at[x['name']](src)}*{sh['누적 배분율']}{r}")
        key = (x['name'], x.get('part'))
        w.fx(C, r, _col(sh['당기 비용']), (f"={sh['누적 비용']}{r}-{sh['누적 비용']}{prev[key]}" if key in prev
                                          else f"={sh['누적 비용']}{r}"))
        prev[key] = r
        by_end.setdefault(x['fy_end'], []).append(f"{_q(SCHED)}!{sh['당기 비용']}{r}")
    if jat is None:
        return
    J = wb[JOURNAL]
    if md == 'cash':
        v = f"{_q(RESULT)}!$C${am['expense']}"
        r = 5
        w.fx(J, r, 3, f'=IF({v}>=0,"주식보상비용","주식보상부채")')
        w.fx(J, r, 4, f"=ABS({v})")
        w.fx(J, r, 5, f'=IF({v}>=0,"주식보상부채","주식보상비용 (환입)")')
        w.fx(J, r, 6, f"=ABS({v})")
        paid = f"{_q(RESULT)}!$C${am['paid']}"
        w.fx(J, jat['당기 중'], 4, f"={paid}")
        w.fx(J, jat['당기 중'], 6, f"={paid}")
        w.fx(J, jat['전기말 주식보상부채'], 3, f"={_q(RESULT)}!$C${am['prev_liab']}")
        w.fx(J, jat['당기 현금 결제액'], 3, f"={paid}")
        w.fx(J, jat['당기말 주식보상부채'], 3, f"={_q(RESULT)}!{total['주식보상부채']}")
        return
    for d, refs_ in by_end.items():
        r = jat.get(d)
        if r is not None and J.cell(r, 3).value == '주식보상비용':
            w.fx(J, r, 4, '=' + '+'.join(refs_))
            w.fx(J, r, 6, f"=D{r}")


def _col(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


def key_cells(run, roots, res):
    """다시 계산했을 때 앱과 같아야 하는 칸 — [(항목, 시트, 칸, 앱 값)]. 조서의 «계산대사» 시트와 시험이 쓴다."""
    from .service import AMOUNT_LABELS
    raw, s = run.raw, run.summary
    am = res['amount']
    out = []
    for key in SO.AMOUNT_KEYS[raw['mode']]:
        out += [(AMOUNT_LABELS[key] + ' (원)', RESULT, f'C{am[key]}', s['amounts_total'][key]),
                (AMOUNT_LABELS[key] + ' (주가 100 기준)', RESULT, f'E{am[key]}', s['amounts_100'][key])]
    for name, row in SOR.lattices(run).items():
        sheet, cell = roots[name]
        out.append((f'{name} 1주당 공정가치', sheet, cell, row['fv']))
    if raw['mode'] == 'mod':
        out.append(('당초 보상원가 (원)', RESULT, f"C{am['grant_cost']}", raw['grant_cost']))
    if raw['mode'] == 'cash':
        out.append(('당기 비용 (원)', RESULT, f"C{am['expense']}", raw['expense']))
    return out


def detail_workbook(run, *, accounting=False, formula=False):
    """상세 조서 — 전환사채 상세 조서와 같은 서식. 값 조서와 수식 조서는 같은 시트·같은 칸이다."""
    w = _Writer(formula)
    wb = Workbook()
    wb.remove(wb.active)
    guide_sheet(w, wb, run)
    refs = assumption_sheet(w, wb, run)
    used = set(wb.sheetnames) | {RESULT, SCHED, JOURNAL, '금리곡선'}
    roots = {}
    for name, row in SOR.lattices(run).items():
        sheet = SOR._sheet_name(f"격자 {name}", used)
        roots[name] = (sheet, lattice_sheet(w, wb, run, row, sheet, name, refs))
    res = result_sheet(w, wb, run)
    sched = schedule_sheet(w, wb, run)
    jat = journal_sheet(w, wb, run) if accounting else None
    rate_sheet(w, wb, run)
    if formula:
        formula_tables(w, wb, run, refs, roots, res, sched, jat)
    wb.so_inputs = w.inputs
    wb.so_key_cells = key_cells(run, roots, res)
    L.polish_wb(wb)
    return wb
