"""Workbook integrity and optional full-grid spreadsheet-engine verification."""
import io
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree as ET


def inspect_workbook(data):
    errors, longest = [], 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        bad = archive.testzip()
        if bad:
            errors.append(f'ZIP 손상: {bad}')
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                with archive.open(name) as stream:
                    for _, node in ET.iterparse(stream, events=('end',)):
                        if node.tag == '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f':
                            size = len(node.text or '')
                            longest = max(longest, size)
                            if size > 8192:
                                errors.append(f'{name}: 수식 {size:,}자 / 한도 8,192자')
                            if '#REF!' in (node.text or ''):
                                errors.append(f'{name}: 잘못된 시트·셀 참조')
                        node.clear()
    return {'errors': errors, 'max_formula_length': longest}


def compare_cells(wb, expected, tol=1e-6):
    """재계산한 통합문서에서 앱 값과 같아야 하는 칸을 견준다 — [(항목, 칸, 앱, 수식, 판정)].

    ``expected`` 는 (항목, 시트, 칸, 앱 값) 목록이다(legacy.formula_key_cells). 허용오차는
    금액 크기에 비례한다(1 이하 1e-6).
    """
    rows = []
    for label, sheet, cell, want in expected:
        have = None
        if sheet in wb.sheetnames:
            v = wb[sheet][cell].value
            have = float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v
        ok = isinstance(have, float) and abs(have - want) <= tol*max(1.0, abs(want))
        rows.append((label, f"{sheet}!{cell}", want, have, "일치" if ok else "불일치"))
    return rows


def recalculate_and_compare(data, executable=None, expected=None):
    """Recalculate the ACTUAL export, without changing inputs or step count.

    ``expected`` — 앱 값과 같아야 하는 칸(legacy.formula_key_cells). 주면 그 칸들을 견준다.
    예전에는 조서 안의 «수식대사» 시트를 읽었는데 그 시트를 조서에서 뺀 뒤로 검사가 늘
    실패했다. 이제 검사하는 쪽이 앱 값을 들고 온다. 옛 조서에 수식대사 시트가 있으면 그것도 읽는다.
    """
    import openpyxl
    exe = executable or os.environ.get('VALUATION_SOFFICE') or shutil.which('libreoffice') or shutil.which('soffice')
    if not exe:
        raise ValueError('수식 재계산 검사에는 LibreOffice가 필요합니다. 설치 후 실행하거나 VALUATION_SOFFICE 경로를 지정하십시오.')
    check = inspect_workbook(data)
    if check['errors']:
        raise ValueError('; '.join(check['errors']))
    with tempfile.TemporaryDirectory(prefix='valuation-xlsx-') as folder:
        base = Path(folder); source = base/'source.xlsx'; source.write_bytes(data)
        out = base/'recalculated'; out.mkdir()
        result = subprocess.run([exe, '-env:UserInstallation='+ (base/'profile').as_uri(),
            '--headless', '--convert-to', 'xlsx', '--outdir', str(out), str(source)],
            capture_output=True, text=True, timeout=1800)
        target = out/source.name
        if result.returncode or not target.exists():
            raise ValueError('엑셀 재계산 실패: '+(result.stderr or result.stdout)[-2000:])
        wb = openpyxl.load_workbook(target, data_only=True, read_only=True)
        rows = []
        if expected:
            rows = compare_cells(wb, expected)
        elif '수식대사' in wb.sheetnames:
            rows = list(wb['수식대사'].values)[1:]
        else:
            wb.close()
            raise ValueError('견줄 앱 값이 없습니다 — expected(legacy.formula_key_cells)를 넘겨 주십시오.')
        errors = [f'{r[0]}: 앱 {r[2] if expected else r[1]} / 수식 {r[3] if expected else r[2]}'
                  for r in rows if r[4] != '일치']
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'e':
                        errors.append(f'{sheet.title}!{cell.coordinate}: {cell.value}')
        wb.close()
        if errors:
            raise ValueError('수식 대사 불일치: '+'; '.join(errors[:20]))
        return {**check, 'comparison': rows, 'engine': 'LibreOffice', 'same_grid': True}
