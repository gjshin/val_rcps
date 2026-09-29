"""Human-reviewed source cards. References inform, never decide accounting policy.

Book excerpts are from the user-supplied 2023 Series 11. The source PDF is not
distributed. AccountingWiki is a convenient mirror/commentary, not the issuer
of K-IFRS. No date-sensitive legal requirements are inferred from the book.
"""
import copy
import hashlib
import json

REVIEWED_ON = '2026-09-28'
BOOK_TITLE = '한국공인회계사회, K-IFRS 실무사례와 해설 Series 11 복합금융상품 (2023)'
BOOK_SHA256 = '7f6bfa85c78f7ea6a333205c5c846b4905fa39e8c755c581c9cc594d1c2b53d6'
OFFICIAL_URL = 'https://www.kasb.or.kr/front/board/ingAccountingList.do'
WIKI = 'https://accountingwiki.co.kr'


def standard(number, paragraph, quote='', appendix=''):
    return dict(kind='기준서 본문·적용지침', title=f'K-IFRS 제{number}호 {paragraph}',
                quote=quote, url=f'{WIKI}/standards/{number}' + (f'/{appendix}' if appendix else ''),
                official_url=OFFICIAL_URL, location=f'문단 {paragraph}',
                note='회계위키 수록 본문 대조. 공식 시행본·개정 적용일은 한국회계기준원에서 최종 확인.',
                checked_on=REVIEWED_ON)


def book(section, page, quote):
    return dict(kind='실무사례·해설', title=BOOK_TITLE, location=f'{section} · 책 {page}쪽 / PDF {page+20}쪽',
                section=section, printed_page=page, pdf_page=page+20, quote=quote,
                sha256=BOOK_SHA256, note='2023년 실무 해설. 기준서의 의무 규정과 구분하여 적용.', checked_on=REVIEWED_ON)


def article(slug, title):
    return dict(kind='회계위키 해설', title=title, url=f'{WIKI}/articles/{slug}', quote='',
                location='해당 주제 해설', note='비공식 해설. 계약과 기준서에 비추어 판단.', checked_on=REVIEWED_ON)


TOPICS = [
    dict(id='fair_value_inputs', title='발행가격·시장자료를 그대로 사용해도 되나요?',
         trigger='all', fields=['d_base', 'S0', 'sig', 'rf_curve', 'cr_curve', 'bs_target'],
         explanation='평가기준일의 시장참여자 가정과 관측자료를 우선 확인합니다. 최초 거래가격이 공정가치라는 전제가 충족되는 경우의 보정과, 목표 금액에 단순히 맞추는 계산을 구분합니다.',
         questions=['거래가 정상적이며 비교 가능한 조건입니까?', '자료의 기준일·출처·조정 이유를 기록했습니까?', '관측되지 않는 변수와 평가값의 민감도를 확인했습니까?'],
         limitation='역산값이 존재하거나 입력자료가 채워졌다는 사실만으로 공정가치 측정의 적정성을 확인할 수 없습니다.',
         sources=[standard('1113', '61, 64, 67', '관련된 관측할 수 있는 투입변수를 최대한으로 사용하고 관측할 수 없는 투입변수를 최소한으로 사용한다.')]),
    dict(id='classification', title='부채·자본 분류는 무엇으로 결정하나요?',
         trigger='all', fields=['view', 'conv_class', 'inst'],
         explanation='상품명이나 계산된 전환권 가치만으로 분류하지 않습니다. 현금 지급의무, 결제 선택권, 자기지분 결제 조건을 계약별로 확인합니다.',
         questions=['누가 현금 상환을 요구할 수 있습니까?', '전환 시 금액과 주식 수가 확정되어 있습니까?', '리픽싱·조건부 결제·통화 관련 예외를 검토했습니까?'],
         limitation='앱의 전환권 분류 선택은 평가자의 입력입니다. 선택값 자체가 회계분류를 입증하지 않습니다.',
         sources=[standard('1032', '15, 16, 22, 25', '계약의 실질과 금융부채, 금융자산, 지분상품의 정의에 따라'),
                  book('2.1.1.3 복합금융상품의 분리', 20, '비파생 복합금융상품은 먼저 부채요소와 자본요소로 분리되어야 한다.'),
                  article('debt-equity-classification-basics', '금융부채와 자본의 분류')]),
    dict(id='redemption_constraint', title='배당가능이익이 없으면 상환권을 어떻게 반영하나요?',
         trigger='rcps', fields=['p_s', 'p_e', 'mat_mode'],
         explanation='계약상 상환의무의 분류와 실제 회수시점·금액의 측정을 따로 검토합니다. 배당가능이익·현금 부족만으로 계약상 의무가 없어지는 것은 아닙니다.',
         questions=['배당가능이익과 지급재원이 확보되는 근거가 있습니까?', '감액배당·차환 계획은 승인 및 실행 가능성이 뒷받침됩니까?', '상환 제한을 무시한 값과 반영한 값의 차이는 얼마입니까?'],
         limitation='DCF 손익계획 한 가지를 모든 주가 경로의 상환 가능 여부에 그대로 적용하는 것은 이론적 제약이 있습니다. 일정 가정 분석은 조건부 시나리오입니다.',
         sources=[standard('1032', 'AG25', '계약상 의무가 무효화되는 것은 아니다.', 'b'),
                  book('3.3.3.2 상환권 행사의 제약', 91, '발행기업이 현실적으로 상환할 수 있는 시점에 대한 고려가 필요하다.'),
                  book('3.3.3.3~4 감액·추정손익 가정', 92, '이익접근법에 적용된 추정손익은 미래에 발생가능한 여러 가지 시나리오 중에 하나로'),
                  article('debt-equity-classification-basics', '상환 의무와 이행 능력의 구분')]),
    dict(id='dividends', title='누적배당과 존속기간 연장은 어떻게 확인하나요?',
         trigger='rcps', fields=['cpn', 'div_mode', 'div_basis', 'd_mat'],
         explanation='누적 여부만 보지 않고 배당 결정권, 미지급액의 상환가액 가산 여부, 자동전환·기간 연장 조건을 각각 읽어야 합니다.',
         questions=['미지급 배당에 별도 이자가 붙습니까?', '배당금은 상환 또는 전환 시 어떻게 정산됩니까?', '배당 미지급 시 연장 기간과 연장 중 권리는 무엇입니까?'],
         limitation='기존 배당 설정만으로 모든 미지급 누적·자동 연장 조항이 반영되지는 않습니다. 계약별 현금흐름 일정을 추가 분석하십시오.',
         sources=[standard('1032', 'AG26', '분배를 발행자의 재량으로 결정하는 경우', 'b'),
                  article('debt-equity-classification-basics', '우선주의 배당 조건과 분류')]),
    dict(id='holder', title='투자자 관점에서도 내재파생상품을 분리하나요?',
         trigger='holder', fields=['view', 'fvpl_whole'],
         explanation='주계약이 제1109호 적용 금융자산인 복합계약은 전체 계약에 금융자산 분류 규정을 적용합니다. 구성요소 계산표를 그대로 회계상 분리액으로 쓰지 않습니다.',
         questions=['주계약이 제1109호 적용 금융자산입니까?', '사업모형·계약상 현금흐름 평가를 기록했습니까?', '독립된 별도 금융상품까지 하나로 합치지는 않았습니까?'],
         limitation='투자자라는 이유만으로 모든 계약을 자동 FVTPL로 확정하지 않습니다.',
         sources=[standard('1109', '4.3.2'), article('compound-initial-recognition', '복합금융상품의 최초 인식')]),
    dict(id='embedded', title='풋·콜을 분리하는 조건은 무엇인가요?',
         trigger='issuer', fields=['p_sep', 'k_sep', 'conv_class', 'fvpl_whole'],
         explanation='금융자산 주계약 외의 복합계약은 내재파생상품의 세 가지 분리 요건을 함께 검토합니다. 상환금액의 상각후원가 근접 여부와 상실이자 보상 예외도 확인합니다.',
         questions=['경제적 특성·위험이 주계약과 밀접하게 관련됩니까?', '동일 조건의 별도 상품이 파생상품 정의를 충족합니까?', '전체 계약이 이미 FVTPL로 측정됩니까?', '상호 배타적인 권리를 하나의 복합내재파생으로 평가할 필요가 있습니까?'],
         limitation='단순 금리 차이 또는 가치 비중만으로 분리 여부를 판정하지 않습니다. 자본요소 분리 전의 주계약 기준도 확인하십시오.',
         sources=[standard('1109', '4.3.3', '다음을 모두 충족하는 경우에만 내재파생상품을 주계약과 분리하여'), standard('1109', 'B4.3.4, B4.3.5(5)', appendix='b'),
                  book('2.1.3.1 분리 조건', 26, '아래의 조건을 모두 충족하는 경우에는 내재파생상품을 주계약과 분리하여'),
                  article('convertible-bonds-issuer', '발행자 전환사채 회계처리')]),
    dict(id='residual', title='자본요소는 어떤 금액에서 차감하나요?',
         trigger='issuer', fields=['conv_class', 'p_sep', 'k_sep'],
         explanation='자본 전환권이 있는 복합상품은 비자본 파생 특성을 포함한 부채요소를 먼저 측정하고 나머지를 자본에 배분합니다. 별도 금융상품을 포함한 거래대금과 복합상품 자체의 공정가치를 구분합니다.',
         questions=['전체 거래대금과 복합상품의 공정가치를 구분했습니까?', '부채요소에 필요한 풋·콜 조건이 포함되어 있습니까?', '발행가격을 공정가치로 볼 수 있는 근거가 있습니까?'],
         limitation='화면의 순차 차감 참고값이 이 회계 배분을 항상 충족하는 것은 아닙니다.',
         sources=[standard('1032', '31, 32', '복합금융상품 전체의 공정가치에서 별도로 결정된 부채요소의 금액을 차감한 나머지 금액을 자본요소에 배분한다.'),
                  book('사례 1108 부채요소 평가', 162, '부채요소에 배분되는 가치는 CU55(=CU57-CU2)이며, 자본요소에 배분되는 가치는 CU5이다.'),
                  book('3.3.2.2 평가방법 적용 시 유의사항', 89, '공정가치로서 신뢰할 수 있는 발행가액 수준을 별도로 추정하고')]),
    dict(id='third_party_call', title='제3자 콜옵션은 별도 상품인가요?',
         trigger='call', fields=['issuer_call', 'k_third', 'k_transfer', 'k_kind', 'k_method', 'k_split'],
         explanation='독립 양도 가능성 또는 다른 거래상대방의 존재를 확인해 회계단위를 정합니다. 발행자 상환권과 제3자가 전환사채를 취득하는 권리는 행사 효과가 다릅니다.',
         questions=['권리자·의무자·지정권자는 누구입니까?', '행사 후 사채가 소멸합니까, 제3자에게 이전됩니까?', '콜 평가방법과 의무보유 가정의 계약상 근거가 있습니까?'],
         limitation='실무사례의 특정 계약 결론을 모든 콜 조항에 일괄 적용하지 않습니다.',
         sources=[standard('1109', '4.3.1'),
                  book('사례 1105 제3자 지정 가능 콜옵션', 158, '해당 콜옵션은 별도의 금융상품으로 인식 및 측정해야 한다.'),
                  book('4.3.2 발행자 콜옵션', 145, '콜조항의 유무(with-without) 가치 비교'),
                  book('4.3.4 제3자 콜옵션', 146, '복합옵션 – 전환사채에 대한 콜옵션'),
                  article('compound-initial-recognition', '별도약정과 회계단위')]),
    dict(id='bdt', title='BDT 금리격자가 필요한가요?',
         trigger='bond', fields=['put_bdt', 'bdt_sig', 'model', 'conv_class'],
         explanation='신용등급, 보장수익률과 위험할인율의 차이, 금리의 시간가치 중요성 및 평가 목적을 함께 고려합니다. 필요하면 확정금리 방식과 금리격자 방식의 차이를 비교합니다.',
         questions=['금리 변동성이 평가에 미치는 영향이 중요합니까?', '금리 변동성의 관측자료와 추정기간을 설명할 수 있습니까?', '전체상품과 부채요소에 다른 모형을 써 생기는 차이를 검토했습니까?'],
         limitation='기존 앱의 5%p·20%·70% 등 수치 경계는 기준서상 의무 기준이 아닙니다. 이 화면은 해당 수치로 적합 여부를 결정하지 않습니다.',
         sources=[book('3.3.1.4 금리 모형 적용 검토', 87, '전문가적 판단에 따라 금리를 확률변수로 고려하는 이자율모형의 적용을'),
                  book('3.3.2.2 평가방법 적용 시 유의사항', 89, '두 모형의 이자율가정에 대한 차이로')]),
    dict(id='refixing', title='주가 리픽싱과 사건형 조정을 구분했나요?',
         trigger='bond', fields=['rfx_mode', 'carry', 'floor', 'K_cap', 'ipo_on'],
         explanation='정기 주가 조정, 저가 신주발행, 합병, 무상증자, IPO 조건은 서로 다른 조항입니다. 조정식·발생시점·상하한과 복수 사건의 적용 순서를 확인합니다.',
         questions=['계약서의 조정 사유를 빠짐없이 목록화했습니까?', '경로를 근사하는 방식의 차이를 확인했습니까?', '미래 발행가격·사건 발생은 확정된 사실입니까, 가정입니까?'],
         limitation='정기 리픽싱을 켰다고 모든 사건형 조정이 반영되는 것은 아닙니다. 사건 분석은 입력한 조건에 한정됩니다.',
         sources=[standard('1032', '16, 22'), article('convertible-bonds-issuer', '전환권 분류와 조정 조항')]),
    dict(id='dilution_backsolve', title='주당가치와 희석효과를 중복 반영하지 않았나요?',
         trigger='bond', fields=['S0', 'base_shares', 'dil_shares', 'bs_target', 'bs_net'],
         explanation='주당가치의 산정 방법에 희석효과가 이미 포함되어 있는지 확인합니다. 역산 시점 이후의 신규 발행·전환·상환은 후속 측정에서 다시 검토합니다.',
         questions=['역산에 사용한 거래가 정상적인 공정가치 거래입니까?', '거래가격에서 별도 권리의 영향을 구분했습니까?', '전기 이후 자본구조가 바뀌었습니까?'],
         limitation='단일 주식 수 비율 조정만으로 여러 종류주식의 청산 배분을 완전히 반영하지 못합니다.',
         sources=[book('3.4.1.3 내재적 희석 반영', 98, '변화된 자본요소에 대한 별도의 희석효과 고려가 필요할 수 있어 주의가 필요하다.')]),
    dict(id='method_change', title='전기와 방법이 달라진 이유가 있나요?',
         trigger='all', fields=['model', 'carry', 'k_method', 'k_split', 'put_bdt'],
         explanation='계약조건·자료·경제환경의 변화와 평가방법 변경 이유를 구분해서 기록합니다. 검토자 변경만으로 방법을 바꾸지 않습니다.',
         questions=['변경 이유와 새 방법이 더 적합한 근거가 있습니까?', '방법 변경 영향과 시장자료 변경 영향을 구분했습니까?'],
         limitation='원인별 변동 분석은 변경 순서에 영향을 받습니다. 합계 대사와 적용 순서를 함께 보관하십시오.',
         sources=[standard('1113', '65', '공정가치 측정을 위해 사용하는 가치평가기법은 일관되게 적용한다.'),
                  book('4.6.2 평가기법 선택 근거 문서화', 151, '판단 근거 등을 제시하고 관련 내용을 문서화할 필요가 있다.')]),
]


def applicable(topic, values):
    trigger = topic['trigger']
    return {'all': True, 'rcps': values.get('inst') == 'RCPS',
            'holder': values.get('view') == 'holder', 'issuer': values.get('view') == 'issuer',
            'bond': values.get('inst') != 'SHA',
            'call': bool(values.get('k_w', 0) or values.get('issuer_call', 0) or values.get('sha_call_e', 0))}[trigger]


def evidence_cards(case, include_all=False):
    values = case.effective()
    key = review_input_key(case)
    rows = []
    for topic in TOPICS:
        if not include_all and not applicable(topic, values):
            continue
        card = copy.deepcopy(topic)
        card['inputs'] = {k: values.get(k) for k in topic['fields']}
        decision = copy.deepcopy(case.judgments.get(topic['id'], {}))
        card['review'] = decision
        card['review_stale'] = bool(decision) and (decision.get('input_key') != key or decision.get('source_version') != source_version())
        card['status'] = '입력·근거 변경 후 재검토' if card['review_stale'] else decision.get('status', '미검토')
        rows.append(card)
    return rows


def judgment_record(case, status, conclusion, rationale, reviewer, contract_clause):
    import datetime as dt
    if status not in {'검토 중', '검토 완료', '해당 없음'}:
        raise ValueError('검토 상태를 확인하십시오.')
    if not reviewer.strip() or not rationale.strip() or (status == '검토 완료' and not conclusion.strip()):
        raise ValueError('검토자·근거를 입력하고, 완료 시 결론을 기록하십시오.')
    return dict(status=status, conclusion=conclusion, rationale=rationale, reviewer=reviewer,
                contract_clause=contract_clause, reviewed_at=dt.datetime.now(dt.timezone.utc).isoformat(),
                source_version=source_version(), input_key=review_input_key(case))


def source_version():
    return hashlib.sha256(json.dumps(TOPICS, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def review_input_key(case):
    payload = {k: getattr(case, k) for k in ['contract', 'market', 'method', 'sources', 'assumptions', 'additional_rights', 'contract_scenarios', 'cashflow_scenarios', 'market_evidence', 'imported_defaults']}
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def evidence_rows(case):
    rows = [['검토 주제', '상태', '적용 시 확인할 내용', '주의할 범위', '근거 구분', '자료명', '문단·쪽', '짧은 발췌', '원문 위치', '대조일', '검토자', '결론', '판단근거', '계약조항']]
    for card in evidence_cards(case):
        review = card['review']
        for src in card['sources']:
            rows.append([card['title'], card['status'], card['explanation'], card['limitation'], src['kind'], src['title'],
                         src['location'], src['quote'], src.get('url', f"사용자 보유 PDF {src.get('pdf_page')}쪽"), src['checked_on'],
                         review.get('reviewer', ''), review.get('conclusion', ''), review.get('rationale', ''), review.get('contract_clause', '')])
    return rows


def attach_evidence(data, case):
    """Append references to a legacy export without repricing or changing formulas."""
    import io
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(io.BytesIO(data))
    ws = wb.create_sheet('판단근거', 0)
    ws.append(['상세 기능 참고조서', '자동 작성 회계분류·분개·설명은 입력 가정에 따른 초안입니다. 아래 검토기록과 계약 원문을 확인하십시오.'])
    ws.append(['건명', case.name, '입력 식별값', case.fingerprint()])
    ws.append([])
    for row in evidence_rows(case): ws.append(row)
    ws.freeze_panes = 'A5'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 32
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str): cell.data_type = 's'
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='17365D')
    output = io.BytesIO(); wb.save(output)
    return output.getvalue()
