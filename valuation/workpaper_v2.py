"""Navigation, run identity and trace evidence for value and formula workbooks."""

import math
from copy import copy
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.utils import get_column_letter
from openpyxl.styles.cell_style import StyleArray
from . import legacy as L
from .explain import VERSION, node_trace, rate_rows, event_rows

# 수식 조서에서 고칠 수 있게 열어 두는 칸 — 엑셀에서 바꾸고 다시 계산해도 앱과 같은 값이 나오는 것이 시험으로
# 확인된 입력만이다 (tests/엑셀입력변경.py · tests/test_distributable_profit.py). 날짜·행사 일정·구간 수·모형·
# 분류처럼 격자의 모양을 정하는 칸과 기준금리·금리곡선은 앱에서 바꾼 뒤 다시 만든다.
EXCEL_INPUTS = {
    '평가기준일 주가',
    '변동성 σ',
    '보통주 배당수익률 δ',
    '표면이자율',
    '우선배당률 (계약)',
    '만기보장수익률',
    '만료 시 상환 보장수익률',
    '조기상환 보장수익률',
    '상환청구 보장수익률',
    # 주주간계약 가정 시트
    '평가기준일 주가 (원)',
    '변동성 σ (연)',
    '보통주 배당수익률 (연, 연속)',
    '풋 가격 가산율 (연)',
    # 콜 대상 주식수는 잠근다 — 같은 주식에 붙은 풋·콜 물량의 나눔과 계산 묶음이 조서를 만들 때 정해진다
    '가격 가산기간 (1 계약 개월 ÷ 12 / 0 실제 일수)',
}
DP_SHEET = '00 배당가능이익 상환'
# 배당가능이익 시트 — 넘긴 상환금 가산율·재원 사용 시작일(월·일)·만기 미상환 처리. 연도별 금액은 아래에서 따로 연다.
# 우선배당률(C6)은 가정 시트의 배당률에서 따라오는 수식이다.
DP_INPUT_CELLS = {(7, 3), (8, 3), (8, 4), (5, 6)}
DP_FIRST_YEAR_ROW = 10


def excel_input(ws, cell):
    """수식 조서에서 고칠 수 있는 노란 입력 칸인가."""
    if ws.title.startswith('가정'):
        return cell.column == 3 and ws.cell(cell.row, 2).value in EXCEL_INPUTS
    if ws.title == DP_SHEET:
        if (cell.row, cell.column) in DP_INPUT_CELLS:
            return True
        # 발생연도별 배당가능이익 금액 (발생연도 칸은 표의 열쇠라 잠근다)
        year = ws.cell(cell.row, 2).value
        return (
            cell.row >= DP_FIRST_YEAR_ROW and cell.column == 3 and isinstance(year, int) and not isinstance(year, bool)
        )
    return False


def add_workpaper_guide(wb, run, *, formula=False):
    from .report import calculation_table
    from .service import AMOUNT_LABELS

    t, s = run.terms, run.summary
    rows = [
        ['항목', '내용'],
        ['평가 건명', run.case.name],
        ['기준일', t.d_base],
        ['상품·회차', f'{t.inst} · {t.tranche}'],
        ['평가관점', '투자자' if t.view == 'holder' else '발행자'],
        ['평가 실행번호', s['run_id']],
        ['문서 개정번호', s['document_id']],
        ['계산 버전', VERSION],
        ['평가시각(UTC)', s['calculated_at']],
        ['문서 생성시각(UTC)', s['generated_at']],
        ['출력 형식', '수식 조서' if formula else '값 조서'],
        [
            '사용 범위',
            (
                '파란 글씨의 노란 칸만 고칠 수 있습니다: 주가·변동성·배당수익률, 표면이자율(우선배당률)·보장수익률, 배당가능이익 시트의 연도별 금액·가산율·재원 사용 시작일·만기 미상환 처리. 날짜·행사 일정·구간 수·모형·기준금리·금리곡선은 앱에서 바꾼 뒤 다시 만드십시오. 시트 보호는 실수 방지용이며 암호 보안이 아닙니다.'
                if formula
                else '평가 당시의 고정값입니다. 입력을 고쳐도 재계산되지 않습니다.'
            ),
        ],
        ['원금 100 기준', '단위가 다른 값을 직접 합산하지 마십시오. 총액 = 원금100 기준 × 평가대상 원금 ÷ 100.'],
        ['분해와 회계', 'TF 주식/현금 결제분, 권리별 증분, 회계상 인식액은 서로 다른 구분입니다.'],
        ['자료 보관', '평가 입력(case.json)·적용입력·시장자료·결과 기록을 조서 묶음과 함께 보관하십시오.'],
        [
            '구조조건 변경',
            '구간 수와 계약일 목록은 생성 시 고정됩니다. 이 파일에서 날짜나 계산간격을 수정하지 마십시오.',
        ],
        ['전체 재계산', '이 파일의 생성만으로 Microsoft Excel 전체 재계산 검증을 의미하지 않습니다.'],
        [
            '색과 유형',
            '노랑 바탕·파랑 글씨: 편집 가능한 입력 / 검정: 계산 / 초록: 다른 시트 참조. 회색 입력은 앱에서 변경합니다.',
        ],
        [
            '입력 변경 후',
            '계산대사 B열·계산근거·금리 적용내역·행사 지급내역·평가자 메모는 생성 당시 기록입니다. 입력 변경 후 Excel 계산값과 최초 앱값의 차이는 정상이며, 변경본의 검증·실행번호는 앱에서 재생성하세요.',
        ],
        ['배당가능이익 미입력 가정', run.case.sources.get('dp_missing_assumption', '해당 없음 또는 미확인')],
        ['목차', '시트명을 누르면 이동합니다.'],
    ]
    guide = calculation_table(wb, '00_안내', rows)
    wb.move_sheet(guide, offset=-wb.sheetnames.index(guide.title))
    guide.column_dimensions['A'].width = 28
    guide.column_dimensions['B'].width = 105
    for row in guide.iter_rows(min_row=2):
        guide.row_dimensions[row[0].row].height = 32
    r = guide.max_row + 1
    for ws in wb:
        if ws is guide:
            continue
        guide.cell(r, 1, ws.title).hyperlink = f"#'{ws.title.replace(chr(39),chr(39)*2)}'!A1"
        guide.cell(r, 1).style = 'Hyperlink'
        r += 1
        if not ws.freeze_panes:
            ws.freeze_panes = 'C2' if ws.max_column > 8 else 'A2'
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape' if ws.max_column > 6 else 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 0 if ws.max_column > 20 else 1
        ws.page_setup.fitToHeight = 0
        if not ws.print_title_rows:
            ws.print_title_rows = '1:2'
        # Existing A1 content is retained; attach a navigation link rather than insert rows.
        if ws['A1'].value is None:
            ws['A1'] = '목차로 돌아가기'
            ws['A1'].hyperlink = "#'00_안내'!A1"
            ws['A1'].style = 'Hyperlink'
    check_rows = [['대사 항목', '생성 당시 앱값(원금100)', '조서 계산값(원금100)', '차이', '출처']]
    refs = (
        []
        if L.is_sha(t) or not formula
        else [dict(zip(('name', 'sheet', 'cell', 'expected'), row)) for row in L.formula_key_cells(t, run.raw)]
    )
    for ref in refs:
        # Each record is supplied by the calculation exporter, not a guessed address.
        if ref['sheet'] not in wb:
            continue
        value = ref['expected']
        loc = f"'{ref['sheet'].replace(chr(39),chr(39)*2)}'!{ref['cell']}"
        check_rows.append([ref['name'], value, None, None, loc])
    check = calculation_table(wb, '계산대사', check_rows)
    for row, ref in enumerate([x for x in refs if x['sheet'] in wb], 2):
        loc = f"'{ref['sheet'].replace(chr(39),chr(39)*2)}'!{ref['cell']}"
        check.cell(row, 3, '=' + loc)
        check.cell(row, 4, f'=C{row}-B{row}')
        check.cell(row, 5).hyperlink = '#' + loc
    if len(check_rows) == 1:
        for key, val in s['amounts_100'].items():
            check.append([AMOUNT_LABELS[key], val, val, 0, '고정값 조서 · 수식 재계산 대사 아님'])
    rates = rate_rows(run)
    calculation_table(wb, '금리 적용내역', [list(rates[0])] + [list(x.values()) for x in rates])
    events = event_rows(run)
    if events:
        calculation_table(wb, '행사 지급내역', [list(events[0])] + [list(x.values()) for x in events])
    if not L.is_sha(t):
        trace = node_trace(run)
        rows = [
            ['항목', '생성 당시 실제 적용값(고정 기록)', '계산 의미'],
            ['평가 실행번호', s['run_id'], '모든 화면·조서 공통'],
            ['노드', '(0, 0)', '콜 반영 전 본체 트리'],
            ['단위', '원금 100', '주당·총액과 구별'],
            ['선택 가치', trace['value'], trace['kind']],
            ['전환가격', trace['conversion_price'], '입력과 적용조건 반영'],
        ]
        if not trace['terminal']:
            for key, lab, meaning in [
                ('q', '상승 확률', '하락 확률 = 1 − 상승 확률'),
                ('delta', '기간(년)', '격자 기간'),
                ('rf', '무위험 선도금리', '연속복리'),
                ('cr', '위험 선도금리', '연속복리'),
                ('df_rf', '무위험 할인계수', 'exp(−무위험 선도금리 × 기간)'),
                ('df_cr', '위험 할인계수', 'exp(−위험 선도금리 × 기간)'),
                ('coupon', '당기 이자·배당', '보유 시 현금결제분에 가산'),
                ('hold', '계속보유가치', '다음 노드 가치의 확률가중 할인'),
                ('reconciliation', '계산근거 대사차이', '설명 재구성값 − 엔진 보유값'),
            ]:
                rows.append([lab, trace[key], meaning])
            for direction in ('up', 'down'):
                for key in ('E', 'B', 'V', 'P'):
                    rows.append(
                        [
                            ('상승' if direction == 'up' else '하락') + ' 자식 ' + key,
                            trace[direction][key],
                            'E 주식결제 / B 현금결제 / V GS 가치 / P GS 전환확률',
                        ]
                    )
        calculation_table(wb, '계산근거', rows)
    # Inspect generated workbooks for stale external dependencies and invalid values.
    if wb._external_links:
        raise ValueError('조서에 외부 파일 연결이 발견되었습니다. 조서를 배포하지 않았습니다.')
    for name in wb.defined_names.values():
        if '[' in (name.attr_text or '') or '#REF!' in (name.attr_text or ''):
            raise ValueError('조서 정의명에 외부 연결 또는 잘못된 참조가 있습니다: ' + name.name)
    yellow_fill, fonts, prots, greys = {}, {}, {}, []

    def recolor(font_id, rgb):
        if (font_id, rgb) not in fonts:
            f = copy(wb._fonts[font_id])
            f.color = rgb
            fonts[font_id, rgb] = wb._fonts.add(f)
        return fonts[font_id, rgb]

    def prot(locked):
        if locked not in prots:
            prots[locked] = wb._protections.add(Protection(locked=locked))
        return prots[locked]

    def grey_fill():
        if not greys:
            greys.append(wb._fills.add(PatternFill('solid', fgColor='EEF1F6')))
        return greys[0]

    for ws in wb:
        for cell in ws._cells.values():
            v = cell.value
            # Legacy writers reuse StyleArray objects; isolate before altering protection.
            cell._style = copy(cell._style) if cell.has_style else StyleArray()
            st_ = cell._style
            if st_.fillId not in yellow_fill:
                f = wb._fills[st_.fillId].fgColor
                yellow_fill[st_.fillId] = f.type == 'rgb' and str(f.rgb)[-6:] == L.INPUT_FILL[-6:]
            yellow = yellow_fill[st_.fillId]
            editable = formula and yellow and cell.data_type != 'f' and excel_input(ws, cell)
            # 글꼴·보호·채우기는 같은 모양끼리 한 번만 만들어 번호로 붙인다 (칸마다 새로 만들면 조서 생성이 몇 배 느려진다).
            if yellow:
                st_.protectionId = prot(not editable)
                st_.fontId = recolor(st_.fontId, '0000FF' if editable else '526176')
                if not editable:
                    st_.fillId = grey_fill()
            if cell.data_type == 'f':
                st_.protectionId = prot(True)
                st_.fontId = recolor(st_.fontId, '008040' if '!' in v else '000000')
            if isinstance(v, (float, int)) and not isinstance(v, bool) and not math.isfinite(v):
                raise ValueError(f'유한하지 않은 조서 숫자: {ws.title}!{cell.coordinate}')
            if cell.data_type == 'f':
                if '#REF!' in v or '[' in v or len(v) - 1 > 8192:
                    raise ValueError(f'지원하지 않는 조서 수식: {ws.title}!{cell.coordinate}')
        if formula:
            ws.protection.sheet = True
            ws.protection.selectLockedCells = False
            ws.protection.selectUnlockedCells = False
            ws.protection.autoFilter = False
            ws.protection.sort = False
    wb.properties.creator = '복합금융상품 평가'
    wb.properties.lastModifiedBy = '복합금융상품 평가'
    wb.properties.title = run.case.name
    wb.properties.description = f'{VERSION} / 평가 {s["run_id"]} / 문서 {s["document_id"]}'
    # Add links to the evidence sheets created after the initial table of contents.
    included = {guide.cell(i, 1).value for i in range(1, guide.max_row + 1)}
    for ws in wb:
        if ws.title not in included:
            rr = guide.max_row + 1
            guide.cell(rr, 1, ws.title).hyperlink = f"#'{ws.title}'!A1"
            guide.cell(rr, 1).style = 'Hyperlink'
